"""A tiny workbook that probes the MM_LOOP forms the loop builder writes, for a real Mind run.

    python scripts/build_loop_forms_probe.py <out.xlsx>

Sheet "Loop":
  Settings block (A1:C7): first = 3, last = 7, column C = the loop value x 10 (a cell, MM_RESULT needs a cell),
    K = MM_LOOP("K", B3, 3, B2, B3)   the five-argument form the builder writes for a counter loop
                                       that does not start at 1 (as in a client model: IndexStart / IndexEnd),
    J = MM_LOOP("J", B3)              the plain form over a single integer cell (dims 1..n),
    sizeK = MM_DIMSIZE("K"), sizeJ = MM_DIMSIZE("J").
  Results block (E1:G8): k = 1..7, K10 = MM_RESULT($C$4, "K", k), J10 = MM_RESULT($C$5, "J", k).
Mind then shows, per block, which indices exist for K (3..7? or 1..5?) and what MM_RESULT returns for
each k: the index base of the five-argument form, which the KB leaves open.
Written through Excel with the MM for Excel add-in so the file carries cached values.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.excel_com import close_quietly, excel_session, open_for_write  # noqa: E402


def main() -> int:
    out = Path(sys.argv[1]).resolve()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loop"
    rows = [
        ("name", "value"),
        ("first", 3),
        ("last", 7),
        ("K", '=MM_LOOP("K",B3,3,B2,B3)', "=B4*10"),
        ("J", '=MM_LOOP("J",B3)', "=B5*10"),
        ("sizeK", '=MM_DIMSIZE("K")', None),
        ("sizeJ", '=MM_DIMSIZE("J")', None),
    ]
    ws["C1"] = "x10"
    for r, row in enumerate(rows, start=1):
        ws.cell(r, 1, row[0])
        ws.cell(r, 2, row[1])
        if len(row) > 2 and row[2] is not None:
            ws.cell(r, 3, row[2])
    # MM_RESULT needs a CELL (an expression such as $B$4*10 fails at runtime in Mind: "null reference")
    ws["E1"], ws["F1"], ws["G1"] = "k", "K10", "J10"
    for k in range(1, 8):
        ws.cell(k + 1, 5, k)
        ws.cell(k + 1, 6, f'=MM_RESULT($C$4,"K",{k})')
        ws.cell(k + 1, 7, f'=MM_RESULT($C$5,"J",{k})')
    tmp = Path(tempfile.mkdtemp(prefix="loop_forms_")) / "loop_forms_src.xlsx"
    wb.save(tmp)
    wb.close()
    with excel_session() as excel:
        book = open_for_write(excel, tmp)
        try:
            excel.CalculateFull()
            book.SaveAs(str(out), FileFormat=51)
            sheet = book.Worksheets("Loop")
            print("Excel (add-in, one pass): K =", sheet.Range("B4").Value, "| J =", sheet.Range("B5").Value, "| sizeK =", sheet.Range("B6").Value, "| sizeJ =", sheet.Range("B7").Value)
            print("F2:F8 =", [sheet.Range(f"F{r}").Value for r in range(2, 9)], "| G2:G8 =", [sheet.Range(f"G{r}").Value for r in range(2, 9)])
        finally:
            close_quietly(book)
    print("written", out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
