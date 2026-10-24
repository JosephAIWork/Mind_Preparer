"""Real recalculation adapter: drives the
locally-installed Excel via win32com to recalculate a copy of the workbook.
Setting calculation flags or saving through openpyxl is explicitly NOT a
valid recalculation per excel/recalculation.md -- this is the trusted
adapter that actually is one.

Never runs macros: AutomationSecurity is forced to msoAutomationSecurityForceDisable
before opening, per instructions/core.md ("never execute VBA... during inspection").
Always operates on a copy, never the user's original file.

Built on app/excel_com.py's `excel_session` like every other Excel
path in the app, with every COM proxy scoped inside an inner function.
Releasing Range/Worksheet proxies *after* `CoUninitialize` poisons the
apartment for a later session in the same process: a recalculation following
an Excel apply then fails with "The interface is unknown" (RPC_S_UNKNOWN_IF). Results carry `ran` so a
run that could not happen is never mistaken for a clean one.
"""
from __future__ import annotations

import re
import time
from pathlib import Path
from typing import Any, Callable

from .perf import record, stage
from .excel_com import MM_XLL_ENV, close_quietly, com_available, excel_session, mm_addin_loaded, open_for_write, save_in_place
from .formula_utils import called_functions, col_to_num, mask_strings, ref_text

# Excel COM's Range.Value returns error cells as these negative integer xlErr*
# constants, never as the display string -- confirmed empirically (Range.Text
# gives the string but only works cell-by-cell, not for a bulk range read).
XL_ERROR_CODES = {
    -2146826281: "#DIV/0!",
    -2146826246: "#N/A",
    -2146826259: "#NAME?",
    -2146826288: "#NULL!",
    -2146826252: "#NUM!",
    -2146826265: "#REF!",
    -2146826273: "#VALUE!",
}
MSO_AUTOMATION_SECURITY_FORCE_DISABLE = 3
XL_CELL_TYPE_FORMULAS = -4123  # Range.SpecialCells(xlCellTypeFormulas, xlErrors)
XL_SPECIAL_ERRORS = 16
ADDIN_NAME_MARKERS = ("mmforexcel", "milliman", "mind")


def _mmforexcel_loaded(excel) -> bool:
    """Is a Mind / MM for Excel add-in loaded in this Excel instance? `excel_session`
    registers the machine's MM for Excel XLL into its background Excel (automation never
    loads the user's startup add-ins by itself), and XLLs show up in `AddIns2` -- checked
    first. The COMAddIns / AddIns checks stay for an add-in installed another way. Without
    any of them, MM_ functions genuinely can't be evaluated by Excel, so a #NAME? on an MM_
    call is not the same kind of finding as one on a native formula."""
    if mm_addin_loaded(excel):
        return True
    try:
        for i in range(1, excel.COMAddIns.Count + 1):
            addin = excel.COMAddIns.Item(i)
            text = f"{addin.ProgId} {addin.Description}".lower()
            if addin.Connect and any(marker in text for marker in ADDIN_NAME_MARKERS):
                return True
    except Exception:
        pass
    try:
        for i in range(1, excel.AddIns.Count + 1):
            addin = excel.AddIns.Item(i)
            if addin.Installed and any(marker in addin.Name.lower() for marker in ADDIN_NAME_MARKERS):
                return True
    except Exception:
        pass
    return False


def _not_run(status: str, message: str) -> dict[str, Any]:
    return {"status": status, "message": message, "formula_errors": [], "addin_gap_errors": [], "ran": False, "mmforexcel_loaded": False}


def _array_formulas(sheet: Any, used_range: Any, cells: list[tuple[int, int, str, str]]) -> dict[tuple[int, int], str]:
    """(row, column) -> the address of the array formula that error cell is part of.

    `Range.HasArray` answers for a whole range at once: False when no cell of
    it belongs to an array formula, True when all do, Null when some do. So a
    sheet without array formulas costs one question, and on a sheet that has
    some only the rows that hold one are looked at cell by cell."""
    try:
        anywhere = used_range.HasArray
    except Exception:
        anywhere = None
    if anywhere is False:
        return {}
    by_row: dict[int, list[int]] = {}
    for row_no, col_no, _error, _formula in cells:
        by_row.setdefault(row_no, []).append(col_no)
    out: dict[tuple[int, int], str] = {}
    for row_no, cols in by_row.items():
        span = sheet.Range(sheet.Cells(row_no, min(cols)), sheet.Cells(row_no, max(cols)))
        try:
            in_row = span.HasArray
        except Exception:
            in_row = None
        finally:
            span = None
        if in_row is False:
            continue
        for col_no in cols:
            cell = sheet.Cells(row_no, col_no)
            try:
                if cell.HasArray:
                    arr = cell.CurrentArray
                    # Dynamic COM dispatch (no gencache) resolves .Address as a no-arg
                    # property, not a callable -- strip the '$' ourselves.
                    out[(row_no, col_no)] = str(arr.Address).replace("$", "")
                    arr = None
            finally:
                cell = None
    return out


def _errors_in_block(rng: Any, found: list[tuple[int, int, str, str]]) -> None:
    """The formula cells of `rng` whose value is an error, read as two arrays (values, formulas)."""
    values = rng.Value
    formulas = rng.Formula
    if values is None:
        return
    # Excel COM returns a scalar for a 1x1 range, a tuple-of-tuples otherwise.
    if not isinstance(values, tuple):
        values = ((values,),)
        formulas = ((formulas,),)
    top_row = int(rng.Row)
    top_col = int(rng.Column)
    # the error cells come out of the two arrays just read, and their
    # addresses are arithmetic. Asking Excel about each one (Cells, Address,
    # HasArray: three round trips a cell) takes ten minutes for a model with
    # 65,883 error cells -- twice, since the original is recalculated too.
    for r_idx, row in enumerate(values):
        for c_idx, val in enumerate(row):
            error_text = XL_ERROR_CODES.get(val) if isinstance(val, int) else None
            if error_text is None:
                continue
            formula_text = formulas[r_idx][c_idx]
            if isinstance(formula_text, str) and formula_text.startswith("="):
                found.append((top_row + r_idx, top_col + c_idx, error_text, formula_text))


def _error_cells(sheet: Any, used_range: Any) -> list[tuple[int, int, str, str]]:
    """(row, column, error, formula) of every formula cell of the sheet showing an error, row by row.

    Excel answers first how many cells of the used range hold an error (one SUMPRODUCT,
    computed inside Excel), and a sheet with none costs nothing more; otherwise Go To Special
    (formulas that are errors) hands over just those cells. Reading the whole used range, values and
    formulas, to find them takes 47 s on a workbook of 5.7 million cells, at every recalculation.
    Whenever Excel cannot answer either question the whole used range is read."""
    found: list[tuple[int, int, str, str]] = []
    try:
        address = str(used_range.Address)
        count = sheet.Evaluate(f"SUMPRODUCT(--ISERROR({address}))")
        count = int(count) if isinstance(count, (int, float)) and not isinstance(count, bool) and count >= 0 else None
    except Exception:
        count = None
    if count == 0:
        return found
    if count is not None:
        special = None
        try:
            special = used_range.SpecialCells(XL_CELL_TYPE_FORMULAS, XL_SPECIAL_ERRORS)
            for i in range(1, int(special.Areas.Count) + 1):
                area = special.Areas(i)
                try:
                    _errors_in_block(area, found)
                finally:
                    area = None
            found.sort(key=lambda e: (e[0], e[1]))
            return found
        except Exception:  # no formula error (the errors are constants), or Excel refused: read it all
            found = []
        finally:
            special = None
    _errors_in_block(used_range, found)
    return found


def _scan(excel, copy_path: Path, save: bool = True, wb: Any = None) -> dict[str, Any]:
    """Open, fully rebuild, save (unless `save` is False), and scan every used range for error values.
    Every COM proxy lives inside this function so it is released before the
    session quits Excel. `wb`: `copy_path` already open in `excel` (the session's Excel
    worker opened it); it is closed here as always."""
    mmforexcel_loaded = _mmforexcel_loaded(excel)
    try:
        if wb is None:
            wb = open_for_write(excel, copy_path)
        with stage("recalc_full_rebuild", file=copy_path.name):
            excel.CalculateFullRebuild()
        if save:
            save_in_place(wb, copy_path)

        formula_errors: list[dict[str, Any]] = []
        addin_gap_errors: list[dict[str, Any]] = []
        t_read = time.perf_counter()
        for sheet in wb.Worksheets:
            used_range = sheet.UsedRange
            sheet_name = sheet.Name
            found = _error_cells(sheet, used_range)
            arrays = _array_formulas(sheet, used_range, found) if found else {}
            for row_no, col_no, error_text, formula_text in found:
                entry = {"sheet": sheet_name, "cell": ref_text(col_no, row_no), "error": error_text, "formula": formula_text}
                array_addr = arrays.get((row_no, col_no))
                if array_addr and ":" in array_addr:
                    entry["array"] = array_addr  # multi-cell array formula: can only be changed as a whole
                is_mm_call = "MM_" in formula_text.upper()
                if error_text == "#NAME?" and is_mm_call and not mmforexcel_loaded:
                    addin_gap_errors.append(entry)
                else:
                    formula_errors.append(entry)
            used_range = None
            sheet = None
        record("recalc_read_errors", time.perf_counter() - t_read, file=copy_path.name, errors=len(formula_errors) + len(addin_gap_errors))
    finally:
        close_quietly(wb)
        wb = None

    # without the add-in every cell that reads an MM_ cell shows #NAME? as well (a Mind
    # loop written into a model: the whole calculation sheet). A #NAME? whose own functions are
    # all known -- Excel's or MM_ -- is downstream of the gap, not a defect of the workbook.
    downstream = 0
    if addin_gap_errors and not mmforexcel_loaded:
        kept: list[dict[str, Any]] = []
        for e in formula_errors:
            if e["error"] == "#NAME?" and not _unknown_functions(e["formula"]):
                addin_gap_errors.append({**e, "downstream": True})
                downstream += 1
            else:
                kept.append(e)
        formula_errors = kept
    if formula_errors:
        status = "ERROR"
        message = f"Recalculated; {len(formula_errors)} genuine formula error cell(s) found."
    elif addin_gap_errors:
        status = "NOT_SUPPORTED"
        direct = len(addin_gap_errors) - downstream
        why = (
            "the MM for Excel add-in is loaded but did not evaluate them (check the add-in's version against Mind's function list)"
            if mmforexcel_loaded
            else (
                "the MM for Excel add-in is not installed on this machine (the app loads it into its own Excel when it is; "
                f"install 'MM for Excel' from Milliman, or point {MM_XLL_ENV} at the XLL) -- Mind's own MM_001 rule expects "
                "this add-in whenever MM functions exist. This is not a workbook defect; it can't be verified as PASS or ERROR without it."
            )
        )
        message = f"Recalculated, but {direct} MM_ function call(s)" + (f" and {downstream} cell(s) that read them" if downstream else "") + f" show #NAME? because {why}"
    else:
        status = "PASS"
        message = "Recalculated successfully; no formula errors found." + (" MM for Excel was loaded: MM_ cells hold the first loop value (the loop itself runs in Mind)." if mmforexcel_loaded else "")
    return {
        "status": status,
        "message": message,
        "formula_errors": formula_errors,
        "addin_gap_errors": addin_gap_errors,
        "ran": True,
        "mmforexcel_loaded": mmforexcel_loaded,
    }


# --- root-cause grouping ---------------------------------------------------------
_REF_TOKEN_RE = re.compile(r"(?:(?:'[^']+'|[A-Za-z0-9_.]+)!)?\$?[A-Za-z]{1,3}\$?\d{1,7}(?::\$?[A-Za-z]{1,3}\$?\d{1,7})?")
_NUMBER_RE = re.compile(r"(?<![A-Za-z_])\d+(?:\.\d+)?")
_A1_CELL_RE = re.compile(r"^([A-Za-z]{1,3})(\d{1,7})$")

# How a cell of the original moved in the prepared version: (sheet, row, column) -> the same
# (app.numbers_check.Moves.forward -- rows / columns inserted, sheets renamed).
Forward = Callable[[str, int, int], tuple[str, int, int]]


def _moved_key(sheet: str, cell: str, forward: Forward | None) -> tuple[str, str]:
    """(sheet, A1 address) of an original cell -> where that cell is in the prepared
    version. A range or anything else that is not one plain address keeps its text
    (its sheet's new name still applies)."""
    cell = cell.replace("$", "")
    if forward is None:
        return sheet, cell
    m = _A1_CELL_RE.match(cell)
    if m is None:
        return forward(sheet, 1, 1)[0], cell
    sheet, row, col = forward(sheet, int(m.group(2)), col_to_num(m.group(1)))
    return sheet, ref_text(col, row)


def classify_errors(errors: list[dict[str, Any]], original_errors: list[dict[str, Any]], original_formulas_with_ref: set[tuple[str, str]] | None = None, forward: Forward | None = None) -> dict[str, Any]:
    """Split recalculation errors of a prepared version against the errors
    of the ORIGINAL after the same full recalculation (the honest baseline --
    cached values can be stale, e.g. a number saved before the column a
    formula pointed to was deleted). An error is pre-existing when the
    original errs at the same sheet+address, or has an erroring formula with
    the same text on the same sheet (row inserts shift addresses), or when the
    original's formula at that address already contained #REF! (broken
    whatever its cached value said).

    `forward` (how the preparation moved the original's cells: rows inserted
    above a grid, sheets renamed) puts each original address where that cell
    is now, so a title row above an erroring grid does not turn its old errors
    into new ones -- nor pass a new error that landed on an old address.
    With `forward` the addresses are exact and the formula-text match is not
    used: a good cell pushed onto an erroring cell's row reads the same text
    (=A10/B10) and would pass a new error as an old one."""
    by_addr = {_moved_key(e.get("sheet", ""), str(e.get("cell", "")), forward) for e in original_errors}
    by_formula = {(e.get("sheet", ""), str(e.get("formula", ""))) for e in original_errors if str(e.get("formula", "")).startswith("=")} if forward is None else set()
    with_ref = {_moved_key(sh, cell, forward) for sh, cell in (original_formulas_with_ref or set())}
    out, pre = [], 0
    for e in errors:
        key_a = (e.get("sheet", ""), str(e.get("cell", "")).replace("$", ""))
        key_f = (e.get("sheet", ""), str(e.get("formula", "")))
        preexisting = key_a in by_addr or key_a in with_ref or (key_f[1].startswith("=") and key_f in by_formula)
        pre += preexisting
        out.append({**e, "preexisting": preexisting})
    return {"compared": True, "message": None, "preexisting": pre, "new": len(errors) - pre, "errors": out}


def formulas_with_broken_refs(original_path: Path) -> set[tuple[str, str]]:
    """(sheet, address) of every formula in the original that contains #REF!."""
    from openpyxl import load_workbook

    out: set[tuple[str, str]] = set()
    t = time.perf_counter()
    try:
        wb = load_workbook(original_path, data_only=False, read_only=True, keep_links=False)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        v = c.value
                        if isinstance(v, str) and v.startswith("=") and "#REF!" in v:
                            out.add((ws.title, c.coordinate))
        finally:
            wb.close()
    except Exception:
        return out
    record("formulas_with_broken_refs", time.perf_counter() - t, file=Path(original_path).name)
    return out


def classify_against_original(original_path: Path, errors: list[dict[str, Any]], forward: Forward | None = None) -> dict[str, Any]:
    """Which recalculation errors were already errors in the ORIGINAL
    upload (the model's own #DIV/0! / #N/A, which Mind takes as they are) and
    which are new (introduced by the preparation, or revealed by it).

    An error cell counts as pre-existing when the original had an error at the
    same sheet+address, or -- because row inserts shift addresses -- an error
    with the same formula text on the same sheet; `forward` moves the original's
    addresses as classify_errors does (and then the text match is not used).
    Reads the original's cached values (what Excel last saved), never
    recalculates it."""
    from openpyxl import load_workbook

    by_addr: set[tuple[str, str]] = set()
    by_formula: set[tuple[str, str]] = set()
    try:
        wb_v = load_workbook(original_path, data_only=True, read_only=True, keep_links=False)
        wb_f = load_workbook(original_path, data_only=False, read_only=True, keep_links=False)
        try:
            for ws_v in wb_v.worksheets:
                ws_f = wb_f[ws_v.title]
                for row_v, row_f in zip(ws_v.iter_rows(), ws_f.iter_rows()):
                    for cv, cf in zip(row_v, row_f):
                        v = cv.value
                        if isinstance(v, str) and v.startswith("#") and (v.endswith("!") or v.endswith("?") or v in ("#N/A", "#NULL!")):
                            by_addr.add(_moved_key(ws_v.title, cv.coordinate, forward))
                            f = cf.value
                            if forward is None and isinstance(f, str) and f.startswith("="):
                                by_formula.add((ws_v.title, f))
        finally:
            wb_v.close()
            wb_f.close()
    except Exception as exc:  # the classification is advisory: an unreadable original marks every error as new
        return {"compared": False, "message": f"could not read the original for comparison: {str(exc)[:120]}", "preexisting": 0, "new": len(errors), "errors": errors}
    out = []
    pre = 0
    for e in errors:
        key_a = (e.get("sheet", ""), str(e.get("cell", "")).replace("$", ""))
        key_f = (e.get("sheet", ""), str(e.get("formula", "")))
        preexisting = key_a in by_addr or (key_f[1].startswith("=") and key_f in by_formula)
        pre += preexisting
        out.append({**e, "preexisting": preexisting})
    return {"compared": True, "message": None, "preexisting": pre, "new": len(errors) - pre, "errors": out}


def formula_signature(formula: str) -> str:
    """The *shape* of a formula: references -> REF, string literals -> "…",
    numeric constants -> N, whitespace removed, upper-cased. Two cells whose
    formulas share a signature almost always share a root cause."""
    if not formula:
        return ""
    masked = mask_strings(formula)
    sig = _REF_TOKEN_RE.sub("REF", masked)
    sig = re.sub(r'"[^"]*"', '"…"', sig)
    sig = _NUMBER_RE.sub("N", sig)
    return re.sub(r"\s+", "", sig).upper()


def _unknown_functions(formula: str) -> list[str]:
    from .validators.excel_functions import EXCEL_FUNCTIONS

    return sorted({fn for fn in called_functions(formula) if fn not in EXCEL_FUNCTIONS and not fn.startswith("MM_")})


def group_errors(formula_errors: list[dict[str, Any]], addin_gap_errors: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Cluster recalculation errors by root cause so the UI can offer 'fix all':
    #NAME? from the same unknown function(s); any error from formulas of the
    same shape; add-in-gap #NAME? by the MM_ function(s) involved."""
    groups: dict[tuple[str, str], dict[str, Any]] = {}

    def _add(key: tuple[str, str], cause: str, kind: str, entry: dict[str, Any]) -> None:
        g = groups.setdefault(key, {"error": key[0], "cause": cause, "kind": kind, "signature": key[1], "cells": []})
        c = {"sheet": entry.get("sheet", ""), "cell": entry.get("cell", ""), "formula": entry.get("formula", ""), "error": entry.get("error", key[0])}
        if entry.get("array"):
            c["array"] = entry["array"]
        g["cells"].append(c)

    for e in formula_errors:
        formula = str(e.get("formula") or "")
        error = str(e.get("error") or "")
        unknown = _unknown_functions(formula) if error == "#NAME?" else []
        if unknown:
            _add((error, "fn:" + ",".join(unknown)), f"#NAME? -- unknown function(s) {', '.join(unknown)} (not an Excel or MM_ function; VBA/add-in or typo)", "formula", e)
        else:
            sig = formula_signature(formula)
            _add((error, "sig:" + sig), f"{error} in formulas of the form {sig or '(empty)'}", "formula", e)
    for e in addin_gap_errors:
        if e.get("downstream"):
            _add(("#NAME?", "mm:downstream"), "#NAME? because the cell reads an MM_ function cell, and MM_ functions need the MMForExcel add-in -- the same environment gap, not a workbook defect", "addin_gap", {**e, "error": "#NAME?"})
            continue
        formula = str(e.get("formula") or "")
        mm = sorted({fn for fn in called_functions(formula) if fn.startswith("MM_")})
        _add(("#NAME?", "mm:" + ",".join(mm)), f"#NAME? because {', '.join(mm) or 'an MM_ function'} needs the MMForExcel add-in -- an environment gap on this machine, not a workbook defect", "addin_gap", {**e, "error": "#NAME?"})

    out = []
    for i, g in enumerate(sorted(groups.values(), key=lambda g: (-len(g["cells"]), g["error"], g["signature"]))):
        arrays = sorted({f"{c['sheet']}!{c['array']}" for c in g["cells"] if c.get("array")})
        if arrays:
            g["arrays"] = arrays
            g["cause"] += f" -- array formula {', '.join(arrays[:4])}{' ...' if len(arrays) > 4 else ''} (Ctrl+Shift+Enter; it can only be changed as a whole)"
        out.append({"id": f"g{i + 1}", **g, "count": len(g["cells"])})
    return out


def recalculate_in(excel: Any, copy_path: Path, wb: Any = None) -> dict[str, Any]:
    """recalculate() in an Excel the caller owns, `copy_path` possibly already open in it as `wb`
    (the session's Excel worker, ahead of the Recalculate step). Errors are raised."""
    copy_path = Path(copy_path).resolve()
    with stage("recalculate", file=copy_path.name, worker=True):
        return _scan(excel, copy_path, save=True, wb=wb)


def recalculate(copy_path: Path, save: bool = True) -> dict[str, Any]:
    """Recalculates `copy_path` in place (it's already a disposable copy)
    and reports any formula-error cells found afterward. `save=False`: the recalculated
    values are not written back -- for a copy nothing reads again (the original's baseline)."""
    if not com_available():
        return _not_run("NOT_SUPPORTED", "pywin32 is not available in this environment; Excel COM recalculation cannot run here.")
    copy_path = Path(copy_path).resolve()
    if not copy_path.is_file():
        return _not_run("ERROR", f"file not found: {copy_path}")
    try:
        with stage("recalculate", file=copy_path.name), excel_session() as excel:
            return _scan(excel, copy_path, save=save)
    except Exception as exc:  # pragma: no cover -- requires a real Excel install to exercise
        return _not_run("ERROR", f"Excel COM recalculation failed: {exc}")
