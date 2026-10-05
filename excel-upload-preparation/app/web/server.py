"""FastAPI backend for the Figma-built "Mind Ready" front-end
(`Mind Copilot Skill/FigmaOutput`, React + Vite).

One endpoint per function of the front-end's `src/services/api.ts`, on top of
the same engine the Streamlit UI uses (app/modes, app/prep, app/chat_context,
app/recalc, app/excel_report, app/change_apply). Nothing here re-implements
a rule or a change; it only maps the engine's dicts onto the front-end's
TypeScript contracts (docs/FIGMA_UI_PROMPT.md) and keeps a per-session
**version lineage**: every write is a new file, verified to open in Excel,
with its change log.

Sessions live in memory (this is a local, single-user tool); their files
live in a temp directory per session. The built front-end (`dist/`) is
served from the same process with an SPA fallback, so one URL does it all:

    python -m app.web.server            (or run_mind_ready_web.bat)
    http://localhost:8600
"""
from __future__ import annotations

import json
import os
import queue
import shutil
import sys
import tempfile
import threading
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from fastapi import Body, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse

PACKAGE_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PACKAGE_ROOT))

from app.autofix import AutoFixConfig, run_autofix  # noqa: E402
from app.autofix_advisor import make_advisor  # noqa: E402
from app.change_apply import convert_output_format  # noqa: E402
from app.chat_context import answer_question, recalculation_context  # noqa: E402
from app.config import load_config  # noqa: E402
from app.excel_com import com_available  # noqa: E402
from app.excel_report import build_report_workbook, build_standalone_report  # noqa: E402
from app.inventory import cell_window, has_vba_project, make_immutable_copy  # noqa: E402
from app.llm import DEFAULT_MODEL_ID, extract_formula, llm_available, suggest_formula_fix  # noqa: E402
from app.modes import fix_incompatible_formulas, plan_mode, prep_mind_loops, structure_fix  # noqa: E402
from app.grid_naming import build_grid_context, suggest_names  # noqa: E402
from app.mind_loop import LoopConfig, run_loop, summarize  # noqa: E402
from app.grids import all_grids, json_safe  # noqa: E402
from app.models import Status, aggregate_status  # noqa: E402
from app.readiness import STATUS_RANK, apply_outcome, apply_record, compute_readiness, plan_counts, prep_progress, recalc_state  # noqa: E402
from app.prep import ASSISTANT_OPS, STRUCTURAL_OPS, apply_operations, deterministic_name, is_weak_name, plan_actions, plan_create_grid_titles, plan_named_areas, standalone_labels  # noqa: E402
from app.validators.kb import documented_flags  # noqa: E402
from app.recalc import classify_against_original, classify_errors, formulas_with_broken_refs, group_errors, recalculate  # noqa: E402
from app.prep import _cells_of as _array_cells  # noqa: E402
from app.prep import array_size_hint  # noqa: E402
from app.rules_engine import RulesEngine  # noqa: E402
from app.progress import APPLY_STAGE_TITLES, STAGE_TITLES, overall_progress  # noqa: E402
from app.sizing import MB, size_gate  # noqa: E402

MODE_MAP = {
    "plan": ("Plan (analyze everything)", plan_mode),
    "prep_loops": ("Prep Mind Loops", prep_mind_loops),
    "fix_formulas": ("Fix Incompatible Formulas", fix_incompatible_formulas),
    "structure_fix": ("Structure Fix", structure_fix),
}
ALLOWED_SUFFIXES = {".xlsx", ".xlsm", ".xlsb"}
DEFAULT_DIST = PACKAGE_ROOT.parent / "FigmaOutput" / "dist"
DIST_DIR = Path(os.environ.get("MIND_READY_DIST", str(DEFAULT_DIST)))
GRID_FIELDS = (
    "sheet", "title_cell", "title", "name", "display_name", "flags", "flag_names", "anchor", "ref",
    "first_row", "last_row", "first_col", "last_col", "n_rows", "n_cols", "header_values", "header_is_all_text",
    "inner_title_cells", "formula_count",
)

app = FastAPI(title="Mind Ready API", version=(PACKAGE_ROOT / "VERSION").read_text(encoding="utf-8").strip() if (PACKAGE_ROOT / "VERSION").exists() else "dev")
_ENGINE: RulesEngine | None = None


def engine() -> RulesEngine:
    global _ENGINE
    if _ENGINE is None:
        _ENGINE = RulesEngine()
    return _ENGINE


@dataclass
class Session:
    id: str
    work_dir: Path
    mode: str
    versions: list[dict[str, Any]] = field(default_factory=list)
    paths: dict[str, Path] = field(default_factory=dict)  # version id -> file
    files: dict[str, Path] = field(default_factory=dict)  # download name -> file
    current_version_id: str = ""
    result: dict[str, Any] | None = None
    plan: list[dict[str, Any]] = field(default_factory=list)
    recalc: dict[str, Any] | None = None  # last real recalculation of the current version
    recalc_version_id: str | None = None
    recalc_path: Path | None = None  # the recalculated copy: freshest cell values for the workbook view
    grid_names: dict[str, str] = field(default_factory=dict)  # "Sheet!Ref" -> assistant-proposed grid name
    created_at: float = field(default_factory=time.time)
    # 1.6.6 -- upload size gate + scan status
    size: dict[str, Any] | None = None  # app.sizing.size_gate of the uploaded file
    ignore_sheets: list[str] = field(default_factory=list)  # sheets the user chose not to scan
    pending: bool = False  # uploaded and inspected, not analysed yet (deferred upload)
    scan: dict[str, Any] = field(default_factory=dict)  # live status of the current/last scan
    # 1.7.2 -- one entry per analysis: how much prep work each version still needs (the Prep gauge)
    plan_history: list[dict[str, Any]] = field(default_factory=list)
    # 1.7.3 -- one entry per Apply: the changes really written into the file (the applied-changes gauge)
    apply_history: list[dict[str, Any]] = field(default_factory=list)
    # 1.7.2 -- the ORIGINAL after a full Excel recalculation: its error cells and its #REF! formulas
    # (computed once per session, the baseline every later recalculation is compared with)
    original_baseline: dict[str, Any] | None = None
    # 1.7.2 -- live status of the current/last Apply (GET /status, field "apply")
    apply_status: dict[str, Any] = field(default_factory=dict)
    lock: threading.Lock = field(default_factory=threading.Lock, repr=False, compare=False)

    def readiness(self) -> dict[str, Any] | None:
        """The verdict for the current version (None before the first analysis)."""
        if not self.result:
            return None
        return compute_readiness(self.result["validation_report"], self.plan, self.recalc, self.current_version_id)

    @property
    def current_path(self) -> Path:
        return self.paths[self.current_version_id]


SESSIONS: dict[str, Session] = {}


# --- helpers -------------------------------------------------------------------------
def _session(session_id: str) -> Session:
    s = SESSIONS.get(session_id)
    if s is None:
        raise HTTPException(status_code=404, detail=f"unknown session {session_id}")
    return s


def _register_file(s: Session, path: Path) -> str:
    """Download name for a file, unique within the session."""
    path = Path(path)
    name = path.name
    if s.files.get(name) not in (None, path):
        stem, suffix = path.stem, path.suffix
        n = 2
        while s.files.get(f"{stem}_v{n}{suffix}") not in (None, path):
            n += 1
        name = f"{stem}_v{n}{suffix}"
    s.files[name] = path
    return name


def _add_version(
    s: Session, path: Path, source: str, label_body: str, change_log: list[dict[str, Any]], verified: bool | None, sha256: str, minor: bool = False
) -> dict[str, Any]:
    """1.7.4: versions are numbered major.minor. A fix made on the Recalculate
    step (`minor`) is the next minor of the current major (v2 -> v2.1 -> v2.2);
    everything else -- and a recalculation after such fixes -- starts the next
    major (v3). Ids follow the number: ver-002, ver-002.1, ..."""
    current = next((v for v in s.versions if v["id"] == s.current_version_id), None)
    if minor and current is not None:
        major = current.get("major", 1)
        n_minor = 1 + max((v.get("minor", 0) for v in s.versions if v.get("major") == major), default=0)
    else:
        major = 1 + max((v.get("major", 0) for v in s.versions), default=0)
        n_minor = 0
    number = f"{major}.{n_minor}" if n_minor else f"{major}"
    vid = f"ver-{major:03d}" + (f".{n_minor}" if n_minor else "")
    download_name = _register_file(s, path)
    version = {
        "id": vid,
        "major": major,
        "minor": n_minor,
        "label": f"v{number} — {label_body}",
        "file_name": download_name,
        "sha256": sha256[:12],
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "source": source,
        "change_log": [_public_op(o) for o in change_log],
        "verified_opens_in_excel": verified,
    }
    s.versions.append(version)
    s.paths[vid] = Path(path)
    s.current_version_id = vid
    return version


def _public_op(o: dict[str, Any]) -> dict[str, Any]:
    out = {k: v for k, v in o.items() if k != "cells"}
    if "cells" in o:
        out["cells"] = list(o["cells"])[:50]
    return out


def _summary(analysis: dict[str, Any]) -> dict[str, Any]:
    wb = analysis["workbooks"][0]
    f = analysis["features"]
    sheets = []
    for sh in wb["sheets"]:
        sheets.append(
            {
                "name": sh["name"],
                "state": sh["state"],
                "dimensions": sh.get("dimensions") or "",
                "grids": [{k: g.get(k) for k in GRID_FIELDS} for g in sh["grids"]],
                "standalone_text_cells": [{"cell": c["cell"], "text": str(c["text"])} for c in sh.get("standalone_text_cells", [])],
                "formula_count": sh.get("formula_count", 0),
                "array_formula_count": sh.get("array_formula_count", 0),
                "protected": bool(sh.get("protected")),
            }
        )
    return {
        "file_name": f"{wb.get('file_name')}.{wb.get('file_type')}",
        "file_type": wb.get("file_type"),
        "sha256": (wb.get("source_sha256") or analysis["source"].get("sha256") or "")[:12],
        "application": f.get("application"),
        "app_version": f.get("app_version"),
        "has_vba": bool(f.get("has_vba")),
        "sheet_count": f.get("sheet_count", 0),
        "ignored_sheet_count": f.get("ignored_sheet_count", 0),
        "total_sheet_count": f.get("total_sheet_count", f.get("sheet_count", 0)),
        "ignored_sheets": [{"name": s["name"], "state": s["state"]} for s in wb.get("ignored_sheets", [])],
        "hidden_sheet_count": f.get("hidden_sheet_count", 0),
        "grid_count": f.get("grid_count", 0),
        "flagged_grid_count": f.get("flagged_grid_count", 0),
        "formula_count": f.get("formula_count", 0),
        "array_formula_count": f.get("array_formula_count", 0),
        "defined_name_count": f.get("defined_name_count", 0),
        "external_link_part_count": f.get("external_link_part_count", 0),
        "mm_functions_used": f.get("mm_functions_used", {}),
        "native_function_usage": dict(sorted(f.get("native_function_usage", {}).items(), key=lambda kv: -kv[1])[:40]),
        "sheets": sheets,
    }


def _delta(previous: dict[str, Any] | None, current: dict[str, Any], version_id: str) -> dict[str, Any]:
    """What changed between two analyses of the same session: findings fixed
    (now PASS), improved (better status or a changed message while still not
    PASS), regressed, and the status counts side by side -- so a fix never
    just *vanishes* from a filtered table."""
    cur = {f["rule_id"]: f for f in current["findings"]}
    if not previous:
        return {"version_id": version_id, "previous_version_id": None, "fixed": [], "improved": [], "regressed": [], "previous_counts": {}, "counts": current["summary"]["status_counts"]}
    prev = {f["rule_id"]: f for f in previous["findings"]}
    fixed, improved, regressed = [], [], []
    for rid, f in cur.items():
        p = prev.get(rid)
        if p is None:
            continue
        before, after = STATUS_RANK.get(p["status"], 0), STATUS_RANK.get(f["status"], 0)
        if p["status"] != "PASS" and f["status"] == "PASS":
            fixed.append({"rule_id": rid, "from": p["status"], "to": f["status"]})
        elif after < before or (after == before and f["status"] != "PASS" and f["message"] != p["message"]):
            improved.append({"rule_id": rid, "from": p["status"], "to": f["status"], "note": "status improved" if after < before else "same status, details changed"})
        elif after > before:
            regressed.append({"rule_id": rid, "from": p["status"], "to": f["status"]})
    return {
        "version_id": version_id,
        "previous_version_id": previous.get("_version_id"),
        "fixed": fixed,
        "improved": improved,
        "regressed": regressed,
        "previous_counts": previous["summary"]["status_counts"],
        "counts": current["summary"]["status_counts"],
    }


# --- scan status (1.6.6) ------------------------------------------------------------------
def _scan_start(s: Session, needs_convert: bool) -> None:
    now = time.time()
    s.scan = {
        "state": "running",
        "stage": None,
        "title": "Starting",
        "message": "Starting",
        "fraction": None,
        "overall": 0.0,
        "sheet": None,
        "rule_id": None,
        "started_at": now,
        "updated_at": now,
        "finished_at": None,
        "stage_started": now,
        "stages": [],
        "error": None,
        "version_id": s.current_version_id,
        "needs_convert": needs_convert,
        "ignore_sheets": list(s.ignore_sheets),
    }


def _progress_for(s: Session):
    def cb(stage: str, message: str, fraction: float | None = None, **facts: Any) -> None:
        sc = s.scan
        now = time.time()
        if sc.get("stage") != stage:
            if sc.get("stage"):
                sc["stages"].append({"stage": sc["stage"], "seconds": round(now - sc["stage_started"], 2)})
            sc["stage"], sc["stage_started"] = stage, now
        sc.update(
            title=STAGE_TITLES.get(stage, stage),
            message=message,
            fraction=fraction,
            updated_at=now,
            sheet=facts.get("sheet"),
            rule_id=facts.get("rule_id"),
        )
        sc["overall"] = overall_progress(stage, fraction, bool(sc.get("needs_convert")))

    return cb


def _scan_finish(s: Session, error: str | None = None) -> None:
    sc = s.scan
    now = time.time()
    if sc.get("stage"):
        sc["stages"].append({"stage": sc["stage"], "seconds": round(now - sc.get("stage_started", now), 2)})
    sc.update(state="error" if error else "ready", finished_at=now, updated_at=now, error=error, overall=sc.get("overall", 0.0) if error else 1.0)
    sc["title"] = "Failed" if error else "Done"
    sc["message"] = error or f"Scan finished in {now - sc.get('started_at', now):.1f} s"
    sc["stage"] = sc.get("stage") if error else "done"


def _scan_snapshot(s: Session) -> dict[str, Any]:
    sc = dict(s.scan) if s.scan else {"state": "pending" if s.pending else "idle", "stage": None, "title": None, "message": None, "fraction": None, "overall": 0.0, "stages": [], "error": None}
    now = time.time()
    started = sc.get("started_at")
    finished = sc.get("finished_at")
    sc["elapsed_s"] = round(((finished or now) - started), 1) if started else 0.0
    sc.pop("stage_started", None)
    sc["state"] = sc.get("state") or ("pending" if s.pending else "idle")
    sc["version_id"] = sc.get("version_id") or s.current_version_id
    sc["apply"] = _apply_snapshot(s)
    return sc


# --- apply status (1.7.2) -----------------------------------------------------------------
def _apply_start(s: Session, total: int, reanalyze: bool) -> None:
    now = time.time()
    s.apply_status = {
        "state": "running",
        "stage": None,
        "title": "Starting",
        "message": "Starting",
        "fraction": None,
        "done": 0,
        "total": total,
        "reanalyze": reanalyze,
        "started_at": now,
        "updated_at": now,
        "finished_at": None,
        "stage_started": now,
        "stages": [],
        "applied": None,
        "failed": None,
        "error": None,
        "from_version_id": s.current_version_id,
    }


def _apply_progress_for(s: Session):
    def cb(stage: str, message: str, fraction: float | None = None, **facts: Any) -> None:
        st = s.apply_status
        now = time.time()
        if st.get("stage") != stage:
            if st.get("stage"):
                st["stages"].append({"stage": st["stage"], "seconds": round(now - st["stage_started"], 2)})
            st["stage"], st["stage_started"] = stage, now
        st.update(title=APPLY_STAGE_TITLES.get(stage, stage), message=message, fraction=fraction, updated_at=now)
        if "done" in facts:
            st["done"] = facts["done"]
        if "total" in facts:
            st["total"] = facts["total"]

    return cb


def _apply_finish(s: Session, res: dict[str, Any] | None = None, error: str | None = None) -> None:
    st = s.apply_status
    now = time.time()
    if st.get("stage"):
        st["stages"].append({"stage": st["stage"], "seconds": round(now - st.get("stage_started", now), 2)})
    st.update(state="error" if error else "done", finished_at=now, updated_at=now, error=error, fraction=None if error else 1.0)
    if res is not None:
        st["applied"], st["failed"] = len(res.get("applied", [])), len(res.get("failed", []))
    st["title"] = "Failed" if error else "Done"
    st["message"] = error or f"{st.get('applied') or 0} change(s) applied" + (f", {st['failed']} not applied" if st.get("failed") else "")
    st["stage"] = st.get("stage") if error else "done"


def _apply_snapshot(s: Session) -> dict[str, Any] | None:
    if not s.apply_status:
        return None
    st = dict(s.apply_status)
    started, finished = st.get("started_at"), st.get("finished_at")
    st["elapsed_s"] = round(((finished or time.time()) - started), 1) if started else 0.0
    st.pop("stage_started", None)
    return st


def _tracked_apply(s: Session, ops: list[dict[str, Any]], reanalyze: bool) -> dict[str, Any]:
    """apply_operations on the current version, every stage and every written
    operation reported into `s.apply_status` (GET /status, field "apply")."""
    _apply_start(s, len(ops), reanalyze)
    try:
        res = apply_operations(s.current_path, s.work_dir / "apply", ops, progress=_apply_progress_for(s))
    except Exception as exc:
        _apply_finish(s, error=f"apply failed: {exc}")
        raise
    if res["status"] == "NOT_APPLICABLE":
        _apply_finish(s, error=res.get("message", "nothing to apply"))
        raise HTTPException(status_code=422, detail=res.get("message", "nothing to apply"))
    _apply_finish(s, res)
    return res


def _status_counts(findings: list[dict[str, Any]]) -> dict[str, int]:
    out: dict[str, int] = {}
    for f in findings:
        st = str(f.get("status") or "UNKNOWN")
        out[st] = out.get(st, 0) + 1
    return out


def _sync_recalc_finding(s: Session) -> None:
    """1.7.4: READY-001 ("successful recalculation") is NOT_SUPPORTED in every
    analysis pass -- the pass never runs Excel -- and, worst status winning,
    it kept the whole report at NOT_SUPPORTED even after a clean Recalculate.
    When the current version has a real recalculation, the finding says what
    it found (PASS when no new error, ERROR otherwise) and the report's status
    and summary are recomputed."""
    if not s.result:
        return
    rc = recalc_state(s.recalc, s.current_version_id)
    if not rc["ran"]:
        return
    report = s.result["validation_report"]
    finding = next((f for f in report["findings"] if f.get("rule_id") == "READY-001"), None)
    if finding is None:
        return
    short = next((v["label"].split(" — ")[0] for v in s.versions if v["id"] == s.current_version_id), s.current_version_id)
    finding["status"] = "PASS" if rc["clean"] else "ERROR"
    finding["message"] = f"Recalculated {short} in Excel: {rc['message']}."
    statuses = [Status(f["status"]) for f in report["findings"]]
    report["status"] = aggregate_status(statuses).value
    summary = report.setdefault("summary", {})
    summary["status_counts"] = _status_counts(report["findings"])
    summary["not_supported_rule_ids"] = [f["rule_id"] for f in report["findings"] if f["status"] == "NOT_SUPPORTED"]
    summary["blocking_rule_ids"] = [f["rule_id"] for f in report["findings"] if f["status"] in ("ERROR", "REQUIRES_USER_INPUT", "NOT_SUPPORTED")]


def _analyze(s: Session, path: Path, progress=None) -> dict[str, Any]:
    label, module = MODE_MAP[s.mode]
    work = s.work_dir / "analysis" / s.current_version_id
    previous = s.result["validation_report"] if s.result else None
    result = module.run(path, work, load_config(), engine=engine(), ignore_sheets=s.ignore_sheets or None, progress=progress)
    if progress is not None:
        progress("plan", "Planning the fixes", None)
    s.plan = plan_actions(result["workbook_analysis"], result["validation_report"], s.grid_names or None)
    report = result["validation_report"]
    # "Fix available" means the current prep plan has an operation for the rule.
    fixable = {rid for a in s.plan if a["count"] > 0 for rid in a["rule_ids"]} | {o["rule_id"] for a in s.plan for o in a["operations"]}
    for f in report["findings"]:
        f["correction_available"] = f["rule_id"] in fixable
    report["_version_id"] = s.current_version_id
    s.result = result
    _sync_recalc_finding(s)  # a re-analysis of a version already recalculated keeps that result
    delta = _delta(previous, report, s.current_version_id)
    summary = _summary(result["workbook_analysis"])
    summary["size"] = s.size
    # 1.7.2: the verdict and the Prep gauge travel with every analysis
    readiness = s.readiness()
    counts = plan_counts(s.plan)
    s.plan_history.append({
        "version_id": s.current_version_id,
        "blocking_ops": readiness["prep"]["blocking_ops"] if readiness else 0,
        "optional_ops": readiness["prep"]["optional_ops"] if readiness else 0,
        "total_ops": sum(counts.values()),
        "blocking_findings": readiness["blocking_count"] if readiness else 0,
        # the repair gauge: every finding still open, by status, so the Fix panel
        # and Findings can show the run going down even when Prep has no work left
        "open_findings": sum(1 for f in report["findings"] if f.get("status") != "PASS"),
        "status_counts": _status_counts(report["findings"]),
        "by_action": counts,
    })
    return {
        "summary": summary,
        "report": {k: v for k, v in report.items() if k != "_version_id"},
        "plan": s.plan,
        "delta": delta,
        "readiness": readiness,
        "prep_progress": prep_progress(s.plan_history, s.apply_history),
    }


def _convert(s: Session, path: Path, progress) -> Path:
    """An .xlsb is converted through Excel first (openpyxl cannot read it);
    the converted file becomes a version of its own."""
    target = "xlsm" if has_vba_project(path) else "xlsx"
    progress("convert", f"Converting .xlsb to .{target} through Excel (openpyxl cannot read .xlsb)", None)
    conv = convert_output_format(path, s.work_dir / "convert", target)
    if conv["status"] != "APPLIED":
        raise HTTPException(status_code=500, detail=conv.get("message", "conversion failed"))
    out = Path(conv["output_path"])
    _add_version(s, out, "convert", f"converted .xlsb to .{target} via Excel", [], None, conv["change_log_entry"]["output_sha256"])
    return out


def _run_scan(s: Session, path: Path | None = None) -> dict[str, Any]:
    """One scan of `path` (default: the current version) with the session's
    ignore list, converting an .xlsb first, reporting every stage into
    `s.scan` (GET /status) -- one at a time per session."""
    if not s.lock.acquire(blocking=False):
        raise HTTPException(status_code=409, detail="a scan is already running for this session")
    try:
        path = Path(path or s.current_path)
        needs_convert = path.suffix.lower() == ".xlsb"
        _scan_start(s, needs_convert)
        progress = _progress_for(s)
        try:
            if needs_convert:
                path = _convert(s, path, progress)
            out = _analyze(s, path, progress)
        except HTTPException as exc:
            _scan_finish(s, error=str(exc.detail))
            raise
        except Exception as exc:
            _scan_finish(s, error=f"analysis failed: {exc}")
            raise HTTPException(status_code=500, detail=f"analysis failed: {exc}") from exc
        _scan_finish(s)
        s.pending = False
        current = next((v for v in s.versions if v["id"] == s.current_version_id), s.versions[-1])
        return {**out, "version": current, "versions": s.versions, "size": s.size, "scan": _scan_snapshot(s)}
    finally:
        s.lock.release()


def _parse_ignore(s: Session, raw: Any) -> list[str]:
    """The sheets to skip, validated against the sheet list read from the
    package (when it could be read): unknown names are refused, and at least
    one sheet must remain to scan."""
    if raw in (None, ""):
        return []
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except json.JSONDecodeError:
            raw = [x.strip() for x in raw.split(",") if x.strip()]
    if not isinstance(raw, list) or not all(isinstance(x, str) for x in raw):
        raise HTTPException(status_code=422, detail="ignore_sheets must be a list of sheet names")
    names = list(dict.fromkeys(x for x in raw if x))
    known = [sh["name"] for sh in (s.size or {}).get("sheets", [])]
    if known:
        unknown = [n for n in names if n not in known]
        if unknown:
            raise HTTPException(status_code=422, detail=f"unknown sheet(s): {', '.join(unknown)}")
        if len(names) >= len(known):
            raise HTTPException(status_code=422, detail="at least one sheet must be scanned")
    return names


def _apply_result(res: dict[str, Any], output_name: str) -> dict[str, Any]:
    return {
        "status": res["status"],
        "method": res.get("method", "excel_com"),
        "output_path": str(res.get("output_path", "")),
        "output_name": output_name,
        "applied": [_public_op(o) for o in res.get("applied", [])],
        "failed": [_public_op(o) for o in res.get("failed", [])],
        "verified_opens_in_excel": res.get("verified_opens_in_excel"),
        "warnings": res.get("warnings", []),
        "message": res.get("message", ""),
    }


def _before_apply(s: Session) -> dict[str, Any]:
    """What the analysis said before an Apply -- the other half of its outcome."""
    return {"plan": s.plan, "report": s.result["validation_report"] if s.result else None, "version_id": s.current_version_id}


def _outcome(s: Session, ops: list[dict[str, Any]], res: dict[str, Any], before: dict[str, Any], reanalyzed: bool) -> dict[str, Any]:
    """The Apply, repair by repair (app.readiness.apply_outcome), against the
    analysis that preceded it and -- when one ran -- the one that followed."""
    after_report = s.result["validation_report"] if reanalyzed and s.result else None
    outcome = apply_outcome(ops, res, before["plan"], before["report"], s.plan if reanalyzed else None, after_report, s.current_version_id, before["version_id"])
    s.apply_history.append(apply_record(outcome))
    return outcome


def _validate_ops(raw_ops: Any) -> list[dict[str, Any]]:
    if not isinstance(raw_ops, list) or not raw_ops:
        raise HTTPException(status_code=422, detail="operations must be a non-empty list")
    ops = []
    for i, o in enumerate(raw_ops, start=1):
        if not isinstance(o, dict) or o.get("op") not in ASSISTANT_OPS | STRUCTURAL_OPS:
            raise HTTPException(status_code=422, detail=f"operation {i}: unknown op {o.get('op') if isinstance(o, dict) else o!r}")
        if not isinstance(o.get("sheet"), str):
            raise HTTPException(status_code=422, detail=f"operation {i}: sheet is required")
        ops.append({"action_id": "frontend", "rule_id": "UI", **o})
    return ops


# --- API ----------------------------------------------------------------------------------
@app.get("/api/health")
def health() -> dict[str, Any]:
    from app.sizing import size_threshold_bytes

    return {
        "version": app.version,
        "excel": com_available(),
        "assistant": llm_available(),
        "rules": len(engine().active_rules()),
        "frontend": DIST_DIR.is_dir(),
        # 1.6.6: deferred upload + sheet selection + scan status are available
        "upload_gate": True,
        "size_threshold_mb": round(size_threshold_bytes(load_config()) / MB, 1),
        # 1.7.2: the assistant's model and POST /chat/stream
        "model": DEFAULT_MODEL_ID,
        "streaming": True,
        # 1.8.0: POST /auto-fix ("Fix all automatically" on the Recalculate step)
        "autofix": True,
    }


@app.post("/api/sessions")
def create_session(file: UploadFile = File(...), mode: str = Form("plan"), defer: str = Form(""), ignore_sheets: str = Form("")) -> dict[str, Any]:
    """Upload a workbook. The file is stored and *inspected* (sizes, sheet
    list -- from the package alone, nothing is parsed) before anything else.

    `defer` truthy: stop there and return `{pending: true, size}` so the
    front-end can ask about skipping sheets when the workbook is above the
    size threshold, then POST .../analyze. Otherwise (default; scripts and
    older front-ends) the file is converted/analysed at once, honouring an
    optional `ignore_sheets` JSON list."""
    if mode not in MODE_MAP:
        raise HTTPException(status_code=422, detail=f"unknown mode {mode}")
    name = Path(file.filename or "workbook.xlsx").name
    if Path(name).suffix.lower() not in ALLOWED_SUFFIXES:
        raise HTTPException(status_code=422, detail="upload an .xlsx, .xlsm or .xlsb file")
    sid = uuid.uuid4().hex[:12]
    work_dir = Path(tempfile.mkdtemp(prefix="mind_ready_"))
    upload_dir = work_dir / "upload"
    upload_dir.mkdir(parents=True)
    raw_path = upload_dir / name
    with raw_path.open("wb") as out:
        shutil.copyfileobj(file.file, out)
    s = Session(id=sid, work_dir=work_dir, mode=mode)
    SESSIONS[sid] = s

    from app.inventory import sha256_of

    version = _add_version(s, raw_path, "upload", "original upload", [], None, sha256_of(raw_path))
    s.size = size_gate(raw_path, load_config())
    s.ignore_sheets = _parse_ignore(s, ignore_sheets)
    if str(defer).strip().lower() in ("1", "true", "yes", "on"):
        s.pending = True
        return {"sessionId": sid, "pending": True, "mode": mode, "size": s.size, "version": version, "versions": s.versions, "scan": _scan_snapshot(s)}
    return {"sessionId": sid, **_run_scan(s, raw_path)}


@app.post("/api/sessions/{session_id}/analyze")
def analyze(session_id: str, payload: dict[str, Any] = Body(default={})) -> dict[str, Any]:
    """Scan a deferred upload (or scan the current version again) with the
    given `ignore_sheets`. Long call: poll GET .../status meanwhile."""
    s = _session(session_id)
    if "ignore_sheets" in payload:
        s.ignore_sheets = _parse_ignore(s, payload.get("ignore_sheets"))
    return {"sessionId": s.id, **_run_scan(s)}


@app.get("/api/sessions/{session_id}/status")
def scan_status(session_id: str) -> dict[str, Any]:
    """Where the current (or last) scan is: state pending|running|ready|error,
    stage + title + message, in-stage fraction, overall progress, the sheet
    or rule being worked on, elapsed seconds, per-stage timings."""
    return _scan_snapshot(_session(session_id))


@app.post("/api/sessions/{session_id}/reanalyze")
def reanalyze(session_id: str, payload: dict[str, Any] = Body(default={})) -> dict[str, Any]:
    s = _session(session_id)
    vid = payload.get("versionId") or s.current_version_id
    if vid not in s.paths:
        raise HTTPException(status_code=404, detail=f"unknown version {vid}")
    s.current_version_id = vid
    return _run_scan(s, s.paths[vid])


@app.post("/api/sessions/{session_id}/apply")
def apply(session_id: str, payload: dict[str, Any] = Body(...)) -> dict[str, Any]:
    s = _session(session_id)
    ops = _validate_ops(payload.get("operations"))
    before = _before_apply(s)
    res = _tracked_apply(s, ops, bool(payload.get("reanalyze", True)))
    output = Path(res["output_path"])
    action_ids = {o.get("action_id") for o in ops}
    source = "assistant" if action_ids == {"assistant"} else "formula" if action_ids == {"formula_replacement"} else "prep"
    who = {"assistant": "Assistant", "formula": "Formula fix", "prep": "Prep"}[source]
    verified = res.get("verified_opens_in_excel")
    label = f"{who}: {len(res.get('applied', []))} change(s) via {'Excel' if res.get('method') == 'excel_com' else 'openpyxl'} · {'verified' if verified else 'unverified' if verified is None else 'FAILED TO OPEN'}"
    minor = payload.get("versionStep") == "minor"  # 1.7.4: a fix on the Recalculate step
    version = _add_version(s, output, source, label, res.get("applied", []), verified, res["change_log_entry"]["output_sha256"], minor=minor)
    out: dict[str, Any] = {"result": _apply_result(res, version["file_name"]), "version": version}
    reanalyzed = bool(payload.get("reanalyze", True)) and res["status"] in ("APPLIED", "PARTIAL")
    if reanalyzed:
        out.update(_run_scan(s, output))
    out["outcome"] = _outcome(s, ops, res, before, reanalyzed)
    out["prep_progress"] = prep_progress(s.plan_history, s.apply_history)  # the gauge, this Apply included
    return out


@app.post("/api/sessions/{session_id}/suggest")
def suggest(session_id: str, payload: dict[str, Any] = Body(...)) -> dict[str, Any]:
    s = _session(session_id)
    if not s.result:
        raise HTTPException(status_code=409, detail="no analysis yet")
    rule_id, sheet, cell = payload.get("ruleId"), payload.get("sheet"), payload.get("cell")
    finding = next((f for f in s.result["validation_report"]["findings"] if f["rule_id"] == rule_id), None)
    if finding is None:
        raise HTTPException(status_code=404, detail=f"no finding for rule {rule_id}")
    formula = next((x["formula"] for x in s.result["workbook_analysis"]["workbooks"][0]["formulas"] if x["sheet"] == sheet and x["cell"] == cell), "")
    res = suggest_formula_fix(finding, formula)
    return {"suggestion": res.get("suggestion"), "formula": extract_formula(res.get("suggestion")), "message": res.get("message")}


@app.post("/api/sessions/{session_id}/grid-names")
def grid_names(session_id: str, payload: dict[str, Any] = Body(default={})) -> dict[str, Any]:
    """Ask the assistant to name the grids whose deterministic name carries no
    meaning ('Cashflows C4', 'I'), and fold the answers into the prep plan.

    Nothing is written: the names only change what `create_grid_titles` would
    *propose*, which the user still reviews and approves. `apply: false`
    returns the comparison without changing the session's plan."""
    s = _session(session_id)
    if not s.result:
        raise HTTPException(status_code=409, detail="no analysis yet")
    analysis = s.result["workbook_analysis"]
    labels = standalone_labels(analysis)
    targets = []
    for g in all_grids(analysis):
        if g.get("name") and not is_weak_name(g.get("name"), g["sheet"]):
            continue
        deterministic, source = deterministic_name(analysis, g, labels)
        if g.get("name") or is_weak_name(deterministic, g["sheet"]):
            targets.append((g, deterministic, source))
    if not targets:
        return {"available": True, "names": [], "message": "every grid already has a meaningful name", "applied": False}
    contexts = [build_grid_context(analysis, g) for g, _, _ in targets]
    res = suggest_names(contexts)
    if not res["available"]:
        return {"available": False, "names": [], "message": res.get("message"), "applied": False}
    proposed = res["names"]
    rows = [
        {
            "grid": f'{g["sheet"]}!{g["ref"]}',
            "sheet": g["sheet"],
            "ref": g["ref"],
            "size": f'{g["n_rows"]}x{g["n_cols"]}',
            "current": g.get("name"),
            "deterministic": deterministic,
            "deterministic_source": source,
            "suggested": proposed.get(f'{g["sheet"]}!{g["ref"]}'),
        }
        for g, deterministic, source in targets
    ]
    applied = bool(payload.get("apply", True))
    if applied:
        s.grid_names.update({k: v for k, v in proposed.items() if v})
        s.plan = plan_actions(analysis, s.result["validation_report"], s.grid_names or None)
    return {"available": True, "names": rows, "message": res.get("message"), "applied": applied, "plan": s.plan if applied else None, "readiness": s.readiness() if applied else None}


# --- the app runs itself against real Mind (1.6.8) ------------------------------------
MIND_LOOPS: dict[str, dict[str, Any]] = {}  # session id -> {state, events, report, started, work_dir}
_MIND_LOOP_LOCK = threading.Lock()


def _loop_running() -> str | None:
    for sid, st in MIND_LOOPS.items():
        if st.get("state") == "running":
            return sid
    return None


@app.post("/api/sessions/{session_id}/mind-loop")
def start_mind_loop(session_id: str, payload: dict[str, Any] = Body(default={})) -> dict[str, Any]:
    """Start the autonomous run-in-Mind loop on the session's current version, in
    a background thread. One loop at a time (it owns the Excel session and the
    Mind browser profile). Poll GET for progress."""
    s = _session(session_id)
    if s.current_path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise HTTPException(status_code=422, detail="Mind accepts .xlsx / .xlsm only")
    with _MIND_LOOP_LOCK:
        busy = _loop_running()
        if busy:
            raise HTTPException(status_code=409, detail=f"a Mind loop is already running (session {busy})")
        work = s.work_dir / "mind_loop" / time.strftime("%Y%m%d_%H%M%S")
        state: dict[str, Any] = {"state": "running", "events": [], "report": None, "started": time.time(), "work_dir": str(work), "error": None}
        MIND_LOOPS[session_id] = state
    cfg = LoopConfig(
        source=s.current_path,
        work_dir=work,
        max_iterations=int(payload.get("maxIterations", 4)),
        use_assistant=bool(payload.get("useAssistant", True)),
        run_model=bool(payload.get("runModel", True)),
        check_numbers=bool(payload.get("checkNumbers", True)),
        enable=list(payload.get("enable", [])),
        disable=list(payload.get("disable", [])),
        delete_projects=bool(payload.get("deleteProjects", True)),
        skip_mind=bool(payload.get("skipMind", False)),
    )

    def progress(e: dict[str, Any]) -> None:
        state["events"].append({"t": time.time(), **e})

    def worker() -> None:
        try:
            state["report"] = run_loop(cfg, progress)
            state["state"] = "done"
        except Exception as exc:  # the thread must never die silently
            state["error"] = str(exc)[:300]
            state["state"] = "error"
            progress({"event": "done", "verdict": "error", "message": f"loop crashed: {exc}"[:300]})

    threading.Thread(target=worker, name=f"mind-loop-{session_id}", daemon=True).start()
    return {"state": "running", "work_dir": str(work)}


@app.get("/api/sessions/{session_id}/mind-loop")
def mind_loop_status(session_id: str, after: int = 0) -> dict[str, Any]:
    """Progress of the session's loop: events after index `after`, the report so
    far (written after every iteration) and the final verdict."""
    _session(session_id)
    st = MIND_LOOPS.get(session_id)
    if not st:
        return {"state": "idle", "events": [], "report": None, "next": 0}
    report = st.get("report")
    if report is None:
        candidate = Path(st["work_dir"]) / "loop_report.json"
        if candidate.is_file():
            try:
                report = json.loads(candidate.read_text(encoding="utf-8"))
            except Exception:
                report = None
    events = st["events"][after:]
    return {"state": st["state"], "events": events, "next": after + len(events), "report": report, "error": st.get("error"), "started": st["started"], "work_dir": st["work_dir"], "summary": summarize(report) if report else None}


def _chat_turn(s: Session, payload: dict[str, Any]) -> tuple[list[dict[str, str]], str, str | None]:
    """(history, question, extra_context) for one chat turn; raises the
    409/422 HTTP errors before any model call."""
    if not s.result:
        raise HTTPException(status_code=409, detail="no analysis yet")
    question = str(payload.get("question") or "").strip()
    if not question:
        raise HTTPException(status_code=422, detail="question is required")
    history = [{"role": m.get("role", "user"), "content": str(m.get("content", ""))} for m in payload.get("history", []) if isinstance(m, dict) and m.get("content")]
    focus = payload.get("focus")
    extra = recalculation_context(s.recalc) if s.recalc else None
    if isinstance(focus, dict) and (focus.get("rule_id") or focus.get("error")):
        # A focused mini chat: naming the rule id and the cell makes
        # app.chat_context.retrieve() pull the finding's observed data and the
        # cell contents into the turn; a recalculation error carries the error
        # value and the formula Excel evaluated.
        where = f" at {focus['sheet']}!{focus['cell']}" if focus.get("sheet") and focus.get("cell") else ""
        cells = [c for c in (focus.get("cells") or []) if isinstance(c, dict) and c.get("sheet") and c.get("cell")]
        if focus.get("kind") == "recalc-group" or len(cells) > 1:
            listed = "; ".join(
                f"{c['sheet']}!{c['cell']} ({c.get('error', '')}) {str(c.get('formula', ''))[:120]}" + (f" [array {c['array']}]" if c.get("array") else "")
                for c in cells[:40]
            )
            arrays = sorted({f"{c['sheet']}!{c['array']}" for c in cells if c.get("array")})
            listed_keys = {f"{c['sheet']}!{c['cell']}" for c in cells}
            bits = []
            for a in arrays:
                sh, ref = a.rsplit("!", 1)
                members = _array_cells(ref)
                quiet = [c for c in members if f"{sh}!{c}" not in listed_keys]
                formula_here = next((str(c.get("formula") or "") for c in cells if c.get("array") == ref and c.get("sheet") == sh), "")
                facts = array_size_hint(ref, formula_here)
                bits.append(
                    f"{a} is ONE array formula (Ctrl+Shift+Enter) over {ref}; all {len(members)} cells share the same formula"
                    + (f", and {', '.join(quiet[:12])} belong to it although they are not in error" if quiet else "")
                    + (f". Facts: {facts}" if facts else "")
                    + "."
                )
            array_note = (
                " IMPORTANT: " + " ".join(bits) + " Excel cannot change part of an array, so the fix must cover ALL cells of the array: "
                "either one set_array_formula on the full range, or one operation for EVERY cell of it (including the cells that are not in error). "
                "Re-entering the same formula changes nothing: the new content must remove the cause -- e.g. an array entered over more cells than its "
                "source range has values must be shrunk (set_array_formula over the cells that have values, clear_cell for the rest) or replaced by "
                "per-cell formulas."
                if arrays else ""
            )
            question = (
                f"[About {len(cells)} recalculation errors sharing one root cause: {focus.get('cause') or focus.get('error', '')}. Cells: {listed}.{array_note} "
                "If you propose a fix, propose operations that fix EVERY listed cell (one operation per cell, or one range operation when the "
                "same formula applies to a contiguous range) and say explicitly that all of them can be fixed together.] "
                f"{question}"
            )
        elif focus.get("kind") == "recalc" or focus.get("error"):
            formula = f"; formula: {str(focus.get('formula'))[:300]}" if focus.get("formula") else ""
            siblings = ""
            for g in (s.recalc or {}).get("groups", []):
                members = [c for c in g["cells"] if not (c["sheet"] == focus.get("sheet") and c["cell"] == focus.get("cell"))]
                if len(members) < len(g["cells"]) and members:
                    siblings = (
                        f" This cell shares its root cause ({g['cause']}) with {len(members)} other cell(s): "
                        + ", ".join(f"{c['sheet']}!{c['cell']}" for c in members[:20])
                        + "; tell the user they can all be fixed together and, if asked to fix, include every one of them."
                    )
                    break
            own_array = focus.get("array")
            for g in (s.recalc or {}).get("groups", []) if not own_array else []:
                own_array = next((c.get("array") for c in g["cells"] if c["sheet"] == focus.get("sheet") and c["cell"] == focus.get("cell") and c.get("array")), None)
                if own_array:
                    break
            others = [c for c in _array_cells(own_array) if c != str(focus.get("cell", "")).upper()] if own_array else []
            own_facts = array_size_hint(own_array, str(focus.get("formula") or "")) if own_array else ""
            array_note = (
                f" This cell is part of the array formula {focus.get('sheet')}!{own_array} (Ctrl+Shift+Enter) together with {', '.join(others[:12])}"
                f"{' ...' if len(others) > 12 else ''} -- all of them share this formula. Excel cannot change part of an array, so a fix must cover every "
                f"cell of {own_array} (one set_array_formula on the range, or one operation per cell, including cells that are not in error); "
                "re-entering the same formula changes nothing -- the new content must remove the cause (shrink an array that is larger than its "
                "source values, or replace it by per-cell formulas)." + (f" Facts: {own_facts}." if own_facts else "")
                if own_array else ""
            )
            question = f"[About the recalculation error {focus.get('error', '')}{where}{formula}.{siblings}{array_note}] {question}"
        else:
            # 1.7.2: a replacement is judged by a full Excel recalculation afterwards -- say so up front
            question = (
                f"[About finding {focus['rule_id']}{where}. If you propose a replacement formula, it must behave like the original on "
                "blank cells, text and errors (e.g. wrap with IF(x=\"\",0,...) when the original tolerated blanks); every new error "
                "after recalculation counts against the fix.] " + question
            )
    return history, question, extra


def _chat_payload(reply: dict[str, Any]) -> dict[str, Any]:
    proposal = None
    if reply.get("proposal") or reply.get("proposal_errors"):
        proposal = {"summary": reply.get("proposal_summary"), "operations": [_public_op(o) for o in reply.get("proposal", [])], "errors": reply.get("proposal_errors", [])}
    return {
        "text": reply.get("text") or f"(no answer: {reply.get('message')})",
        "proposal": proposal,
        "provenance": {"context_chars": reply.get("context_chars", 0), "detail_chars": reply.get("detail_chars", 0), "lookups": reply.get("lookups", 0)},
    }


@app.post("/api/sessions/{session_id}/chat")
def chat(session_id: str, payload: dict[str, Any] = Body(...)) -> dict[str, Any]:
    s = _session(session_id)
    history, question, extra = _chat_turn(s, payload)
    reply = answer_question(history, question, s.result["workbook_analysis"], s.result["validation_report"], s.plan, extra_context=extra)
    return _chat_payload(reply)


@app.post("/api/sessions/{session_id}/chat/stream")
def chat_stream(session_id: str, payload: dict[str, Any] = Body(...)) -> StreamingResponse:
    """The same turn as /chat, streamed as NDJSON (one JSON object per line):
    {"type": "phase", "phase": "answer"|"lookup"|"repair", "n"} before every
    model call, {"type": "delta", "text"} for each chunk of the reply, then
    {"type": "done", <the /chat payload>} or {"type": "error", "detail"}."""
    s = _session(session_id)
    history, question, extra = _chat_turn(s, payload)
    events: queue.Queue[dict[str, Any] | None] = queue.Queue()

    def work() -> None:
        try:
            reply = answer_question(history, question, s.result["workbook_analysis"], s.result["validation_report"], s.plan, extra_context=extra, on_event=events.put)
            events.put({"type": "done", **_chat_payload(reply)})
        except Exception as exc:  # the stream has started: report in-band
            events.put({"type": "error", "detail": f"{type(exc).__name__}: {exc}"})
        finally:
            events.put(None)

    threading.Thread(target=work, name=f"chat-{session_id}", daemon=True).start()

    def lines():
        while True:
            ev = events.get()
            if ev is None:
                return
            yield json.dumps(ev, ensure_ascii=False, default=str) + "\n"

    return StreamingResponse(lines(), media_type="application/x-ndjson", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


def _original_baseline(s: Session, original: Path) -> dict[str, Any] | None:
    """The original upload after the same full Excel recalculation: its error
    cells (and its formulas that already contain #REF!). Computed once per
    session on a fresh copy; None when Excel could not recalculate it, in which
    case the caller falls back to the original's cached values."""
    if s.original_baseline is None:
        copy_path, _ = make_immutable_copy(original, s.work_dir / "recalc_original")
        res = recalculate(copy_path)
        ran = bool(res.get("ran", res.get("status") != "NOT_SUPPORTED"))
        s.original_baseline = {
            "ran": ran,
            "errors": list(res.get("formula_errors") or []) + list(res.get("addin_gap_errors") or []) if ran else [],
            "ref_formulas": formulas_with_broken_refs(original),
            "message": res.get("message"),
        }
    return s.original_baseline if s.original_baseline.get("ran") else None


@app.post("/api/sessions/{session_id}/recalculate")
def recalc(session_id: str) -> dict[str, Any]:
    s = _session(session_id)
    copy_path, _ = make_immutable_copy(s.current_path, s.work_dir / "recalc")
    res = recalculate(copy_path)
    return _recalc_out(s, res, copy_path)


def _recalc_out(s: Session, res: dict[str, Any], copy_path: Path, close_round: bool = True) -> dict[str, Any]:
    """What a recalculation of the current version means for the session: the
    errors against the original's, the root-cause groups, READY-001, the verdict.
    `res` has the shape of app.recalc.recalculate -- the automatic fixer (1.8.0)
    ends with a full recalculation of its own and hands it in the same way."""
    out = {
        "status": res["status"],
        "message": res["message"],
        "version_id": s.current_version_id,
        "ran": bool(res.get("ran", res["status"] != "NOT_SUPPORTED" and not str(res.get("message", "")).startswith("Excel COM recalculation failed"))),
        "formula_errors": [
            {"sheet": e.get("sheet", ""), "cell": e.get("cell", ""), "error": e.get("error", ""), "formula": e.get("formula", ""), **({"array": e["array"]} if e.get("array") else {})}
            for e in res.get("formula_errors", [])
        ],
        "addin_gap_errors": [{"sheet": e.get("sheet", ""), "cell": e.get("cell", ""), "formula": e.get("formula", "")} for e in res.get("addin_gap_errors", [])],
    }
    out["groups"] = group_errors(out["formula_errors"], out["addin_gap_errors"])
    # 1.7.2: errors the original already had are the model's own, not the preparation's.
    # The baseline is the ORIGINAL after the same full recalculation (once per
    # session), never its cached values -- those can be stale.
    original = next((s.paths[v["id"]] for v in s.versions if s.paths[v["id"]].suffix.lower() in (".xlsx", ".xlsm")), None)
    if original is not None and out["formula_errors"]:
        base = _original_baseline(s, original)
        if base is not None:
            cls = classify_errors(out["formula_errors"], base["errors"], base["ref_formulas"])
        else:
            cls = classify_against_original(original, out["formula_errors"])
            cls["message"] = (cls.get("message") or "") + " -- compared with the original's cached values, its own recalculation did not run"
        out["formula_errors"] = cls["errors"]
        out["preexisting_errors"], out["new_errors"] = cls["preexisting"], cls["new"]
        out["compared_with_original"] = cls["compared"]
        if cls.get("message"):
            out["message"] = f"{out['message']} ({cls['message'].strip(' -')})"
    else:
        out["preexisting_errors"], out["new_errors"] = 0, len(out["formula_errors"])
        out["compared_with_original"] = original is not None
    # 1.7.4: a recalculation after fixes (v2.1, v2.2, ...) closes the round: the
    # same file becomes the next major version (v3), and the recalculation is its own
    current = next((v for v in s.versions if v["id"] == s.current_version_id), None)
    if close_round and out["ran"] and current is not None and current.get("minor"):
        short = current["label"].split(" — ")[0]
        version = _add_version(s, s.current_path, "recalculate", f"recalculated {short}", [], current.get("verified_opens_in_excel"), current["sha256"])
        out["version"] = version
        out["version_id"] = version["id"]
    s.recalc = out
    s.recalc_version_id = s.current_version_id
    _sync_recalc_finding(s)
    if s.result:
        out["report"] = {k: v for k, v in s.result["validation_report"].items() if k != "_version_id"}
        out["plan"] = s.plan
    s.recalc_path = copy_path
    out["readiness"] = s.readiness()  # 1.7.2: a clean recalculation is what turns the verdict green
    return out


# --- 1.8.0: fix all automatically ----------------------------------------------------------------
AUTOFIX: dict[str, dict[str, Any]] = {}  # session id -> live state of the current / last run
_AUTOFIX_PUBLIC = ("state", "stage", "title", "message", "fraction", "pass_no", "errors", "errors_before", "kept", "started", "updated", "finished", "error", "version_id", "use_assistant", "keep_good_values")
_AUTOFIX_TITLES = {
    "start": "Starting", "read": "Reading the workbook", "find": "Finding where the errors start", "advise": "Asking the assistant",
    "fix": "Fixing", "pass": "Fixing", "confirm": "Confirming", "save": "Saving", "verify": "Checking the file opens", "reanalyze": "Re-analyzing the fixed version", "done": "Done",
}


@app.post("/api/sessions/{session_id}/auto-fix")
def start_autofix(session_id: str, payload: dict[str, Any] = Body(default={})) -> dict[str, Any]:
    """Fix every formula error of the current version that can be fixed without
    changing a value that is good today (app/autofix.py), in a background
    thread: one Excel session that finds the cells where errors start, fixes
    them, recalculates and rolls back whatever moved a good value, pass after
    pass. Poll GET for progress; the result is a new version, already
    recalculated and re-analysed.

    Body (all optional): use_assistant (default true: the assistant says which
    value an error cell should fall back to, and why), keep_good_values (default
    true -- the owner's rule; false also fixes the errors a formula is hiding,
    and lists every value that moved), time_budget_min (default 120)."""
    s = _session(session_id)
    if not s.result:
        raise HTTPException(status_code=409, detail="no analysis yet")
    if s.current_path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise HTTPException(status_code=422, detail="the automatic fixer works on .xlsx / .xlsm")
    if not com_available():
        raise HTTPException(status_code=409, detail="the automatic fixer needs Excel on this machine")
    running = AUTOFIX.get(session_id)
    if running and running.get("state") == "running":
        raise HTTPException(status_code=409, detail="the automatic fixer is already running for this session")
    use_assistant = bool(payload.get("use_assistant", True))
    keep = bool(payload.get("keep_good_values", True))
    budget = max(1.0, float(payload.get("time_budget_min", 120))) * 60
    advisor = make_advisor() if use_assistant else None
    state: dict[str, Any] = {
        "state": "running", "stage": "start", "title": "Starting", "message": "Opening the workbook in Excel", "fraction": None,
        "pass_no": 0, "errors": None, "errors_before": None, "kept": 0, "started": time.time(), "updated": time.time(), "finished": None,
        "error": None, "stop": False, "version_id": s.current_version_id, "use_assistant": advisor is not None, "keep_good_values": keep,
        "result": None, "version": None, "recalc": None, "analysis": None,
    }
    AUTOFIX[session_id] = state
    source = s.current_path

    def progress(stage: str, message: str, fraction: float | None = None, **facts: Any) -> None:
        state.update(stage=stage, title=_AUTOFIX_TITLES.get(stage, stage), message=message, fraction=fraction, updated=time.time())
        if facts.get("pass_no") is not None:
            state["pass_no"] = facts["pass_no"]
        if facts.get("errors") is not None:
            state["errors"] = facts["errors"]
            if state["errors_before"] is None:
                state["errors_before"] = facts["errors"]
        if stage == "pass" and facts.get("kept") is not None:
            state["kept"] += int(facts["kept"])

    def worker() -> None:
        try:
            cfg = AutoFixConfig(time_budget_s=budget, keep_good_values=keep, advisor=advisor, progress=progress, should_stop=lambda: bool(state["stop"]))
            res = run_autofix(source, s.work_dir / "autofix", cfg)
            if res.get("status") == "error":
                raise RuntimeError(res.get("message") or "the automatic fixer failed")
            if res.get("output_path"):
                output = Path(res["output_path"])
                left = int(res.get("errors_after") or 0)
                verified = res.get("verified_opens_in_excel")
                label = (
                    f"Auto-fix: {res.get('cells_rewritten', 0)} cell(s) fixed · "
                    + ("recalculated clean" if left == 0 else f"{left} error(s) left for a person")
                    + ("" if verified else " · NOT VERIFIED")
                )
                state["version"] = _add_version(s, output, "autofix", label, res.get("applied", []), verified, res["change_log_entry"]["output_sha256"])
                rec = _recalc_out(s, res["recalc"], output, close_round=False)
                progress("reanalyze", "Checking the fixed version against the rules", None)
                state["analysis"] = _run_scan(s, output)
                rec["readiness"] = s.readiness()
                rec.pop("report", None)
                rec.pop("plan", None)
                state["recalc"] = rec
            state["result"] = {k: v for k, v in res.items() if k not in ("applied", "change_log_entry", "recalc", "log")}
            state["result"]["log"] = res.get("log", [])[-200:]
            state.update(state="done", stage="done", title="Done", message=res.get("summary") or "", finished=time.time(), updated=time.time())
        except HTTPException as exc:  # the re-analysis refused (a scan already running)
            state.update(state="error", error=str(exc.detail)[:400], finished=time.time(), updated=time.time())
        except Exception as exc:  # the thread must never die silently
            state.update(state="error", error=f"{type(exc).__name__}: {exc}"[:400], finished=time.time(), updated=time.time())

    threading.Thread(target=worker, name=f"autofix-{session_id}", daemon=True).start()
    return _autofix_public(s, state, full=False)


def _autofix_public(s: Session, state: dict[str, Any], full: bool) -> dict[str, Any]:
    out = {k: state.get(k) for k in _AUTOFIX_PUBLIC}
    out["elapsed_s"] = round((state.get("finished") or time.time()) - state["started"], 1)
    if state.get("stage") == "reanalyze" and state.get("state") == "running":
        scan = _scan_snapshot(s)
        out["message"] = f"{scan.get('title') or 'Re-analyzing'}: {scan.get('message') or ''}".strip(": ")
        out["fraction"] = scan.get("overall")
    out["version"] = state.get("version")
    result = state.get("result")
    if result is not None:
        out["result"] = result if full else {k: v for k, v in result.items() if k not in ("fixes", "left", "good_values_changed_list", "log")}
    if full:
        out["recalc"] = state.get("recalc")
        out["analysis"] = state.get("analysis")
    return out


@app.get("/api/sessions/{session_id}/auto-fix")
def autofix_status(session_id: str, full: int = 0) -> dict[str, Any]:
    """Where the automatic fixer is (state running | done | error | idle, the pass,
    the errors left). `full=1` once it is done adds the whole result (every fix
    with its reason, what is left and why), the recalculation and the new analysis."""
    s = _session(session_id)
    state = AUTOFIX.get(session_id)
    if not state:
        return {"state": "idle"}
    return _autofix_public(s, state, full=bool(full))


@app.post("/api/sessions/{session_id}/auto-fix/stop")
def stop_autofix(session_id: str) -> dict[str, Any]:
    """Ask a running fixer to stop after the round it is in. What it has fixed so far is kept."""
    s = _session(session_id)
    state = AUTOFIX.get(session_id)
    if not state or state.get("state") != "running":
        return {"state": (state or {}).get("state", "idle")}
    state["stop"] = True
    return _autofix_public(s, state, full=False)


@app.get("/api/sessions/{session_id}/readiness")
def readiness(session_id: str) -> dict[str, Any]:
    """1.7.2: the one verdict -- blocked / unverified / ready -- with what stands
    in the way and the next step, plus the Prep gauge history."""
    s = _session(session_id)
    r = s.readiness()
    if r is None:
        raise HTTPException(status_code=409, detail="no analysis yet")
    return {"readiness": r, "prep_progress": prep_progress(s.plan_history, s.apply_history)}


@app.get("/api/sessions/{session_id}/cells")
def cells(session_id: str, sheet: str, cell: str, rows: int = 3, cols: int = 3) -> dict[str, Any]:
    """What is in the workbook around a cell: contents (values) and formulas,
    for the per-error workbook view. Values come from the last recalculation
    of the current version when there is one, else from the analysis copy."""
    s = _session(session_id)
    if not s.result:
        raise HTTPException(status_code=409, detail="no analysis yet")
    fresh = s.recalc_path if (s.recalc_path and s.recalc_version_id == s.current_version_id and Path(s.recalc_path).is_file()) else None
    return cell_window(s.result["workbook_analysis"], sheet, cell, rows=rows, cols=cols, values_path=fresh)


# --- Grid Namer (1.7.0) ---------------------------------------------------------------
@app.get("/api/mind-flags")
def mind_flags() -> dict[str, Any]:
    """The documented Mind title flags (references/mind-flags.yaml), for the
    Grid Namer's flag picker. Header-cell markers and alias spellings are
    left out -- they are not written into a grid title."""
    flags = [
        {"name": spec["name"], "meaning": spec.get("meaning", "")}
        for spec in documented_flags().values()
        if spec.get("where", "title") == "title" and not spec.get("alias_of")
    ]
    return {"flags": sorted(flags, key=lambda f: f["name"].lower())}


@app.get("/api/sessions/{session_id}/sheet-cells")
def sheet_cells(session_id: str, sheet: str, max_rows: int = 400, max_cols: int = 60) -> dict[str, Any]:
    """The whole (used) area of one sheet of the current version, for the Grid
    Namer's workbook view: per cell the calculated value when the file carries
    one, else the formula text. Values come from the last recalculation of the
    current version when there is one, else from the analysis copy."""
    s = _session(session_id)
    if not s.result:
        raise HTTPException(status_code=409, detail="no analysis yet")
    analysis = s.result["workbook_analysis"]
    if sheet not in {sh["name"] for sh in analysis["workbooks"][0]["sheets"]}:
        raise HTTPException(status_code=404, detail=f"no sheet {sheet!r}")
    fresh = s.recalc_path if (s.recalc_path and s.recalc_version_id == s.current_version_id and Path(s.recalc_path).is_file()) else None
    path = Path(fresh or analysis["source"]["copy_path"])
    max_rows, max_cols = max(1, min(int(max_rows), 2000)), max(1, min(int(max_cols), 200))

    import openpyxl

    raw_wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    val_wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws_raw, ws_val = raw_wb[sheet], val_wb[sheet]
        if not hasattr(ws_raw, "iter_rows"):
            raise HTTPException(status_code=422, detail=f"{sheet} is a chart sheet -- there are no cells to show")
        total_rows, total_cols = int(ws_raw.max_row or 0), int(ws_raw.max_column or 0)
        rows: list[list[Any]] = []
        for raw_row, val_row in zip(
            ws_raw.iter_rows(min_row=1, max_row=min(total_rows, max_rows) or 1, max_col=min(total_cols, max_cols) or 1, values_only=True),
            ws_val.iter_rows(min_row=1, max_row=min(total_rows, max_rows) or 1, max_col=min(total_cols, max_cols) or 1, values_only=True),
        ):
            rows.append([json_safe(v if v is not None else raw) for raw, v in zip(raw_row, val_row)])
        while rows and all(v is None or v == "" for v in rows[-1]):
            rows.pop()
        width = 0
        for row in rows:
            for i in range(len(row) - 1, width - 1, -1):
                if row[i] is not None and row[i] != "":
                    width = i + 1
                    break
        rows = [row[:width] for row in rows]
    finally:
        raw_wb.close()
        val_wb.close()
    return {
        "sheet": sheet,
        "rows": rows,
        "n_rows": len(rows),
        "n_cols": width if rows else 0,
        "total_rows": total_rows,
        "total_cols": total_cols,
        "truncated": total_rows > max_rows or total_cols > max_cols,
        "values_from": "recalculation" if fresh else "analysis copy",
    }


@app.post("/api/sessions/{session_id}/grid-namer")
def grid_namer(session_id: str, payload: dict[str, Any] = Body(...)) -> dict[str, Any]:
    """The Grid Namer's Submit: each user-selected area becomes a '#Name /Flags'
    grid title (written with the reference-safety rules -- a refused area comes
    back in `skipped` with the reason); with nameRest (default) every other
    untitled grid is titled by the existing conventions in the same apply. The
    result is a new verified version of the workbook."""
    s = _session(session_id)
    if not s.result:
        raise HTTPException(status_code=409, detail="no analysis yet")
    raw_areas = payload.get("areas")
    if not isinstance(raw_areas, list) or not raw_areas:
        raise HTTPException(status_code=422, detail="areas must be a non-empty list")
    areas = []
    for i, a in enumerate(raw_areas, start=1):
        if not isinstance(a, dict) or not isinstance(a.get("sheet"), str) or not isinstance(a.get("ref"), str) or not str(a.get("name") or "").strip():
            raise HTTPException(status_code=422, detail=f"area {i}: sheet, ref and name are required")
        flags = a.get("flags") or []
        if not isinstance(flags, list):
            raise HTTPException(status_code=422, detail=f"area {i}: flags must be a list of flag names")
        areas.append({"sheet": a["sheet"], "ref": a["ref"], "name": str(a["name"]), "flags": [str(f) for f in flags]})

    analysis, report = s.result["workbook_analysis"], s.result["validation_report"]
    manual = plan_named_areas(analysis, areas)
    ops = list(manual["ops"])
    skipped = list(manual["skipped"])
    if not ops:
        raise HTTPException(status_code=422, detail="no area could be named: " + ("; ".join(skipped[:8]) or "nothing resolved"))
    auto_count = 0
    if bool(payload.get("nameRest", True)):
        auto_ops, auto_skipped = plan_create_grid_titles(
            analysis, report, s.grid_names or None,
            exclude=manual["exclude"], reserved=manual["reserved"],
            pre_titled=manual["titled_cells"], pre_inserted=manual["inserted_rows"],
        )
        ops += auto_ops
        skipped += auto_skipped
        auto_count = len(auto_ops)

    before = _before_apply(s)
    res = _tracked_apply(s, ops, bool(payload.get("reanalyze", True)))
    output = Path(res["output_path"])
    verified = res.get("verified_opens_in_excel")
    label = f"Grid Namer: {len(manual['named'])} named area(s) + {auto_count} automatic title op(s) · {'verified' if verified else 'unverified' if verified is None else 'FAILED TO OPEN'}"
    version = _add_version(s, output, "grid_namer", label, res.get("applied", []), verified, res["change_log_entry"]["output_sha256"])
    out: dict[str, Any] = {
        "result": _apply_result(res, version["file_name"]),
        "version": version,
        "named": manual["named"],
        "skipped": skipped,
        "manual_ops": len(manual["ops"]),
        "auto_ops": auto_count,
    }
    reanalyzed = bool(payload.get("reanalyze", True)) and res["status"] in ("APPLIED", "PARTIAL")
    if reanalyzed:
        out.update(_run_scan(s, output))
    out["outcome"] = _outcome(s, ops, res, before, reanalyzed)
    out["prep_progress"] = prep_progress(s.plan_history, s.apply_history)  # the gauge, this Apply included
    return out


@app.post("/api/sessions/{session_id}/reports")
def reports(session_id: str) -> dict[str, Any]:
    s = _session(session_id)
    if not s.result:
        raise HTTPException(status_code=409, detail="no analysis yet")
    analysis = s.result["workbook_analysis"]
    report = s.result["validation_report"]
    copy_path = Path(analysis["source"]["copy_path"])
    stem = s.current_path.stem
    out_dir = s.work_dir / "reports" / s.current_version_id
    standalone = build_standalone_report(report, out_dir / f"{stem}_mind_readiness_report.xlsx", rules_by_id=engine().rules, source_name=s.current_path.name)
    built = build_report_workbook(copy_path, report, out_dir / f"{stem}_with_report{s.current_path.suffix}", rules_by_id=engine().rules, source_name=s.current_path.name)
    return {
        "standalone_name": _register_file(s, standalone),
        "workbook_name": _register_file(s, Path(built.path)),
        "method": built.method,
        "verified_opens_in_excel": built.verified_opens_in_excel,
        "warnings": built.warnings,
    }


@app.get("/api/sessions/{session_id}/versions")
def versions(session_id: str) -> list[dict[str, Any]]:
    return _session(session_id).versions


@app.get("/api/sessions/{session_id}/files/{file_name}")
def download(session_id: str, file_name: str):
    s = _session(session_id)
    path = s.files.get(file_name)
    if path is None or not Path(path).is_file():
        raise HTTPException(status_code=404, detail=f"no file {file_name} in this session")
    media = "application/vnd.ms-excel.sheet.macroEnabled.12" if Path(path).suffix.lower() == ".xlsm" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return FileResponse(str(path), media_type=media, filename=file_name)


# --- built front-end ------------------------------------------------------------------------
if DIST_DIR.is_dir():
    from fastapi.staticfiles import StaticFiles

    if (DIST_DIR / "assets").is_dir():
        app.mount("/assets", StaticFiles(directory=str(DIST_DIR / "assets")), name="assets")

    @app.get("/{full_path:path}", include_in_schema=False)
    def spa(full_path: str):
        if full_path.startswith("api/"):
            raise HTTPException(status_code=404)
        candidate = DIST_DIR / full_path
        if full_path and candidate.is_file():
            return FileResponse(str(candidate))
        return FileResponse(str(DIST_DIR / "index.html"))

else:

    @app.get("/", include_in_schema=False)
    def no_frontend():
        return JSONResponse({"detail": f"front-end build not found at {DIST_DIR}; run `npm run build` in FigmaOutput or set MIND_READY_DIST"}, status_code=404)


def _answers_as_mind_ready(port: int, timeout: float = 20) -> bool:
    """A Mind Ready server already answers on `port` (slowly, when it is in
    the middle of analysing a large workbook -- hence the long timeout)."""
    import urllib.request

    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/health", timeout=timeout) as res:
            body = json.loads(res.read().decode("utf-8"))
    except Exception:
        return False
    return isinstance(body, dict) and "version" in body and "rules" in body


def _port_is_taken(port: int) -> bool:
    import socket

    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.settimeout(1)
        return sock.connect_ex(("127.0.0.1", port)) == 0


def main() -> None:
    import uvicorn

    port = int(os.environ.get("MIND_READY_PORT", "8600"))
    url = f"http://localhost:{port}"
    # Started a second time (F5 while the .bat window is open, or the reverse): say so and
    # open the app, instead of failing on "only one usage of each socket address".
    if _port_is_taken(port):
        print(f"Port {port} is already in use. Checking whether it is Mind Ready (up to 20 s) ...")
        if _answers_as_mind_ready(port):
            import webbrowser

            print(f"Mind Ready is already running: {url}")
            print("Opening it in the browser. To restart it (after a code change), stop the running one first:")
            print("close its window, or press the red square / Shift+F5 in VS Code, then start again.")
            webbrowser.open(url)
            return
        print(f"Whatever uses port {port} did not answer as Mind Ready.")
        print(f"- If Mind Ready is already running and busy with a large workbook, it is at {url}: wait, then reload the page.")
        print("- Otherwise close the program that uses the port, or choose another one: set MIND_READY_PORT=8601 and start again.")
        raise SystemExit(1)
    uvicorn.run("app.web.server:app", host="127.0.0.1", port=port, reload=False)


if __name__ == "__main__":
    main()
