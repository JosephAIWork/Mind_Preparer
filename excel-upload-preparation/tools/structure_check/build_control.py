"""A control workbook whose right answers are known, to test the structure detector."""
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L

wb = openpyxl.Workbook()
inp = wb.active
inp.title = "Inputs"
inp["A1"], inp["B1"], inp["C1"] = "Scenario", "S1", "S2"
inp["D1"] = "S3"
for i, (a, b, c) in enumerate([(0.02, 0.03, 0.01)] * 1, start=2):
    inp[f"A{i}"], inp[f"B{i}"], inp[f"C{i}"], inp[f"D{i}"] = "Rate", a, b, c

# 1. Three product sheets: same layout, own assumptions; Prod_B has one formula override.
assump = {"A": (100, 10, 0.05), "B": (200, 6, 0.04), "C": (50, 16, 0.07)}
for p, (pol, prem, lapse) in assump.items():
    ws = wb.create_sheet(f"Prod_{p}")
    ws["A1"], ws["B1"] = "Policies", pol
    ws["A2"], ws["B2"] = "Premium", prem
    ws["A3"], ws["B3"] = "Lapse", lapse
    ws["A5"] = "Year"
    ws["A6"], ws["A7"], ws["A8"], ws["A9"] = "In force", "Premium", "Expense", "Profit"
    for y in range(1, 21):
        c = L(y + 1)
        prev = L(y)
        ws[f"{c}5"] = y
        ws[f"{c}6"] = "=$B$1" if y == 1 else f"={prev}6*(1-$B$3)"
        ws[f"{c}7"] = f"={c}6*$B$2"
        ws[f"{c}8"] = f"={c}7*0.1" if p != "B" else f"={c}7*0.12+1"
        ws[f"{c}9"] = f"={c}7-{c}8"
tot = wb.create_sheet("Total")
tot["A1"] = "Profit"
for y in range(1, 21):
    c = L(y + 1)
    tot[f"{c}1"] = f"=Prod_A!{c}9+Prod_B!{c}9+Prod_C!{c}9"

# 2. Three independent scenario blocks side by side (4 columns each), 10 years down.
sc = wb.create_sheet("Scen")
for k in range(3):
    c0 = 1 + k * 5            # block width 4 + one empty column
    A, B, C, D = (L(c0 + i) for i in range(4))
    sc[f"{A}1"] = f"Scenario {k + 1}"
    sc[f"{A}2"], sc[f"{B}2"], sc[f"{C}2"], sc[f"{D}2"] = "Year", "Rate", "Fund", "Interest"
    for t in range(1, 11):
        r = t + 2
        sc[f"{A}{r}"] = t
        sc[f"{B}{r}"] = f"=Inputs!{L(2 + k)}$2"
        sc[f"{C}{r}"] = "=1000" if t == 1 else f"={C}{r - 1}+{D}{r - 1}"
        sc[f"{D}{r}"] = f"={C}{r}*{B}{r}"
sc["P2"] = "Total fund"
for t in range(1, 11):
    sc[f"P{t + 2}"] = f"=C{t + 2}+H{t + 2}+M{t + 2}"

# 3. Chained year sheets: each opens with the previous year's closing (a sequence, not copies).
for i, yr in enumerate((2024, 2025, 2026)):
    ws = wb.create_sheet(f"Y{yr}")
    ws["A1"], ws["A2"], ws["A3"], ws["A4"] = "Opening", "Premium", "Claims", "Closing"
    ws["B1"] = 500 if i == 0 else f"=Y{yr - 1}!B4"
    ws["B2"], ws["B3"] = 120 + 10 * i, 80 + 5 * i
    ws["B4"] = "=B1+B2-B3"
    for m in range(1, 13):
        c = L(m + 2)
        ws[f"{c}1"] = "=B1" if m == 1 else f"={L(m + 1)}4"
        ws[f"{c}2"] = "=$B$2/12"
        ws[f"{c}3"] = "=$B$3/12"
        ws[f"{c}4"] = f"={c}1+{c}2-{c}3"

# 4. A plain one-policy time projection (filled across, recurrent): nothing to flag.
pj = wb.create_sheet("Proj")
pj["A1"], pj["A2"], pj["A3"] = "Month", "Reserve", "Release"
for m in range(1, 61):
    c = L(m + 1)
    pj[f"{c}1"] = m
    pj[f"{c}2"] = "=1000" if m == 1 else f"={L(m)}2-{L(m)}3"
    pj[f"{c}3"] = f"={c}2*0.02"

out = Path(__file__).parent / "control.xlsx"
wb.save(out)
print(out)
