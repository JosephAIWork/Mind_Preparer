"""2026-10-26: FRM-009 stayed blocking on a model after Do everything -- all 512,796 cells it
flagged were skipped as "an array formula: rewrite it by hand". They were one-cell dynamic arrays
(<f t="array" ref="J222" cm="1">): Excel 365 stores most formulas with INDEX that way. A one-cell
array is now rewritten like any formula (Formula2 keeps it one: .Formula would add an @), and the
512,792 cells of the output table -- 8 columns, one formula filled down 64,099 rows each -- are
written as 8 fills instead of one write per cell."""
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app.config import load_config
from app.excel_com import close_quietly, com_available, excel_session, open_for_write, save_in_place
from app.inventory import build_analysis
from app.numbers_check import Moves, compare_workbooks
from app.prep import _apply_with_openpyxl, _fill_down_blocks, _op, _write_formulas, apply_operations, plan_actions
from app.rules_engine import RulesEngine
from app.validators import run_rule
from app.validators.mind_values import SEVERAL_CELLS, several_cells

needs_excel = pytest.mark.skipif(not com_available(), reason="needs a real Excel")


def test_only_an_array_over_several_cells_is_left_by_hand():
    assert not several_cells({"cell": "J222", "array_ref": "J222"})
    assert not several_cells({"cell": "J222", "array_ref": "$J$222:$J$222"})
    assert not several_cells({"cell": "J222"})
    assert several_cells({"cell": "G2", "array_ref": "G2:G3"})


def _ops(rows, formula, after, sheet="Out", col="B"):
    return [_op("set_formula", "empty_cells_as_zero", "FRM-009", sheet, cell=f"{col}{r}", before=formula.format(r=r), after=after.format(r=r), note=f"compares $A{r} (empty)") for r in rows]


def test_a_formula_filled_down_is_one_operation_per_unbroken_block():
    before, after = "=IF($A{r}>0,+INDEX($H$2:$I$5,$A{r},1)*$C{r},0)", "=IF(N($A{r})>0,+INDEX($H$2:$I$5,$A{r},1)*$C{r},0)"
    ops = (
        _ops([2, 4, 5, 6, 7], before, after)                                                   # row 3 is not planned: B2 alone, then B4:B7
        + _ops([8], "=IF($A{r}>0,+INDEX($H$2:$I$5,$A{r},2)*$C{r},0)", "=IF(N($A{r})>0,+INDEX($H$2:$I$5,$A{r},2)*$C{r},0)")  # another formula
        + _ops([9], "=IF($A{r}>0,GETPIVOTDATA(\"x\",$P$1),0)", "=IF(N($A{r})>0,GETPIVOTDATA(\"x\",$P$1),0)")               # stays alone
        + _ops([2, 3], "=IF($A{r}>0,1,2)", "=IF(N($A{r})>0,1,2)", col="J")
        + _ops([2, 3], "=IF($A{r}>0,SUM(1:1),2)", "=IF(N($A{r})>0,SUM(1:1),2)", col="K")            # a fill would move 1:1: stays one per cell
        + _ops([2, 3], "=IF($A{r}>0,[1]Data!A1,2)", "=IF(N($A{r})>0,[1]Data!A1,2)", col="L")        # another workbook's cell: the same
    )
    analysis = {"workbooks": [{"formulas": [{"sheet": "Out", "cell": f"B{r}", "array_ref": f"B{r}"} for r in range(2, 10)]}]}
    out = {o["cell"]: o for o in _fill_down_blocks(ops, analysis)}
    assert set(out) == {"B2", "B4:B7", "B8", "B9", "J2:J3", "K2", "K3", "L2", "L3"}
    block = out["B4:B7"]
    assert block["block_cells"] == 4 and block["one_cell_array"] is True
    assert (block["before"], block["after"]) == (before.format(r=4), after.format(r=4))
    assert (block["before_last"], block["after_last"]) == (before.format(r=7), after.format(r=7))
    assert block["note"] == "compares $A4 (empty) -- the same formula down to B7 (4 cells), written as one fill"
    assert out["B2"]["one_cell_array"] and "block_cells" not in out["B2"]
    assert "one_cell_array" not in out["J2:J3"] and out["J2:J3"]["block_cells"] == 2   # a plain formula joins plain ones only


def test_a_block_that_does_not_come_out_as_planned_is_put_back():
    class Cell:
        Formula2 = "=IF(N($A9)>0,+@INDEX($H$2:$I$5,$A9,1)*$C9,0)"

    class Count:
        def __init__(self, n):
            self.Count = n

    class Block:
        Rows, Columns = Count(8), Count(1)
        written: list = []

        def __setattr__(self, name, value):
            self.written.append((name, value))

        def Cells(self, r, c):
            assert (r, c) == (8, 1)
            return Cell()

    o = {"op": "set_formula", "one_cell_array": True, "before": "=IF($A2>0,1,0)", "after": "=IF(N($A2)>0,1,0)", "after_last": "=IF(N($A9)>0,+INDEX($H$2:$I$5,$A9,1)*$C9,0)"}
    rng = Block()
    with pytest.raises(ValueError, match="put back as it was"):
        _write_formulas(rng, o, "Out", "B2:B9")
    assert rng.written == [("Formula2", o["after"]), ("Formula2", o["before"])]


def test_without_excel_each_cell_of_a_block_gets_its_own_row(tmp_path):
    path = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Out"
    for r in range(2, 6):
        wb.active[f"B{r}"] = f"=IF($A{r}>0,1,2)"
    wb.save(path)
    block = _fill_down_blocks(_ops(range(2, 6), "=IF($A{r}>0,1,2)", "=IF(N($A{r})>0,1,2)"), {"workbooks": [{"formulas": []}]})
    applied, failed = _apply_with_openpyxl(path, block)
    assert [o["cell"] for o in applied] == ["B2:B5"] and not failed
    ws = openpyxl.load_workbook(path).active
    assert [ws[f"B{r}"].value for r in range(2, 6)] == [f"=IF(N($A{r})>0,1,2)" for r in range(2, 6)]


def _workbook(path: Path) -> Path:
    """Out!B and E: one-cell dynamic arrays (Formula2 with INDEX), F2 a Ctrl+Shift+Enter cell, G2:G3
    an array over two cells, J plain formulas; every one compares column A, empty but for A3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Out"
    ws["A1"], ws["B1"], ws["C1"] = "Index", "Adjusted", "Base"
    ws["A3"] = 2
    for r in range(2, 10):
        ws[f"C{r}"] = r
    for i, r in enumerate(range(2, 6)):
        ws[f"H{r}"], ws[f"I{r}"] = 1 + i / 10, 2 + i / 10
    wb.save(path)
    wb.close()
    with excel_session() as excel:
        book = open_for_write(excel, path)
        try:
            sh = book.Worksheets("Out")
            sh.Range("B2:B9").Formula2 = "=IF($A2>0,+INDEX($H$2:$I$5,$A2,1)*$C2,0)"
            sh.Range("E2:E9").Formula2 = "=IF($A2>0,+INDEX($H$2:$I$5,$A2,2)*$C2,0)"
            sh.Range("F2").FormulaArray = "=IF($A2>0,SUM($H$2:$H$5*$C2),0)"
            sh.Range("G2:G3").FormulaArray = "=IF($A$2>0,{1;2},0)"
            sh.Range("J2:J9").Formula = "=IF($A2>0,1,2)"
            excel.CalculateFullRebuild()
            save_in_place(book, path)
        finally:
            close_quietly(book)
    return path


def _sheet_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("xl/worksheets/sheet1.xml").decode("utf-8")


@needs_excel
def test_one_cell_arrays_are_rewritten_as_one_cell_arrays_and_blocks_as_one_fill(tmp_path):
    src = _workbook(tmp_path / "arrays.xlsx")
    assert re.search(r'<c r="B9"[^>]*cm="1"[^>]*><f t="array" ref="B9"', _sheet_xml(src))   # what Excel 365 does with INDEX: the case itself
    analysis = build_analysis(src, tmp_path / "w", "a")
    finding = run_rule(RulesEngine().get("FRM-009"), analysis, load_config())
    sites = {s["cell"]: s for s in finding["observed"]["sites"]}
    assert set(sites) == {"B2", "B4", "B5", "B6", "B7", "B8", "B9", "E2", "E4", "E5", "E6", "E7", "E8", "E9", "F2", "G2", "J2", "J4", "J5", "J6", "J7", "J8", "J9"}
    assert sites["B9"]["suggested_formula"] == "=IF(N($A9)>0,+INDEX($H$2:$I$5,$A9,1)*$C9,0)"
    assert sites["G2"]["suggested_formula"] is None and sites["G2"]["issue"] == SEVERAL_CELLS

    plan = {a["id"]: a for a in plan_actions(analysis, {"findings": [finding]})}["empty_cells_as_zero"]
    ops = {o["cell"]: o for o in plan["operations"]}
    assert set(ops) == {"B2", "B4:B9", "E2", "E4:E9", "F2", "J2", "J4:J9"}
    assert ops["B4:B9"]["one_cell_array"] and ops["B4:B9"]["block_cells"] == 6 and ops["F2"]["one_cell_array"] and "one_cell_array" not in ops["J4:J9"]
    assert plan["skipped"] == [f"Out!G2: {SEVERAL_CELLS}"]

    out = apply_operations(src, tmp_path / "prep", plan["operations"])
    assert out["status"] == "APPLIED" and not out["failed"] and len(out["applied"]) == 7
    done = Path(out["output_path"])
    with excel_session() as excel:
        book = open_for_write(excel, done)
        try:
            sh = book.Worksheets("Out")
            assert sh.Range("B9").Formula2 == "=IF(N($A9)>0,+INDEX($H$2:$I$5,$A9,1)*$C9,0)"   # no @: still the formula it was
            assert sh.Range("E5").Formula2 == "=IF(N($A5)>0,+INDEX($H$2:$I$5,$A5,2)*$C5,0)"
            assert sh.Range("B3").Formula2 == "=IF($A3>0,+INDEX($H$2:$I$5,$A3,1)*$C3,0)"     # not planned: as it was
            assert sh.Range("F2").HasArray and sh.Range("F2").Formula == "=IF(N($A2)>0,SUM($H$2:$H$5*$C2),0)"
            assert sh.Range("J9").Formula == "=IF(N($A9)>0,1,2)"
            assert str(sh.Range("G2").CurrentArray.Address) == "$G$2:$G$3"
        finally:
            book.Saved = True
            close_quietly(book)
    xml = _sheet_xml(done)
    assert re.search(r'<c r="B9"[^>]*cm="1"[^>]*><f t="array" ref="B9"', xml) and re.search(r'<c r="E4"[^>]*cm="1"[^>]*><f t="array" ref="E4"', xml)

    res = compare_workbooks(src, done, Moves())
    assert res["ran"] and res["clean"] is True
    after = run_rule(RulesEngine().get("FRM-009"), build_analysis(done, tmp_path / "after", "b"), load_config())
    assert [s["cell"] for s in after["observed"]["sites"]] == ["G2"]   # the array over two cells only
