"""1.6.0: the FastAPI backend behind the Figma-built front-end (app/web/server.py),
exercised through FastAPI's TestClient against the synthetic fixtures. The
assistant call is mocked; the apply/report paths use Excel when it is
available (they fall back to openpyxl otherwise, which the test tolerates)."""
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import chat_context
from app.excel_com import com_available
from app.web import server


@pytest.fixture(scope="module")
def client():
    return TestClient(server.app)


def _upload(client, path: Path, mode: str = "plan") -> dict:
    with path.open("rb") as f:
        res = client.post("/api/sessions", files={"file": (path.name, f, "application/octet-stream")}, data={"mode": mode})
    assert res.status_code == 200, res.text
    return res.json()


def test_health_reports_engine_facts(client):
    body = client.get("/api/health").json()
    assert body["rules"] == 101 and isinstance(body["excel"], bool) and isinstance(body["assistant"], bool)


def test_upload_returns_the_front_end_contracts(client, flagged_model_broken_xlsx):
    body = _upload(client, flagged_model_broken_xlsx)
    assert set(body) >= {"sessionId", "summary", "report", "plan", "version"}
    summary, report, plan, version = body["summary"], body["report"], body["plan"], body["version"]
    assert summary["file_name"] == "flagged_model_broken.xlsx" and summary["file_type"] == "xlsx"
    assert {s["name"] for s in summary["sheets"]} == {"Model", "Settings", "Outputs"}
    grid = next(g for s in summary["sheets"] for g in s["grids"] if g["anchor"] == "A4")
    assert grid["title"] == "#Assumptions /Reorder /Inpt" and sorted(grid["flag_names"]) == ["inpt", "reorder"]
    assert set(grid) >= {"display_name", "ref", "header_values", "inner_title_cells", "formula_count"}
    assert report["status"] == "NOT_SUPPORTED" and report["summary"]["finding_count"] == 101
    assert all({"rule_id", "status", "confidence", "evidence", "location", "message", "readiness_impact", "correction_available", "source"} <= set(f) for f in report["findings"])
    assert {a["id"] for a in plan} >= {"flag_spelling", "loop_name_case", "create_grid_titles"}
    assert version["id"] == "ver-001" and version["source"] == "upload" and version["label"].startswith("v1")


def test_mode_filters_rules_and_bad_inputs_are_rejected(client, plain_grid_xlsx):
    body = _upload(client, plain_grid_xlsx, mode="fix_formulas")
    categories = {f["rule_id"].split("-")[0] for f in body["report"]["findings"]}
    assert categories <= {"FILE", "READY", "FRM", "FORMULA"}
    with plain_grid_xlsx.open("rb") as f:
        assert client.post("/api/sessions", files={"file": ("x.csv", f, "text/csv")}, data={"mode": "plan"}).status_code == 422
    with plain_grid_xlsx.open("rb") as f:
        assert client.post("/api/sessions", files={"file": ("x.xlsx", f, "application/octet-stream")}, data={"mode": "nope"}).status_code == 422
    assert client.get("/api/sessions/does-not-exist/versions").status_code == 404


def test_apply_creates_a_verified_version_and_reanalyzes(client, flagged_model_broken_xlsx):
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    plan = {a["id"]: a for a in body["plan"]}
    ops = plan["flag_spelling"]["operations"] + plan["special_headers"]["operations"]
    res = client.post(f"/api/sessions/{sid}/apply", json={"operations": ops, "reanalyze": True})
    assert res.status_code == 200, res.text
    out = res.json()
    assert out["result"]["status"] == "APPLIED"
    assert out["result"]["output_name"].endswith(".xlsx") and out["version"]["id"] == "ver-002" and out["version"]["source"] == "prep"
    assert len(out["version"]["change_log"]) == len(ops)
    # the re-analysis of the new file is included and reflects the corrections
    statuses = {f["rule_id"]: f["status"] for f in out["report"]["findings"]}
    assert statuses["FLG-001"] == "PASS" and statuses["PAR-002"] == "PASS"
    # download by the name the front-end was given
    dl = client.get(f"/api/sessions/{sid}/files/{out['result']['output_name']}")
    assert dl.status_code == 200 and dl.content[:2] == b"PK"
    versions = client.get(f"/api/sessions/{sid}/versions").json()
    assert [v["id"] for v in versions] == ["ver-001", "ver-002"]
    # applying again works on the produced file and yields v3 with a distinct download name
    res2 = client.post(f"/api/sessions/{sid}/apply", json={"operations": plan["loop_name_case"]["operations"], "reanalyze": False}).json()
    assert res2["version"]["id"] == "ver-003" and res2["result"]["output_name"] != out["result"]["output_name"]
    assert client.post(f"/api/sessions/{sid}/apply", json={"operations": [{"op": "explode", "sheet": "Model"}]}).status_code == 422


def test_status_carries_the_apply_progress(client, flagged_model_broken_xlsx):
    """1.7.2: GET /status says what the Apply did stage by stage, so the
    front-end shows the real step and the changes written, not a timer."""
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    assert client.get(f"/api/sessions/{sid}/status").json()["apply"] is None  # no Apply yet
    ops = {a["id"]: a for a in body["plan"]}["flag_spelling"]["operations"]
    out = client.post(f"/api/sessions/{sid}/apply", json={"operations": ops, "reanalyze": True})
    assert out.status_code == 200, out.text
    status = client.get(f"/api/sessions/{sid}/status").json()
    ap = status["apply"]
    assert ap["state"] == "done" and ap["stage"] == "done" and ap["error"] is None and ap["reanalyze"] is True
    assert ap["total"] == len(ops) and ap["done"] == len(ops) and ap["applied"] == len(ops) and ap["failed"] == 0
    assert [s["stage"] for s in ap["stages"]] == ["apply_copy", "apply_write", "apply_save", "apply_verify", "apply_log"]
    assert ap["elapsed_s"] >= 0 and ap["finished_at"] >= ap["started_at"]
    # the re-analysis that followed is the scan of the same status, started after the Apply
    assert status["state"] == "ready" and status["version_id"] == "ver-002" and status["started_at"] >= ap["started_at"]


def test_chat_returns_reply_and_validated_proposal(client, flagged_model_broken_xlsx, monkeypatch):
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]

    def fake_chat(messages, system_prompt, **kwargs):
        assert any("Model" in m["content"] for m in messages) or "Model" in system_prompt
        return {"available": True, "text": 'Renaming.\n```changes\n{"summary": "rename", "operations": [{"op": "rename_sheet", "sheet": "Settings", "new_name": "Config"}, {"op": "set_value", "sheet": "Nope", "cell": "A1", "value": 1}]}\n```', "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    res = client.post(f"/api/sessions/{sid}/chat", json={"history": [{"role": "user", "content": "hi"}, {"role": "assistant", "content": "hello", "note": True}], "question": "rename Settings to Config"})
    assert res.status_code == 200, res.text
    reply = res.json()
    assert reply["text"] == "Renaming." and reply["provenance"]["context_chars"] > 1000
    assert [o["op"] for o in reply["proposal"]["operations"]] == ["rename_sheet"] and reply["proposal"]["summary"] == "rename"
    assert len(reply["proposal"]["errors"]) == 1
    assert client.post(f"/api/sessions/{sid}/chat", json={"history": [], "question": "   "}).status_code == 422


def test_chat_stream_sends_phases_deltas_then_the_chat_payload(client, flagged_model_broken_xlsx, monkeypatch):
    """1.7.2: /chat/stream is the same turn as /chat, as NDJSON lines."""
    import json

    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    answer = 'Renaming.\n```changes\n{"summary": "rename", "operations": [{"op": "rename_sheet", "sheet": "Settings", "new_name": "Config"}]}\n```'

    def fake_chat(messages, system_prompt, **kwargs):
        if kwargs.get("on_delta"):
            kwargs["on_delta"](answer[:5])
            kwargs["on_delta"](answer[5:])
        return {"available": True, "text": answer, "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    turn = {"history": [], "question": "rename Settings to Config", "focus": {"rule_id": "RES-002", "sheet": "Model", "cell": "J6"}}
    res = client.post(f"/api/sessions/{sid}/chat/stream", json=turn)
    assert res.status_code == 200 and res.headers["content-type"].startswith("application/x-ndjson")
    events = [json.loads(line) for line in res.text.splitlines() if line.strip()]
    assert [e["type"] for e in events] == ["phase", "delta", "delta", "done"]
    assert events[0] == {"type": "phase", "phase": "answer", "n": 0}
    assert events[1]["text"] + events[2]["text"] == answer
    done = {k: v for k, v in events[-1].items() if k != "type"}
    assert done == client.post(f"/api/sessions/{sid}/chat", json=turn).json()
    assert done["text"] == "Renaming." and [o["op"] for o in done["proposal"]["operations"]] == ["rename_sheet"]
    # validation errors stay plain HTTP errors (nothing is streamed)
    assert client.post(f"/api/sessions/{sid}/chat/stream", json={"history": [], "question": "  "}).status_code == 422
    assert client.post("/api/sessions/nope/chat/stream", json={"history": [], "question": "hi"}).status_code == 404


def test_chat_stream_reports_a_failed_turn_in_band(client, flagged_model_broken_xlsx, monkeypatch):
    import json

    sid = _upload(client, flagged_model_broken_xlsx)["sessionId"]

    def boom(messages, system_prompt, **kwargs):
        raise ValueError("gateway exploded")

    monkeypatch.setattr(chat_context, "chat_completion", boom)
    res = client.post(f"/api/sessions/{sid}/chat/stream", json={"history": [], "question": "hi"})
    events = [json.loads(line) for line in res.text.splitlines() if line.strip()]
    assert events[-1]["type"] == "error" and "gateway exploded" in events[-1]["detail"]
    assert client.get("/api/health").json()["streaming"] is True


def test_reports_and_recalculate_endpoints(client, plain_grid_xlsx):
    body = _upload(client, plain_grid_xlsx)
    sid = body["sessionId"]
    rep = client.post(f"/api/sessions/{sid}/reports").json()
    assert rep["standalone_name"].endswith("_mind_readiness_report.xlsx") and rep["workbook_name"].endswith(".xlsx")
    assert client.get(f"/api/sessions/{sid}/files/{rep['standalone_name']}").content[:2] == b"PK"
    rc = client.post(f"/api/sessions/{sid}/recalculate").json()
    assert rc["status"] in ("PASS", "ERROR", "NOT_SUPPORTED", "WARNING")
    assert isinstance(rc["formula_errors"], list) and isinstance(rc["addin_gap_errors"], list)
# --- 1.6.1: re-analysis delta + finding-focused chat -------------------------------------


def test_apply_and_reanalyze_report_a_delta_and_fixability(client, flagged_model_broken_xlsx):
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    findings = {f["rule_id"]: f for f in body["report"]["findings"]}
    # "Fix available" now reflects the prep plan, not the YAML flag
    assert findings["FLG-001"]["correction_available"] is True and findings["PAR-002"]["correction_available"] is True
    assert findings["UNQ-001"]["correction_available"] is False
    plan = {a["id"]: a for a in body["plan"]}
    out = client.post(f"/api/sessions/{sid}/apply", json={"operations": plan["special_headers"]["operations"], "reanalyze": True}).json()
    delta = out["delta"]
    fixed = {d["rule_id"] for d in delta["fixed"]}
    assert {"PAR-002", "PRJ-002"} <= fixed and delta["previous_version_id"] == "ver-001" and delta["version_id"] == "ver-002"
    assert delta["previous_counts"]["ERROR"] > delta["counts"]["ERROR"]
    # re-analyzing the same version again changes nothing
    again = client.post(f"/api/sessions/{sid}/reanalyze", json={"versionId": "ver-002"}).json()
    assert again["delta"]["fixed"] == [] and again["delta"]["regressed"] == [] and [v["id"] for v in again["versions"]] == ["ver-001", "ver-002"]
    # going back to v1 shows the regression honestly
    back = client.post(f"/api/sessions/{sid}/reanalyze", json={"versionId": "ver-001"}).json()
    assert {d["rule_id"] for d in back["delta"]["regressed"]} >= {"PAR-002", "PRJ-002"}


def test_chat_focus_pulls_the_finding_into_the_turn(client, flagged_model_broken_xlsx, monkeypatch):
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    seen = {}

    def fake_chat(messages, system_prompt, **kwargs):
        seen["turn"] = messages[-1]["content"]
        return {"available": True, "text": "explained", "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    res = client.post(f"/api/sessions/{sid}/chat", json={"history": [], "question": "Explain this finding", "focus": {"rule_id": "RES-002", "sheet": "Model", "cell": "J6"}})
    assert res.status_code == 200 and res.json()["text"] == "explained"
    assert seen["turn"].startswith("[About finding RES-002 at Model!J6.")
    assert "behave like the original on blank cells" in seen["turn"]  # 1.7.2: a replacement is judged by the recalculation
    assert "## Finding RES-002" in seen["turn"] and "J6 = =MM_RESULT(J5,\"scenario\",1)" in seen["turn"]
# --- 1.6.2: recalculation errors in the Fix panel -------------------------------------------


def test_chat_focus_on_a_recalculation_error_gets_the_recalc_context(client, flagged_model_broken_xlsx, monkeypatch):
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    # plant a recalculation result as the recalc endpoint would (no Excel needed here)
    server.SESSIONS[sid].recalc = {
        "status": "ERROR", "message": "Recalculated; 1 genuine formula error cell(s) found.", "version_id": "ver-001",
        "formula_errors": [{"sheet": "Model", "cell": "L5", "error": "#NAME?", "formula": "=SUM(L4,MM_SETSIZE(3,1))"}],
        "addin_gap_errors": [{"sheet": "Model", "cell": "J5", "formula": '=MM_LOOP("Scenario", A5:A7)'}],
    }
    seen = {}

    def fake_chat(messages, system_prompt, **kwargs):
        seen["turn"] = messages[-1]["content"]
        seen["system"] = system_prompt
        return {"available": True, "text": "explained", "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    res = client.post(f"/api/sessions/{sid}/chat", json={"history": [], "question": "Explain this error", "focus": {"kind": "recalc", "rule_id": "READY-001", "sheet": "Model", "cell": "L5", "error": "#NAME?", "formula": "=SUM(L4,MM_SETSIZE(3,1))"}})
    assert res.status_code == 200 and res.json()["text"] == "explained"
    assert seen["turn"].startswith("[About the recalculation error #NAME? at Model!L5; formula: =SUM(L4,MM_SETSIZE(3,1)).")
    assert "<recalculation>" in seen["system"] and "Model!L5 #NAME?" in seen["system"] and "add-in is not installed" in seen["system"]
    assert "L5 = =SUM(L4,MM_SETSIZE(3,1))" in seen["turn"]  # the cell itself was retrieved


def test_recalculate_endpoint_records_the_version(client, plain_grid_xlsx):
    body = _upload(client, plain_grid_xlsx)
    sid = body["sessionId"]
    rc = client.post(f"/api/sessions/{sid}/recalculate").json()
    assert rc["version_id"] == "ver-001"
    assert server.SESSIONS[sid].recalc is rc or server.SESSIONS[sid].recalc == rc


def test_chat_focus_on_a_group_asks_for_a_fix_for_every_cell(client, flagged_model_broken_xlsx, monkeypatch):
    from app.recalc import group_errors

    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    errors = [{"sheet": "Model", "cell": "L6", "error": "#VALUE!", "formula": "=L5+1"}, {"sheet": "Model", "cell": "L7", "error": "#VALUE!", "formula": "=L6+1"}]
    server.SESSIONS[sid].recalc = {"status": "ERROR", "message": "2 errors", "version_id": "ver-001", "ran": True, "formula_errors": errors, "addin_gap_errors": [], "groups": group_errors(errors, [])}
    seen = {}

    def fake_chat(messages, system_prompt, **kwargs):
        seen["turn"] = messages[-1]["content"]
        seen["system"] = system_prompt
        return {"available": True, "text": "one fix", "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    group = server.SESSIONS[sid].recalc["groups"][0]
    res = client.post(f"/api/sessions/{sid}/chat", json={"history": [], "question": "Propose one fix for all cells", "focus": {"kind": "recalc-group", "rule_id": "READY-001", "error": "#VALUE!", "cause": group["cause"], "cells": group["cells"]}})
    assert res.status_code == 200
    assert seen["turn"].startswith("[About 2 recalculation errors sharing one root cause")
    assert "Model!L6" in seen["turn"] and "Model!L7" in seen["turn"] and "fix EVERY listed cell" in seen["turn"]
    assert "root causes (errors that can be fixed together)" in seen["system"]
    # a single-cell focus mentions its siblings
    res = client.post(f"/api/sessions/{sid}/chat", json={"history": [], "question": "Explain", "focus": {"kind": "recalc", "rule_id": "READY-001", "sheet": "Model", "cell": "L6", "error": "#VALUE!", "formula": "=L5+1"}})
    assert res.status_code == 200 and "shares its root cause" in seen["turn"] and "Model!L7" in seen["turn"]


def test_chat_focus_on_an_array_group_tells_the_model_to_fix_the_whole_array(client, flagged_model_broken_xlsx, monkeypatch):
    from app.recalc import group_errors

    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    errors = [{"sheet": "Model", "cell": f"L{r}", "error": "#N/A", "formula": "=A5:A6*2", "array": "L6:L8"} for r in (6, 7, 8)]
    server.SESSIONS[sid].recalc = {"status": "ERROR", "message": "3 errors", "version_id": "ver-001", "ran": True, "formula_errors": errors, "addin_gap_errors": [], "groups": group_errors(errors, [])}
    seen = {}

    def fake_chat(messages, system_prompt, **kwargs):
        seen["turn"] = messages[-1]["content"]
        return {"available": True, "text": "ok", "message": None}

    monkeypatch.setattr(chat_context, "chat_completion", fake_chat)
    group = server.SESSIONS[sid].recalc["groups"][0]
    assert group["arrays"] == ["Model!L6:L8"]
    res = client.post(f"/api/sessions/{sid}/chat", json={"history": [], "question": "fix", "focus": {"kind": "recalc-group", "rule_id": "READY-001", "error": "#N/A", "cause": group["cause"], "cells": group["cells"]}})
    assert res.status_code == 200 and "Model!L6:L8 is ONE array formula" in seen["turn"] and "set_array_formula" in seen["turn"] and "[array L6:L8]" in seen["turn"]
    assert "Facts: the array L6:L8 has 3x1 cells" in seen["turn"] and "set_array_formula on L6:L7 with the same formula and clear_cell L8" in seen["turn"]
    res = client.post(f"/api/sessions/{sid}/chat", json={"history": [], "question": "fix", "focus": {"kind": "recalc", "rule_id": "READY-001", "sheet": "Model", "cell": "L7", "error": "#N/A", "formula": "=A5:A6*2"}})
    assert res.status_code == 200 and "part of the array formula Model!L6:L8" in seen["turn"] and "every cell of L6:L8" in seen["turn"]


@pytest.mark.skipif(not com_available(), reason="needs Excel")
def test_recalculate_reports_the_array_a_cell_belongs_to(client, array_formula_xlsx):
    body = _upload(client, array_formula_xlsx)
    rc = client.post(f"/api/sessions/{body['sessionId']}/recalculate").json()
    assert rc["ran"] is True
    members = [e for e in rc["formula_errors"] if e["cell"] in ("C7", "C8", "C9")]
    assert len(members) == 3 and all(e["array"] == "C5:C9" for e in members)
    assert all("array" not in e for e in rc["formula_errors"] if e["cell"] == "E7")
    g = next(g for g in rc["groups"] if g["count"] == 3)
    assert g["arrays"] == ["Arr!C5:C9"] and g["cells"][0]["array"] == "C5:C9" and "array formula Arr!C5:C9" in g["cause"]
    # after a recalculation the workbook view shows the computed values (the errors) and the array's formula
    win = client.get(f"/api/sessions/{body['sessionId']}/cells", params={"sheet": "Arr", "cell": "C7", "rows": 0, "cols": 0}).json()
    c7 = win["rows"][0]["cells"][0]
    assert win["values_from"] == "recalculation" and c7["value"] == "#N/A" and c7["error"] is True and c7["formula"] == "{=A5:A6*2}" and c7["array"] == "C5:C9"


def test_cells_window_shows_contents_and_formulas(client, flagged_model_broken_xlsx):
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    win = client.get(f"/api/sessions/{sid}/cells", params={"sheet": "Model", "cell": "L5", "rows": 1, "cols": 1}).json()
    assert win["sheet"] == "Model" and win["focus"] == "L5" and win["columns"] == ["K", "L", "M"] and [r["row"] for r in win["rows"]] == [4, 5, 6]
    focus = [c for r in win["rows"] for c in r["cells"] if c["focus"]]
    assert len(focus) == 1 and focus[0]["ref"] == "L5" and focus[0]["formula"] == "=SUM(L4,MM_SETSIZE(3,1))"
    title = client.get(f"/api/sessions/{sid}/cells", params={"sheet": "Model", "cell": "A3", "rows": 0, "cols": 0}).json()
    assert title["rows"][0]["cells"][0]["value"] == "#Assumptions /Reorder /Inpt" and title["rows"][0]["cells"][0]["formula"] is None
    assert title["values_from"] == "analysis copy"
    bad = client.get(f"/api/sessions/{sid}/cells", params={"sheet": "Nope", "cell": "A1"}).json()
    assert "not found" in bad["error"]


# --- assistant grid naming (1.6.7) ----------------------------------------------------


def test_grid_names_endpoint_proposes_names_and_folds_them_into_the_plan(client, section_heading_grids_xlsx, monkeypatch):
    """The endpoint asks the assistant only about grids whose deterministic name
    is meaningless, and the accepted names change what the titles action writes."""
    body = _upload(client, section_heading_grids_xlsx)
    session_id = body["sessionId"]

    seen: dict = {}

    def fake_chat(messages, system_prompt, **kw):
        seen["prompt"] = messages[0]["content"]
        return {"available": True, "text": '{"CF!B8:C8": "Section Two Heading"}', "message": None}

    monkeypatch.setattr(server, "suggest_names", lambda contexts, **kw: __import__("app.grid_naming", fromlist=["x"]).suggest_names(contexts, completion=fake_chat, **kw))

    res = client.post(f"/api/sessions/{session_id}/grid-names", json={"apply": True})
    assert res.status_code == 200, res.text
    out = res.json()
    assert out["available"] and out["applied"]
    rows = {r["grid"]: r for r in out["names"]}
    # the block with a real heading above it is not sent for renaming
    assert "CF!C3:D5" not in rows
    # the fallback-named one is, and carries both names for review
    assert rows["CF!B8:C8"]["deterministic"] == "CF B8"
    assert rows["CF!B8:C8"]["suggested"] == "Section Two Heading"
    titles = {o["after"] for a in out["plan"] if a["id"] == "create_grid_titles" for o in a["operations"] if o["op"] == "set_value"}
    assert "#Section Two Heading" in titles and "#CF B8" not in titles
    assert "#Demographic Assumptions" in titles  # heading-derived name is untouched


def test_grid_names_endpoint_survives_an_unavailable_assistant(client, section_heading_grids_xlsx, monkeypatch):
    body = _upload(client, section_heading_grids_xlsx)
    monkeypatch.setattr(server, "suggest_names", lambda contexts, **kw: {"available": False, "names": {}, "message": "no secret.key found nearby"})
    out = client.post(f"/api/sessions/{body['sessionId']}/grid-names", json={}).json()
    assert out["available"] is False and out["names"] == [] and out["applied"] is False


# --- 1.7.1: upload size gate, sheet skipping, scan status ------------------------------------


def test_deferred_upload_reports_size_and_sheets_then_scans_without_the_skipped_ones(client, flagged_model_broken_xlsx, monkeypatch):
    monkeypatch.setenv("MIND_READY_SIZE_THRESHOLD_MB", "0.001")  # everything is "large"
    with flagged_model_broken_xlsx.open("rb") as f:
        res = client.post("/api/sessions", files={"file": (flagged_model_broken_xlsx.name, f, "application/octet-stream")}, data={"mode": "plan", "defer": "1"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["pending"] is True and body["mode"] == "plan" and "summary" not in body
    size = body["size"]
    assert size["above_threshold"] is True and size["measure"] == "decompressed" and size["decompressed_bytes"] > size["file_bytes"] > 0
    assert [s["name"] for s in size["sheets"]] == ["Model", "Settings", "Outputs"] and all(s["bytes"] > 0 for s in size["sheets"])
    assert body["version"]["id"] == "ver-001" and body["scan"]["state"] == "pending"
    sid = body["sessionId"]
    assert client.get(f"/api/sessions/{sid}/status").json()["state"] == "pending"
    # the ignore list is validated against the sheet list read from the package
    assert client.post(f"/api/sessions/{sid}/analyze", json={"ignore_sheets": ["Nope"]}).status_code == 422
    assert client.post(f"/api/sessions/{sid}/analyze", json={"ignore_sheets": ["Model", "Settings", "Outputs"]}).status_code == 422
    assert client.post(f"/api/sessions/{sid}/analyze", json={"ignore_sheets": 5}).status_code == 422
    assert client.get(f"/api/sessions/{sid}/status").json()["state"] == "pending"  # refused requests never started a scan
    out = client.post(f"/api/sessions/{sid}/analyze", json={"ignore_sheets": ["Outputs"]})
    assert out.status_code == 200, out.text
    out = out.json()
    assert set(out) >= {"sessionId", "summary", "report", "plan", "delta", "version", "versions", "size", "scan"}
    summary = out["summary"]
    assert [s["name"] for s in summary["sheets"]] == ["Model", "Settings"]
    assert summary["ignored_sheets"] == [{"name": "Outputs", "state": "visible"}]
    assert (summary["sheet_count"], summary["ignored_sheet_count"], summary["total_sheet_count"]) == (2, 1, 3)
    assert summary["size"]["above_threshold"] is True
    assert not any(f["location"].get("sheet") == "Outputs" for f in out["report"]["findings"])
    assert out["report"]["summary"]["finding_count"] == 101  # every rule still runs
    scan = client.get(f"/api/sessions/{sid}/status").json()
    assert scan["state"] == "ready" and scan["overall"] == 1.0 and scan["error"] is None and scan["elapsed_s"] >= 0
    assert {s["stage"] for s in scan["stages"]} >= {"copy", "load", "inventory", "names", "rules", "report", "plan"}
    assert scan["ignore_sheets"] == ["Outputs"] and scan["needs_convert"] is False and scan["version_id"] == "ver-001"
    assert out["scan"]["state"] == "ready"
    # a re-analysis keeps the ignore list and is tracked the same way
    again = client.post(f"/api/sessions/{sid}/reanalyze", json={}).json()
    assert again["summary"]["ignored_sheet_count"] == 1 and again["scan"]["state"] == "ready" and again["version"]["id"] == "ver-001"
    # scanning again with another list re-analyses the current version
    swapped = client.post(f"/api/sessions/{sid}/analyze", json={"ignore_sheets": ["Settings"]}).json()
    assert [s["name"] for s in swapped["summary"]["sheets"]] == ["Model", "Outputs"] and swapped["delta"]["previous_version_id"] == "ver-001"
    health = client.get("/api/health").json()
    assert health["upload_gate"] is True and health["size_threshold_mb"] == pytest.approx(0.001, abs=0.1)


def test_one_shot_upload_carries_size_facts_and_honours_ignore_sheets(client, flagged_model_broken_xlsx, plain_grid_xlsx, monkeypatch):
    monkeypatch.delenv("MIND_READY_SIZE_THRESHOLD_MB", raising=False)
    body = _upload(client, plain_grid_xlsx)
    assert "pending" not in body and body["size"]["above_threshold"] is False and body["summary"]["ignored_sheet_count"] == 0
    assert body["scan"]["state"] == "ready" and body["summary"]["size"]["threshold_mb"] == 25.0 and body["versions"][0]["id"] == "ver-001"
    with flagged_model_broken_xlsx.open("rb") as f:
        res = client.post("/api/sessions", files={"file": (flagged_model_broken_xlsx.name, f, "application/octet-stream")}, data={"mode": "plan", "ignore_sheets": '["Settings"]'})
    assert res.status_code == 200, res.text
    out = res.json()
    assert [s["name"] for s in out["summary"]["sheets"]] == ["Model", "Outputs"] and out["summary"]["ignored_sheets"][0]["name"] == "Settings"
    assert client.get(f"/api/sessions/{out['sessionId']}/status").json()["ignore_sheets"] == ["Settings"]
    with flagged_model_broken_xlsx.open("rb") as f:
        assert client.post("/api/sessions", files={"file": ("x.xlsx", f, "application/octet-stream")}, data={"mode": "plan", "ignore_sheets": '["Nope"]'}).status_code == 422


# --- 1.7.2: the readiness verdict and the Prep gauge travel with every analysis ------------
def test_analysis_carries_the_verdict_levels_and_gauge(client, flagged_model_broken_xlsx):
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    rd, pp = body["readiness"], body["prep_progress"]
    assert rd["state"] in ("blocked", "unverified", "ready") and rd["version_id"] == "ver-001"
    assert set(rd) >= {"headline", "blocking", "blocking_count", "manual_blocking_count", "optional_count", "prep", "recalc", "next_step"}
    assert rd["recalc"]["ran"] is False and rd["state"] != "ready"
    for b in rd["blocking"]:
        assert b["fix"] in ("prep", "prep_skipped", "assistant")
    assert all(f.get("priority") in ("REQUIRED", "RECOMMENDED", "INFORMATIONAL") for f in body["report"]["findings"])
    assert all(a["level"] in ("blocking", "optional") for a in body["plan"])
    assert pp["entries"][-1]["version_id"] == "ver-001" and pp["stalled"] is False
    last = pp["entries"][-1]
    assert last["open_findings"] == sum(1 for f in body["report"]["findings"] if f["status"] != "PASS")
    assert sum(last["status_counts"].values()) == len(body["report"]["findings"])
    # the same verdict on demand
    got = client.get(f"/api/sessions/{sid}/readiness").json()
    assert got["readiness"]["state"] == rd["state"] and got["prep_progress"]["entries"] == pp["entries"]
    assert client.get("/api/sessions/nope/readiness").status_code == 404


def test_apply_says_what_it_did_repair_by_repair(client, flagged_model_broken_xlsx):
    """1.7.3: the Apply answer carries its outcome -- per repair, the changes
    written, the rule before and after, and what is left -- so the front-end
    never has to guess whether a blocking problem was resolved."""
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    plan = {a["id"]: a for a in body["plan"]}
    ops = plan["special_headers"]["operations"] + plan["flag_spelling"]["operations"]
    out = client.post(f"/api/sessions/{sid}/apply", json={"operations": ops, "reanalyze": True}).json()
    oc = out["outcome"]
    assert oc["reanalyzed"] is True and oc["version_id"] == "ver-002" and oc["previous_version_id"] == "ver-001"
    assert oc["sent"] == len(ops) and oc["applied"] == len(ops) and oc["failed"] == 0
    by_id = {a["id"]: a for a in oc["actions"]}
    assert set(by_id) == {"special_headers", "flag_spelling"}
    assert by_id["flag_spelling"]["verdict"] == "resolved" and by_id["flag_spelling"]["level"] == "optional"
    # two of the three rules it serves pass; the third still fails on fewer cells: said as it is
    headers = by_id["special_headers"]
    assert headers["verdict"] == "partial" and headers["level"] == "blocking" and headers["planned_after"] == 0
    assert "EXP-003 still ERROR" in headers["summary"] and "by hand" in headers["summary"]
    assert headers["applied"] == plan["special_headers"]["count"] == headers["planned_before"]
    moves = {r["rule_id"]: (r["from"], r["to"]) for r in headers["rules"]}
    assert moves["PAR-002"] == ("ERROR", "PASS") and "PAR-002 ERROR -> PASS" in headers["summary"]
    assert {"PAR-002", "PRJ-002"} <= set(oc["resolved_blocking"])
    assert oc["blocking_after"] == out["readiness"]["blocking_count"]
    assert oc["blocking_after"] == oc["blocking_before"] - len(oc["resolved_blocking"]) + sum(1 for r in oc["remaining_blocking"] if r["new"])
    assert {r["rule_id"] for r in oc["remaining_blocking"]} == {b["rule_id"] for b in out["readiness"]["blocking"]}
    assert oc["headline"].startswith(f"{len(ops)} of {len(ops)} changes written")
    # without a re-analysis only the writes are known, and the outcome says so
    raw = client.post(f"/api/sessions/{sid}/apply", json={"operations": plan["loop_name_case"]["operations"], "reanalyze": False}).json()
    assert raw["outcome"]["reanalyzed"] is False and raw["outcome"]["blocking_after"] is None
    assert [a["verdict"] for a in raw["outcome"]["actions"]] == ["written"]


def test_the_gauge_counts_the_changes_every_apply_wrote(client, flagged_model_broken_xlsx):
    """1.7.3: prep_progress carries one entry per Apply with the changes really
    written into the file, by level, and the totals of the session."""
    body = _upload(client, flagged_model_broken_xlsx)
    sid = body["sessionId"]
    assert body["prep_progress"]["applies"] == [] and body["prep_progress"]["written"]["applied"] == 0
    plan = {a["id"]: a for a in body["plan"]}
    first = plan["special_headers"]["operations"]
    out = client.post(f"/api/sessions/{sid}/apply", json={"operations": first, "reanalyze": True}).json()
    (a1,) = out["prep_progress"]["applies"]
    assert a1["version_id"] == "ver-002" and a1["previous_version_id"] == "ver-001"
    assert (a1["applied"], a1["failed"], a1["blocking_applied"], a1["optional_applied"]) == (len(first), 0, len(first), 0)
    assert {"PAR-002", "PRJ-002"} <= set(a1["resolved_blocking"])
    second = plan["flag_spelling"]["operations"]
    out2 = client.post(f"/api/sessions/{sid}/apply", json={"operations": second, "reanalyze": True}).json()
    progress = out2["prep_progress"]
    assert [a["version_id"] for a in progress["applies"]] == ["ver-002", "ver-003"]
    assert progress["written"]["applied"] == len(first) + len(second)
    assert progress["written"]["blocking_applied"] == len(first) and progress["written"]["optional_applied"] == len(second)
    # the same figures on demand
    assert client.get(f"/api/sessions/{sid}/readiness").json()["prep_progress"]["written"] == progress["written"]


def test_recalculate_step_fixes_are_minor_versions_and_a_recalculation_closes_the_round(client, plain_grid_xlsx):
    """1.7.4: v1 -> Prep v2 -> Recalculate-step fixes v2.1, v2.2 -> recalculate v3 -> fix v3.1."""
    sid = _upload(client, plain_grid_xlsx)["sessionId"]

    def apply(value, step):
        op = {"op": "set_value", "sheet": "Data", "cell": "Z99", "value": value, "after": value}
        out = client.post(f"/api/sessions/{sid}/apply", json={"operations": [op], "reanalyze": False, "versionStep": step})
        assert out.status_code == 200, out.text
        return out.json()["version"]

    v2 = apply("a", "major")
    assert (v2["id"], v2["label"].split(" — ")[0]) == ("ver-002", "v2")
    assert [apply(x, "minor")["label"].split(" — ")[0] for x in ("b", "c")] == ["v2.1", "v2.2"]
    if not com_available():
        pytest.skip("recalculation needs Excel")
    rec = client.post(f"/api/sessions/{sid}/recalculate").json()
    assert rec["version"]["id"] == "ver-003" and rec["version"]["label"] == "v3 — recalculated v2.2" and rec["version_id"] == "ver-003"
    assert "version" not in client.post(f"/api/sessions/{sid}/recalculate").json()  # nothing to close: v3 stays v3
    assert apply("d", "minor")["label"].startswith("v3.1 — ")


def test_a_clean_recalculation_turns_ready_001_into_a_pass(client, plain_grid_xlsx):
    """1.7.4: READY-001 is NOT_SUPPORTED in every analysis pass (no Excel run there),
    which kept the whole report NOT_SUPPORTED after a clean Recalculate."""
    body = _upload(client, plain_grid_xlsx)
    sid = body["sessionId"]
    ready = lambda report: next(f for f in report["findings"] if f["rule_id"] == "READY-001")  # noqa: E731
    assert ready(body["report"])["status"] == "NOT_SUPPORTED"
    if not com_available():
        pytest.skip("recalculation needs Excel")
    rec = client.post(f"/api/sessions/{sid}/recalculate").json()
    assert rec["ran"]
    expected = "PASS" if rec["readiness"]["recalc"]["clean"] else "ERROR"
    assert ready(rec["report"])["status"] == expected and "READY-001" not in rec["report"]["summary"]["not_supported_rule_ids"]
    # a re-analysis of the same version keeps what the recalculation found
    again = client.post(f"/api/sessions/{sid}/reanalyze", json={}).json()
    assert ready(again["report"])["status"] == expected


def test_a_clean_recalculation_turns_ready_001_to_pass_and_lifts_the_report(client, plain_grid_xlsx):
    """1.7.4: READY-001 is NOT_SUPPORTED in every analysis pass (it never runs
    Excel) and used to hold the whole report at NOT_SUPPORTED after a clean Recalculate."""
    body = _upload(client, plain_grid_xlsx)
    ready = next(f for f in body["report"]["findings"] if f["rule_id"] == "READY-001")
    assert ready["status"] == "NOT_SUPPORTED"
    if not com_available():
        pytest.skip("recalculation needs Excel")
    sid = body["sessionId"]
    rec = client.post(f"/api/sessions/{sid}/recalculate").json()
    if rec["formula_errors"]:
        pytest.skip("fixture recalculates with errors")
    report = rec["report"]
    ready = next(f for f in report["findings"] if f["rule_id"] == "READY-001")
    assert ready["status"] == "PASS" and "Recalculated v1" in ready["message"]
    assert "READY-001" not in report["summary"]["not_supported_rule_ids"]
    assert report["status"] != "NOT_SUPPORTED" or report["summary"]["not_supported_rule_ids"]
    # a re-analysis of the same version keeps the recalculation's verdict
    again = client.post(f"/api/sessions/{sid}/reanalyze", json={}).json()
    assert next(f for f in again["report"]["findings"] if f["rule_id"] == "READY-001")["status"] == "PASS"


def test_auto_fix_runs_in_the_background_and_leaves_a_clean_recalculated_version(client, tmp_path):
    """1.8.0: POST /auto-fix -> poll GET -> a new major version whose recalculation is the fixer's own,
    READY-001 PASS, and the manual way still there (the Fix panel's /apply is untouched)."""
    import time

    import openpyxl

    if not com_available():
        pytest.skip("the automatic fixer needs Excel")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "#Rates"
    ws["A2"], ws["B2"], ws["C2"] = "Premium", "Policies", "Average"
    for r, (a, b) in enumerate([(100, 4), (50, 0), (30, 3)], start=3):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = a, b, f"=A{r}/B{r}"
    ws["C6"] = "=SUM(C3:C5)"
    path = tmp_path / "rates.xlsx"
    wb.save(path)
    wb.close()

    sid = _upload(client, path)["sessionId"]
    assert client.get(f"/api/sessions/{sid}/auto-fix").json() == {"state": "idle"}
    started = client.post(f"/api/sessions/{sid}/auto-fix", json={"use_assistant": False}).json()
    assert started["state"] == "running" and started["use_assistant"] is False and started["keep_good_values"] is True
    assert client.post(f"/api/sessions/{sid}/auto-fix", json={}).status_code == 409  # one at a time
    # ... and nothing else moves the workbook meanwhile: its version, recalculation and analysis come at the end
    assert client.post(f"/api/sessions/{sid}/recalculate").status_code == 409
    assert client.post(f"/api/sessions/{sid}/apply", json={"operations": [{"op": "set_value", "sheet": "Data", "cell": "Z9", "after": 1}]}).status_code == 409
    deadline = time.time() + 300
    while True:
        st = client.get(f"/api/sessions/{sid}/auto-fix").json()
        if st["state"] != "running":
            break
        assert time.time() < deadline, st
        time.sleep(0.5)
    assert st["state"] == "done", st
    assert "fixes" not in st["result"] and st["result"]["status"] == "clean"  # the light answer while polling
    full = client.get(f"/api/sessions/{sid}/auto-fix?full=1").json()
    result = full["result"]
    assert result["errors_before"] == 2 and result["errors_after"] == 0 and result["cells_rewritten"] == 1
    assert result["fixes"][0]["after"] == "=IFERROR(A4/B4,0)" and result["left"] == []
    version = full["version"]
    assert version["source"] == "autofix" and version["label"].startswith("v2 — Auto-fix: 1 cell(s) fixed · recalculated clean")
    assert version["change_log"][0]["action_id"] == "autofix"
    # the run's own recalculation is the session's, for the new version; the analysis is the new version's
    assert full["recalc"]["ran"] is True and full["recalc"]["formula_errors"] == [] and full["recalc"]["version_id"] == version["id"]
    ready = next(f for f in full["analysis"]["report"]["findings"] if f["rule_id"] == "READY-001")
    assert ready["status"] == "PASS"
    assert client.get(f"/api/sessions/{sid}/readiness").json()["readiness"]["recalc"]["clean"] is True
    assert [v["id"] for v in client.get(f"/api/sessions/{sid}/versions").json()][-1] == version["id"]
    # nothing left to fix: a second run changes nothing and makes no version
    client.post(f"/api/sessions/{sid}/auto-fix", json={"use_assistant": False})
    while client.get(f"/api/sessions/{sid}/auto-fix").json()["state"] == "running":
        time.sleep(0.5)
    again = client.get(f"/api/sessions/{sid}/auto-fix?full=1").json()
    assert again["result"]["status"] == "unchanged" and again["version"] is None
    assert client.get("/api/health").json()["autofix"] is True


def test_numbers_check_says_what_the_preparation_changed_and_the_fixer_does_not_hide_it(client, tmp_path):
    """1.8.0: POST /numbers-check compares the current version with the original, cell by cell
    through the rows Prep inserted; an error that was a value in the original is not given a fallback."""
    import time

    import openpyxl

    if not com_available():
        pytest.skip("needs Excel")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "#Rates"
    ws["A2"], ws["B2"], ws["C2"] = "Premium", "Policies", "Average"
    for r, (a, b) in enumerate([(100, 4), (50, 0), (30, 3)], start=3):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = a, b, f"=A{r}/B{r}"   # C4 is an error of the model's own
    path = tmp_path / "rates.xlsx"
    wb.save(path)
    wb.close()

    sid = _upload(client, path)["sessionId"]
    first = client.post(f"/api/sessions/{sid}/numbers-check").json()
    assert first["ran"] is False and "original" in first["verdict"]
    # a row inserted above the table: every value is where it was, one row lower
    ok = client.post(f"/api/sessions/{sid}/apply", json={"operations": [{"op": "insert_row", "sheet": "Data", "row": 2}], "reanalyze": False}).json()
    assert ok["version"]["parent"] == "ver-001"
    numbers = client.post(f"/api/sessions/{sid}/numbers-check").json()
    assert numbers["ran"] and numbers["clean"] is True and numbers["moves"]["row_inserts"] == 1 and numbers["counts"]["error_to_error"] == 1
    # then a change that breaks a formula that worked: B6 (was B5 = 3) emptied -> C6 = 30/0
    client.post(f"/api/sessions/{sid}/apply", json={"operations": [{"op": "set_value", "sheet": "Data", "cell": "B6", "after": 0}], "reanalyze": False})
    numbers = client.post(f"/api/sessions/{sid}/numbers-check").json()
    assert numbers["clean"] is False and numbers["counts"]["good_to_error"] == 1 and numbers["regressions"] == 1
    assert numbers["samples"]["good_to_error"][0] == {"original": "Data!C5", "now": "Data!C6", "was": "10", "is": "#DIV/0!"}
    # the fixer: the model's own error (C5, was C4) is fixed; the one the changes made (C6) stays, with what it was
    client.post(f"/api/sessions/{sid}/auto-fix", json={"use_assistant": False})
    deadline = time.time() + 300
    while client.get(f"/api/sessions/{sid}/auto-fix").json()["state"] == "running":
        assert time.time() < deadline
        time.sleep(0.5)
    full = client.get(f"/api/sessions/{sid}/auto-fix?full=1").json()
    result = full["result"]
    assert full["state"] == "done" and result["status"] == "partial" and result["regression_cells"] == 1
    assert [g["cell"] for g in result["fixes"]] == ["C5"] and [g["cell"] for g in result["left"]] == ["C6"]
    assert "it was 10 in the original workbook" in result["left"][0]["why"]
    assert result["numbers_check"]["clean"] is False and result["numbers_check"]["counts"]["good_to_error"] == 1
