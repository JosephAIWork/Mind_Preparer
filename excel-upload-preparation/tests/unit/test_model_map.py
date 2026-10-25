"""Model map (1.9.0): where a workbook's inputs, loop driver, run settings and outputs are.

The macro parser is tested on VBA written here in the shapes the four client models use
(a row-per-record loop with sheet bounds; a named driver with named bounds and a paste per
item; nested scenario / year / cluster loops over named lists; a sheet-building loop). The
four models themselves are checked when their local copies are present (never committed)."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from local_usecases import usecase_or_skip

from app.inventory import build_analysis
from app.model_map import build_model_map, describe, parse_macros, sheet_role


def _names(*pairs: tuple[str, str]) -> list[dict]:
    return [{"name": n, "value": v, "local_sheet_id": None} for n, v in pairs]


ROW_LOOP_STYLE = '''
Sub Button2_Click()
first_row = Worksheets("Calculation").Cells(5, 20).Value
last_row = Worksheets("Calculation").Cells(6, 20).Value
'Clear the output
Worksheets("OUTPUT").Range("A" & first_row & ":I100000").ClearContents
For ins = first_row To (last_row)
    Application.StatusBar = "MP : " & ins
    Worksheets("Calculation").Range("A2:NI2") = Worksheets("Policy Data").Range("A" & ins & ":NI" & ins).Value
    Worksheets("OUTPUT").Range("A" & ins & ":J" & ins) = Worksheets("Calculation").Range("AY35:BH35").Value
Next
End Sub
'''

COUNTER_STYLE = '''
Sub RunPolicies()
Dim MacroStart As Long: MacroStart = Range("MacroStart").Value     'first pol
Dim MacroEnd As Long: MacroEnd = Range("MacroEnd").Value
If Range("DeleteOld").Value = "Yes" Then
    Range("ByGroupOutput").ClearContents
End If
For polNo = MacroStart To MacroEnd
    Range("POLICY_ROW").Value = polNo
    If Range("Run_All?").Value = "No" Then
        GoTo SkipPolicy
    End If
    Application.Calculate
    Range("ByGroupPaster").Offset(0, 10 * (Temp2 - 1)).Value2 = Range("AggregateOutput").Value2
SkipPolicy:
Next polNo
Call QuickPaste
End Sub
Sub SaveTotal()
Dim SaveAs_txt As String:  SaveAs_txt = Range("SaveNameTotal")
    Sheets(Array("Model Output Pasted")).Copy
    ActiveWorkbook.SaveAs Filename:=SaveAs_txt, FileFormat:=52
End Sub
Sub QuickPaste()
    Range("ModelOutput").Copy
    Sheets("Model Output Pasted").Select
    Range("PasteOutput").Select
        Selection.PasteSpecial Paste:=xlPasteValues, Operation:=xlNone, SkipBlanks _
        :=False, Transpose:=False
End Sub
'''

LISTS_STYLE = '''
Sub baseline_projections()
Dim paste_start_cell As Range
For i_scenario = Range("start_1").Value2 To Range("end_1").Value2
    Range("active_scenario").Value2 = Range("scenario_start").Offset(0, i_scenario - 1).Value2
    For i_year = 1 To 5
        Range("active_year").Value2 = Range("proj_year").Value2 + i_year - 1
        For i_cluster = 1 To 20
            Range("active_cluster").Value2 = Range("Cluster1").Offset(i_cluster - 1, 0).Value2
            Range("GroupA_paste_start").Offset((i_year - 1) * 5, (i_cluster - 1) * 3).Value2 = Range("copyA_final").Value2
            Set paste_start_cell = Range("paste_output").Cells(1, 1).Offset(236 * (tracker - 1), 0)
            paste_start_cell.Resize(Range("copy_output").Rows.Count, Range("copy_output").Columns.Count).Value2 = Range("copy_output").Value2
        Next i_cluster
    Next i_year
Next i_scenario
End Sub
Sub combined_macro()
Call baseline_projections
End Sub
'''

BUILDER_STYLE = '''
Sub Macro1()
For i = 11 To 40
    Sheets("Line 1").Copy Before:=Sheets(">>")
    Sheets("Line 1 (2)").Name = "Line " & i
    Range("G2").Value = "Line " & i
Next i
End Sub
'''


def test_a_row_per_record_loop_with_sheet_bounds_is_read():
    v = parse_macros([("Module1", ROW_LOOP_STYLE)], [])
    assert len(v["loops"]) == 1
    loop = v["loops"][0]
    assert loop["var"] == "ins" and loop["from"]["cell"] == "Calculation!T5" and loop["to"]["cell"] == "Calculation!T6"
    d = loop["drivers"][0]
    assert d["cell"] == "Calculation!A2:NI2" and d["set_from"] == "'Policy Data'!A<ins>:NI<ins>"
    o = loop["outputs"][0]
    assert o["to"] == "OUTPUT!A<ins>:J<ins>" and o["from"] == "Calculation!AY35:BH35" and o["per_item"]
    assert v["cleared"][0]["cell"] == "OUTPUT!A<first_row>:I100000"
    assert v["entry_points"] == ["Button2_Click"]


def test_a_named_driver_with_named_bounds_switches_and_a_paste_per_item():
    names = _names(("MacroStart", "Macro!$B$1"), ("MacroEnd", "Macro!$B$2"), ("POLICY_ROW", "Projection!$C$3"), ("DeleteOld", "Macro!$B$3"), ("Run_All?", "Macro!$B$4"),
                   ("ByGroupPaster", "Macro!$V$13:$AE$140"), ("AggregateOutput", "Macro!$A$11:$J$138"), ("ByGroupOutput", "Macro!$V$13:$JU$140"),
                   ("ModelOutput", "'Model Output'!$M$9:$EK$48"), ("PasteOutput", "'Model Output Pasted'!$A$1:$DY$40"), ("SaveNameTotal", "'Saved paths'!$C$9"))
    v = parse_macros([("Module1", COUNTER_STYLE)], names)
    loop = v["loops"][0]
    assert loop["from"] == {"text": "MacroStart = Range(\"MacroStart\").Value", "kind": "cell", "cell": "Macro!B1", "name": "MacroStart"}
    assert loop["to"]["cell"] == "Macro!B2"
    assert loop["drivers"][0]["cell"] == "Projection!C3" and loop["drivers"][0]["name"] == "POLICY_ROW" and loop["drivers"][0]["how"] == "the loop value itself"
    assert loop["outputs"][0]["to"] == "Macro!V13:AE140" and loop["outputs"][0]["from"] == "Macro!A11:J138" and loop["outputs"][0]["per_item"] is False
    sw = {s["name"]: s for s in v["switches"]}
    assert sw["Run_All?"]["cell"] == "Macro!B4" and sw["Run_All?"]["compared_with"] == ["No"] and sw["DeleteOld"]["compared_with"] == ["Yes"]
    assert v["cleared"][0]["name"] == "ByGroupOutput"
    assert v["pastes"] == [{"from_name": "ModelOutput", "from": "'Model Output'!M9:EK48", "in": "QuickPaste", "to_name": "PasteOutput", "to": "'Model Output Pasted'!A1:DY40"}]
    assert v["entry_points"] == ["RunPolicies"]
    assert v["saves"] == [{"in": "SaveTotal", "path_from": "'Saved paths'!C9", "name": "SaveNameTotal"}]


def test_nested_loops_over_named_lists_and_a_block_per_pass():
    names = _names(("start_1", "Output!$G$2"), ("end_1", "Output!$G$3"), ("active_scenario", "'Model Inputs'!$C$8"), ("scenario_start", "'Model Inputs'!$E$8"),
                   ("active_year", "'Model Inputs'!$C$173"), ("proj_year", "'Assumption Setting'!$D$8"), ("active_cluster", "'Model Inputs'!$C$174"), ("Cluster1", "'Validation list'!$B$7"),
                   ("GroupA_paste_start", "'Cost Baseline'!$V$12"), ("copyA_final", "'Cost Calculation'!$DI$13:$DK$118"), ("paste_output", "Output!$B$9:$U$244"), ("copy_output", "'Output working'!$AB$13:$AU$248"))
    v = parse_macros([("Module1", LISTS_STYLE)], names)
    loops = {l["var"]: l for l in v["loops"]}
    assert loops["i_scenario"]["from"]["cell"] == "Output!G2" and loops["i_scenario"]["to"]["cell"] == "Output!G3"
    assert loops["i_year"]["from"] == {"text": "1", "kind": "literal", "value": 1} and loops["i_cluster"]["to"]["value"] == 20
    assert loops["i_scenario"]["drivers"][0]["cell"] == "'Model Inputs'!C8" and loops["i_scenario"]["drivers"][0]["set_from_name"] == "scenario_start"
    assert loops["i_year"]["drivers"][0]["cell"] == "'Model Inputs'!C173"
    assert loops["i_cluster"]["drivers"][0]["cell"] == "'Model Inputs'!C174" and loops["i_cluster"]["drivers"][0]["set_from_name"] == "Cluster1"
    outs = {(o["to_name"], o["from_name"]) for l in v["loops"] for o in l["outputs"]}
    assert ("GroupA_paste_start", "copyA_final") in outs and ("paste_output", "copy_output") in outs
    assert v["entry_points"] == ["combined_macro"]


def test_a_loop_that_builds_sheets_is_not_a_run_loop():
    v = parse_macros([("Module2", BUILDER_STYLE)], [])
    loop = v["loops"][0]
    assert loop["builds_sheets"] and loop["copies_sheet"] == "Line 1"
    assert loop["drivers"][0]["cell"] == "'Line 1'!G2"
    assert v["entry_points"] == []


def test_sheet_roles_by_name():
    assert sheet_role("Inputs_Model") == "input" and sheet_role("OUTPUT") == "output" and sheet_role("Calculation") == "calculation"
    assert sheet_role(">> Inputs", "A1:A1") == "divider" and sheet_role("Data inputs-->") == "divider"
    assert sheet_role("Other assumptions") == "input" and sheet_role("Reporting_Line Total") == "output" and sheet_role("Saved paths") == "output"


def test_a_workbook_without_macros_gets_dropdown_switches_named_settings_and_key_cells(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    inp = wb.active
    inp.title = "Inputs"
    inp["B3"], inp["C3"] = "Basis", "SCR"
    dv = DataValidation(type="list", formula1='"SCR,MSR"')
    inp.add_data_validation(dv)
    dv.add("C3")
    inp["B4"], inp["C4"] = "Last year", 2016
    inp["B6"], inp["C6"] = "Threshold", 0.001
    wb.defined_names["Last_Year"] = DefinedName("Last_Year", attr_text="Inputs!$C$4")
    calc = wb.create_sheet("Calc")
    calc["G2"] = "Line 1"
    for r in range(4, 40):
        calc[f"B{r}"] = f'="x#"&$G$2&{r}'
        calc[f"C{r}"] = f"=Last_Year+{r}"
    rep = wb.create_sheet("Reporting")
    rep["B2"] = "=SUM(Calc!C4:C39)"
    p = tmp_path / "nomacro.xlsx"
    wb.save(p)
    wb.close()
    analysis = build_analysis(p, tmp_path / "w", "m")
    mp = build_model_map(analysis)
    assert mp["has_macros"] is False and mp["loop"]["found"] is False
    assert mp["loop"]["key_cells"][0]["cell"] == "Calc!G2" and mp["loop"]["key_cells"][0]["readers"] == 36
    sw = mp["run_settings"]["switches"]
    assert sw[0]["cell"] == "Inputs!C3" and sw[0]["options"] == ["SCR", "MSR"] and sw[0]["label"] == "Basis" and sw[0]["value"] == "SCR"
    named = {n["name"]: n for n in mp["run_settings"]["named_settings"]}
    assert named["Last_Year"]["cell"] == "Inputs!C4" and named["Last_Year"]["readers"] == 36
    lab = {x["label"]: x for x in mp["run_settings"]["labelled"]}
    assert lab["Threshold"]["cell"] == "Inputs!C6" and lab["Threshold"]["value"] == 0.001
    assert [s["sheet"] for s in mp["inputs"]["sheets"]] == ["Inputs"] and [s["sheet"] for s in mp["outputs"]["sheets"]] == ["Reporting"]
    text = describe(mp)
    assert "key cell  Calc!G2" in text and "switch    Inputs!C3" in text


# --- the four client models, when their local copies are here ---------------------------------------
# Their files and the sheet / cell / name values expected of them are kept next to the copies
# (tests/local_usecases.py), never in the repo.
def _map_of(code: str, tmp_path: Path) -> tuple[dict, dict]:
    path, n = usecase_or_skip(code)
    return build_model_map(build_analysis(path, tmp_path / "w", code)), n


def test_uc1_map_a_row_per_record_loop(tmp_path):
    mp, n = _map_of("UC-1", tmp_path)
    lp = mp["loop"]
    assert lp["found"] and lp["drivers"][0]["cell"] == n["driver"] and lp["drivers"][0]["set_from"] == n["driver_from"]
    bounds = {b["role"]: b for b in lp["bounds"]}
    assert bounds["first"]["cell"] == n["first_cell"] and bounds["first"]["label"] == n["first_label"] and bounds["last"]["cell"] == n["last_cell"] and bounds["last"]["formula"] == n["last_formula"]
    assert lp["entry_points"] == n["entry_points"]
    assert mp["inputs"]["per_item"][0]["cell"] == n["driver_from"]
    assert {s["sheet"] for s in mp["inputs"]["sheets"]} >= set(n["input_sheets"])
    outs = {o["cell"]: o for o in mp["outputs"]["cells"]}
    assert outs[n["output"]]["from"] == n["output_from"]
    assert any(x["sheet"] == n["settings_sheet"] for x in mp["run_settings"]["labelled"])


def test_uc4_map_no_run_loop_but_dropdowns_and_named_settings(tmp_path):
    mp, n = _map_of("UC-4", tmp_path)
    lp = mp["loop"]
    assert lp["found"] is False
    assert {b["cell"] for b in lp["item_cells"]} == set(n["item_cells"])
    sw = {s["name"]: s for s in mp["run_settings"]["switches"] if s.get("name")}
    for name, (cell, options) in n["named_switches"].items():
        assert sw[name]["cell"] == cell and sw[name]["options"] == options
    cell, options = n["option_switch"]
    assert any(s["cell"] == cell and s["options"] == options for s in mp["run_settings"]["switches"])
    named = {x["name"]: x for x in mp["run_settings"]["named_settings"]}
    assert all(named[name]["cell"] == cell for name, cell in n["named_settings"].items())
    assert {s["sheet"] for s in mp["inputs"]["sheets"]} == set(n["input_sheets"])
    assert {s["sheet"] for s in mp["outputs"]["sheets"]} == set(n["output_sheets"])


def test_uc2_map_a_counter_driver_with_named_bounds(tmp_path):
    mp, n = _map_of("UC-2", tmp_path)
    lp = mp["loop"]
    d = {x["name"]: x for x in lp["drivers"]}
    assert all(d[name]["cell"] == cell for name, cell in n["drivers"].items())
    bounds = {b["role"]: b for b in lp["bounds"]}
    assert bounds["first"]["name"] == n["first_name"] and bounds["first"]["cell"] == n["first_cell"] and bounds["last"]["cell"] == n["last_cell"]
    assert lp["entry_points"] == n["entry_points"]
    sw = {s["name"]: s for s in mp["run_settings"]["switches"] if s.get("name")}
    assert all(sw[name]["cell"] == cell for name, cell in n["switches"].items())
    drop = sw[n["dropdown_switch"]]
    assert drop["how"].startswith("a dropdown in the cell") and drop["options"] == n["dropdown_options"]
    outs = {o["name"]: o for o in mp["outputs"]["cells"] if o.get("name")}
    assert all(outs[name]["from_name"] == source for name, source in n["outputs_from"].items())
    assert n["cleared_output"] in outs
    tables = {t["name"] for t in mp["inputs"]["tables"]}
    assert n["table"] in tables or any(s["sheet"] == n["table_sheet"] for s in mp["inputs"]["sheets"])
    assert any(sv["name"] == n["save"] for sv in mp["outputs"]["saves"])


def test_uc3_map_nested_loops_over_named_lists(tmp_path):
    mp, n = _map_of("UC-3", tmp_path)
    lp = mp["loop"]
    d = {x["name"]: x for x in lp["drivers"]}
    assert all(d[name]["cell"] == cell for name, cell in n["drivers"].items())
    assert all(d[name]["set_from_name"] == source for name, source in n["set_from"].items())
    bounds = {(b["role"], b["name"]) for b in lp["bounds"]}
    assert all((role, name) in bounds for role, name in n["bounds"])
    assert n["entry_point"] in lp["entry_points"]
    outs = {o["name"]: o for o in mp["outputs"]["cells"] if o.get("name")}
    assert all(outs[name]["from_name"] == source for name, source in n["outputs_from"].items())
    sw = {s["cell"]: s for s in mp["run_settings"]["switches"]}
    assert all(cell in sw for cell in n["switch_cells"])
    assert {s["sheet"] for s in mp["outputs"]["sheets"]} >= set(n["output_sheets"])
