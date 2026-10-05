"""1.6.8: a grid title must never change a computed value. Found by the
autonomous Mind loop's numbers gate on the real Shlomo model: a title written
into column A of the policy table was counted by `=COUNTA(CoverageDetails!A:A)`
and the model picked a different policy; the company name in Information!A1
was "moved" into a title while five sheets read `=Information!A1`."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app import prep
from app.inventory import build_analysis


def _wb(path: Path, sheets: dict[str, dict[str, object]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for ref, v in cells.items():
            ws[ref] = v
    wb.save(path)
    wb.close()
    return path


def _titles(analysis):
    ops, skipped = prep.plan_create_grid_titles(analysis, {"findings": []})
    return ops, skipped


def test_no_row_insert_and_no_title_on_a_sheet_counted_by_whole_column_formulas(tmp_path):
    """The policy table starts on row 1 and another sheet sizes the model with
    COUNTA over its column A: inserting a row or writing a title there would
    change that count."""
    path = _wb(tmp_path / "m.xlsx", {
        "Data": {"A1": "Id", "B1": "Value", "A2": 1, "B2": 10, "A3": 2, "B3": 20},
        "Calc": {"B2": "Size", "C2": "=COUNTA(Data!A:A)"},
    })
    analysis = build_analysis(path, tmp_path / "w", "t")
    index = prep.reference_index(analysis)
    assert prep.whole_reference_to(index, "Data") == "Calc!C2"
    ops, skipped = _titles(analysis)
    assert not any(o["sheet"] == "Data" for o in ops), ops
    assert any("Data!A1:B3" in s and "counts or indexes over whole columns" in s and "Calc!C2" in s for s in skipped), skipped


def test_no_title_inside_a_range_a_formula_reads(tmp_path):
    """An INDEX/MATCH lookup over Parameters!C10:C40 must not start returning a
    '#Title' written into that range."""
    cells = {"B9": "Translation", "C9": "Key"}  # a titled-looking header row with a lookup range below
    for r in range(10, 20):
        cells[f"B{r}"] = f"k{r}"
        cells[f"C{r}"] = f"v{r}"
    cells["B24"] = "x"
    cells["C24"] = 1
    cells["B25"] = "y"
    cells["C25"] = 2
    path = _wb(tmp_path / "m.xlsx", {
        "Parameters": cells,
        "Calc": {"A1": "=INDEX(Parameters!$C$10:$C$40,MATCH(\"k12\",Parameters!$B$10:$B$40,0))"},
    })
    analysis = build_analysis(path, tmp_path / "w", "t")
    ops, skipped = _titles(analysis)
    # the block at B24:C25 lies inside the lookup range: no title may be written above it (B23 is inside C10:C40? no -- B column) -> B23 is read by the MATCH range
    assert not any(o["sheet"] == "Parameters" and o.get("cell") == "B23" for o in ops), ops
    assert any("B24:C25" in s and "is read by Calc!A1" in s for s in skipped), skipped


def test_a_label_something_reads_is_never_cleared(tmp_path):
    """The label's text still names the grid, but the cell stays put."""
    path = _wb(tmp_path / "m.xlsx", {
        "Info": {"B2": "Company Name", "B4": "Key", "C4": "Value", "B5": "a", "C5": 1},
        "Other": {"A1": "=Info!B2"},
    })
    analysis = build_analysis(path, tmp_path / "w", "t")
    ops, skipped = _titles(analysis)
    titles = {o["after"] for o in ops if o["op"] == "set_value"}
    assert "#Company Name" in titles
    assert not any(o["op"] == "clear_cell" for o in ops), ops
    assert any("Info!B2" in s and "label kept in place" in s and "Other!A1" in s for s in skipped), skipped


def test_a_caption_a_formula_reads_is_not_turned_into_a_title(tmp_path):
    # a caption directly above a table that starts in the same column: read as a
    # one-column grid B2:B4 swallowing the table's first column (case A)
    path = _wb(tmp_path / "m.xlsx", {
        "S": {"B2": "Rates", "B3": 1, "C3": 2, "B4": 3, "C4": 4},
        "T": {"A1": "=S!B2"},
    })
    analysis = build_analysis(path, tmp_path / "w", "t")
    ops, skipped = _titles(analysis)
    assert not any(o.get("cell") == "B2" and str(o.get("after", "")).startswith("#") for o in ops)
    assert any("S!B2" in s and "caption is read by T!A1" in s for s in skipped), skipped


def test_reference_index_reads_defined_names_and_skips_broken_ones(tmp_path):
    path = _wb(tmp_path / "m.xlsx", {"S": {"A1": 1, "A2": 2}})
    wb = openpyxl.load_workbook(path)
    from openpyxl.workbook.defined_name import DefinedName

    wb.defined_names["KeyDate"] = DefinedName("KeyDate", attr_text="S!$A$1")
    wb.defined_names["Broken"] = DefinedName("Broken", attr_text="S!#REF!")
    wb.save(path)
    wb.close()
    analysis = build_analysis(path, tmp_path / "w", "t")
    index = prep.reference_index(analysis)
    assert prep.referenced_by(index, "S", 1, 1) == "name KeyDate"
    assert prep.referenced_by(index, "S", 2, 1) is None


@pytest.mark.skipif(not __import__("app.excel_com", fromlist=["com_available"]).com_available(), reason="needs Excel")
def test_titles_leave_every_computed_value_unchanged_end_to_end(tmp_path):
    """Apply the safe titles through Excel and prove, with the loop's own
    numbers gate, that no value moved."""
    from app import mind_loop as ml

    path = _wb(tmp_path / "m.xlsx", {
        "Data": {"A1": "Id", "B1": "Value", "A2": 1, "B2": 10, "A3": 2, "B3": 20},
        "Calc": {"B2": "Size", "C2": "=COUNTA(Data!A:A)", "B4": "Total", "C4": "=SUM(Data!B:B)", "B7": "Section", "B8": "k", "C8": "=C4*2"},
    })
    analysis = build_analysis(path, tmp_path / "w", "t")
    ops, _ = _titles(analysis)
    assert ops
    out = prep.apply_operations(path, tmp_path / "prep", ops)
    assert out["status"] == "APPLIED"
    gate = ml.numbers_gate(path, Path(out["output_path"]), out["applied"], analysis, tmp_path / "numbers")
    assert gate["ran"] and gate["match"], gate


# --- 1.8.0: what Excel does not move along when a row is inserted ---------------------------
def test_no_row_insert_on_a_sheet_inside_a_3d_reference(tmp_path):
    """=SUM('LoB 1:>>'!B2) reads B2 on every sheet from 'LoB 1' to '>>', by
    position. A row inserted on 'LoB 1' alone shifts that sheet and the total
    reads another line (the PVFP model: 2,462 computed values changed)."""
    lob = {"A1": "Line", "B1": "2024", "A2": "Premium", "B2": 100, "A3": "Claims", "B3": 60}
    path = _wb(tmp_path / "m.xlsx", {
        "LoB 1": dict(lob), "LoB 2": dict(lob), ">>": {},
        "Total": {"A2": "Premium", "B2": "=SUM('LoB 1:>>'!B2)", "A3": "Claims", "B3": "=SUM('LoB 1:>>'!B3)"},
        "Other": {"A1": "Id", "B1": "Value", "A2": 1, "B2": 10},
    })
    analysis = build_analysis(path, tmp_path / "w", "t")
    index = prep.reference_index(analysis)
    assert set(index["positional"]) == {"LoB 1", "LoB 2", ">>"} and "Total!B2" in index["positional"]["LoB 1"]
    assert "3-D reference" in prep.insert_blocker(index, "LoB 2") and prep.insert_blocker(index, "Other") is None
    assert prep.referenced_by(index, "LoB 2", 2, 2) == "Total!B2"  # the cells it reads are known on every sheet of the range
    ops, skipped = _titles(analysis)
    assert not any(o["op"] == "insert_row" and o["sheet"] in ("LoB 1", "LoB 2") for o in ops), ops
    assert any(o["op"] == "insert_row" and o["sheet"] == "Other" for o in ops), ops  # a sheet nothing reads by position still gets its title row
    assert any(s.startswith("LoB 1!") and "3-D reference" in s and "Total!B2" in s for s in skipped), skipped


def test_no_row_insert_on_a_sheet_an_indirect_addresses_as_text(tmp_path):
    """=INDIRECT("'"&A1&"'!B"&A2) is text: Excel does not move it when a row is
    inserted on the sheet it names -- named here by the cell A1, and by a text
    piece ("'CoC "&n) for the CoC sheets."""
    block = {"A1": "Line", "B1": "2024", "A2": "Premium", "B2": 100, "A3": "Claims", "B3": 60}
    path = _wb(tmp_path / "m.xlsx", {
        "LoB 1": dict(block), "CoC 1": dict(block), "CoC 2": dict(block), "Free": dict(block),
        "Report": {
            "A1": "LoB 1", "A2": 2,
            "C1": '=INDIRECT("\'"&A1&"\'!B"&A2)',
            "C2": '=INDIRECT("\'CoC "&A2&"\'!B3")',
        },
    })
    analysis = build_analysis(path, tmp_path / "w", "t")
    index = prep.reference_index(analysis)
    assert set(index["text_addressed"]) == {"LoB 1", "CoC 1", "CoC 2"}
    assert "INDIRECT" in prep.insert_blocker(index, "CoC 2") and prep.insert_blocker(index, "Free") is None
    ops, skipped = _titles(analysis)
    inserted = {o["sheet"] for o in ops if o["op"] == "insert_row"}
    assert "Free" in inserted and not inserted & {"LoB 1", "CoC 1", "CoC 2"}, (inserted, skipped)
    assert any(s.startswith("LoB 1!") and "address built as text" in s and "Report!C1" in s for s in skipped), skipped


def test_an_indirect_that_names_no_sheet_anyone_can_tell_blocks_every_sheet(tmp_path):
    block = {"A1": "Line", "B1": "2024", "A2": "Premium", "B2": 100}
    path = _wb(tmp_path / "m.xlsx", {
        "One": dict(block), "Two": dict(block),
        "Report": {"A1": 3, "C1": '=INDIRECT("\'"&CHAR(64+A1)&"\'!B2")', "E1": '=INDIRECT("B"&A1)'},
    })
    index = prep.reference_index(build_analysis(path, tmp_path / "w", "t"))
    assert set(index["text_addressed"]) == {"One", "Two", "Report"}  # C1: any sheet; E1: its own sheet
