"""1.8.0: FMT-002's cell-level action (theme colours -> explicit RGB) works by
blocks of same-coloured cells instead of cell by cell -- ~10 COM calls a cell
made one Apply take 37 minutes on a 77,000-cell model. The result must be the
same, cell for cell."""
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color

from app import prep
from app.excel_com import close_quietly, com_available, excel_session, open_for_write, save_in_place

needs_excel = pytest.mark.skipif(not com_available(), reason="needs a real Excel")


def _book(path: Path) -> list[str]:
    """A 12 x 6 block in theme colours, with a differently coloured column in
    the middle, one cell with no fill, one with an RGB colour already, and a
    stray cell far away. -> the cells of the operation."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    cells = []
    for r in range(2, 14):
        for c in range(2, 8):
            cell = ws.cell(r, c, r * c)
            theme = 4 if c != 5 else 9                       # column E: another theme colour
            cell.font = Font(color=Color(theme=1 if c != 5 else 0))
            if not (r == 7 and c == 3):                      # C7: no fill at all
                cell.fill = PatternFill("solid", fgColor=Color(theme=theme, tint=0.4 if r % 2 else 0.0))
            cells.append(cell.coordinate)
    ws["D9"].font = Font(color="FF112233")                   # explicit already
    ws["K30"] = "note"
    ws["K30"].font = Font(color=Color(theme=5))
    cells.append("K30")
    wb.save(path)
    wb.close()
    return cells


def _colours(path: Path) -> dict[str, tuple]:
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb["Calc"]
        out = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                font, fill = c.font.color, c.fill.fgColor if c.fill and c.fill.fill_type else None
                out[c.coordinate] = (
                    (font.type, font.rgb if font.type == "rgb" else font.theme) if font is not None else None,
                    (fill.type, fill.rgb if fill.type == "rgb" else fill.theme) if fill is not None else None,
                )
        return out
    finally:
        wb.close()


def _run(path: Path, cells: list[str], blockwise: bool) -> None:
    with excel_session() as excel:
        def go() -> None:
            wb = open_for_write(excel, path)
            try:
                ws = wb.Worksheets("Calc")
                prep._explicit_colors(ws, {"op": "explicit_colors", "sheet": "Calc", "cells": cells}, 1, 1, None, blockwise=blockwise)
                ws = None
                save_in_place(wb, path)
            finally:
                close_quietly(wb)

        go()


@needs_excel
def test_blocks_of_one_colour_give_the_same_result_as_cell_by_cell(tmp_path):
    a, b = tmp_path / "blocks.xlsx", tmp_path / "cells.xlsx"
    cells = _book(a)
    _book(b)
    before = _colours(a)
    assert before["B2"][0][0] == "theme" and before["B2"][1][0] == "theme"
    _run(a, cells, blockwise=True)
    _run(b, cells, blockwise=False)
    blocks, one_by_one = _colours(a), _colours(b)
    assert blocks == one_by_one
    assert all(v[1] is None or v[1][0] == "rgb" for k, v in blocks.items()), blocks  # every fill is explicit now
    assert blocks["E2"][0][0] == "rgb"                    # ... and the fonts Excel turns explicit (it keeps the theme's own text colour)
    assert blocks["C7"][1] is None                       # no fill stays no fill
    assert blocks["D9"][0] == ("rgb", "FF112233")        # an explicit colour inside a block is kept
    assert blocks["B2"] != blocks["E2"] and blocks["B2"][1] != blocks["B3"][1]  # the mixed block was split, not flattened


def test_a_mixed_range_is_split_and_a_single_cell_never_is():
    class Fake:
        def __init__(self, font, index, fill=None):
            self.Font = type("F", (), {"Color": font})()
            self.Interior = type("I", (), {"ColorIndex": index, "Color": fill})()

    assert prep._recolor(Fake(None, 5, 255), single=False) is False          # several font colours
    assert prep._recolor(Fake(0, None), single=False) is False               # several fills
    assert prep._recolor(Fake(0, 5, None), single=False) is False            # one palette index, several real colours
    block = Fake(255, 5, 65280)
    assert prep._recolor(block, single=False) is True and block.Font.Color == 255 and block.Interior.Color == 65280
    assert prep._recolor(Fake(None, None), single=True) is True              # rich text: left alone, as before
    assert prep._recolor(Fake(0, prep.XL_COLOR_INDEX_NONE), single=False) is True  # no fill: only the font
