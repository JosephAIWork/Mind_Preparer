"""Mind loop (1.10.0): the macro's run loop as MM_LOOP / MM_RESULT formulas.

The plans are built on small workbooks written here, with the macros the four client models
use (read through a patched `read_vba`, since an .xlsx carries none): a row-per-record loop
(UC-1), a counter driver with named bounds (UC-2), a list walked with Offset and a derived
year (UC-3). UC-1 itself is checked when its local copy is present."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from local_usecases import usecase_or_skip
from openpyxl.workbook.defined_name import DefinedName

from app import model_map
from app.inventory import build_analysis
from app.loop_builder import build_loop_plan, describe, loop_name_from
from app.model_map import build_model_map

ROW_LOOP_STYLE = '''
Sub Run_All()
first_row = Worksheets("Calculation").Cells(5, 20).Value
last_row = Worksheets("Calculation").Cells(6, 20).Value
Worksheets("OUTPUT").Range("A" & first_row & ":B100000").ClearContents
For ins = first_row To last_row
    Worksheets("Calculation").Range("A2:C2") = Worksheets("Policy Data").Range("A" & ins & ":C" & ins).Value
    Worksheets("OUTPUT").Range("A" & ins & ":B" & ins) = Worksheets("Calculation").Range("E10:F10").Value
Next
End Sub
'''

COUNTER_STYLE = '''
Sub RunPolicies()
Dim MacroStart As Long: MacroStart = Range("MacroStart").Value
Dim MacroEnd As Long: MacroEnd = Range("MacroEnd").Value
For polNo = MacroStart To MacroEnd
    Range("POLICY_ROW").Value = polNo
    Application.Calculate
    Range("ByGroupPaster").Offset(0, 10 * (Temp2 - 1)).Value2 = Range("AggregateOutput").Value2
Next polNo
End Sub
'''

COUNTED_STYLE = '''
Sub RunYears()
For i = 1 To 5
    Range("active_year").Value2 = Range("proj_year").Value2 + i - 1
    Range("out_block").Value2 = Range("calc_block").Value2
Next i
End Sub
'''

LISTS_STYLE = '''
Sub combined()
For i_s = 1 To 3
    Range("active_scenario").Value2 = Range("scenario_start").Offset(0, i_s - 1).Value2
    For i_c = 1 To 2
        Range("active_cluster").Value2 = Range("Cluster1").Offset(i_c - 1, 0).Value2
        Worksheets("Out").Range("A" & i_c & ":B" & i_c) = Worksheets("Calc").Range("E10:F10").Value
    Next i_c
Next i_s
End Sub
'''


def _row_loop_book(tmp_path: Path, last_as_formula: bool = True) -> Path:
    wb = openpyxl.Workbook()
    base = wb.active
    base.title = "Policy Data"
    base.append(["Policy", "Loan", "Person"])
    for i in range(1, 4):
        base.append([100 + i, 10 * i, 5 + i])
    calc = wb.create_sheet("Calculation")
    calc["A2"], calc["B2"], calc["C2"] = 103, 30, 8
    calc["S5"], calc["T5"] = "First row", 1
    calc["S6"], calc["T6"] = "Last row", "=N5" if last_as_formula else 4
    calc["N5"] = "=COUNTA('Policy Data'!A:A)"
    calc["E10"], calc["F10"] = "=A2*2", "=B2+C2"
    out = wb.create_sheet("OUTPUT")
    for r in range(1, 5):
        out.append([r, r])
    out["A5"] = "left from an earlier run"
    p = tmp_path / "row_loop.xlsx"
    wb.save(p)
    wb.close()
    return p


def _plan_for(path: Path, tmp_path: Path, vba: str, monkeypatch, **kw) -> dict:
    monkeypatch.setattr(model_map, "read_vba", lambda _p: [("Module1", vba)])
    analysis = build_analysis(path, tmp_path / "w", path.stem)
    mp = build_model_map(analysis)
    assert mp["loop"]["found"], mp["loop"]
    return build_loop_plan(analysis, mp, **kw)


def test_loop_names_are_plain_words():
    assert loop_name_from("Policy Data") == "Policy_Data" and loop_name_from("Tariffs to run!") == "Tariffs_to_run" and loop_name_from("") == "Loop"


def test_a_row_per_record_loop_becomes_mm_loop_per_column_and_mm_result_per_row(tmp_path, monkeypatch):
    # the last bound is a formula with no stored value in a file openpyxl wrote: the plan asks
    plan = _plan_for(_row_loop_book(tmp_path), tmp_path, ROW_LOOP_STYLE, monkeypatch)
    lp = plan["loops"][0]
    assert lp["status"] == "needs_decision" and "first or last value" in lp["why"] and plan["operations"] == []
    # with the last value given, everything is placed
    plan = _plan_for(_row_loop_book(tmp_path), tmp_path, ROW_LOOP_STYLE, monkeypatch, last=4)
    lp = plan["loops"][0]
    assert lp["status"] == "ready" and lp["name"] == "Policy_Data" and lp["kind"] == "row_of_table" and (lp["first"], lp["last"], lp["count"]) == (2, 4, 3) and lp["header_skipped"] == 1
    assert lp["first_from"] == {"kind": "cell", "cell": "Calculation!T5", "name": None, "value": 1, "formula": None} and lp["last_from"]["kind"] == "chosen"
    assert [c["formula"] for c in lp["cells"]] == [f"=MM_LOOP(\"Policy_Data\",'Policy Data'!{col}2:INDEX('Policy Data'!{col}:{col},MM_LASTROW('Policy Data'!A2)))" for col in "ABC"]  # follows the table in Mind
    assert [c["range_fixed"] for c in lp["cells"]] == [f"'Policy Data'!{col}2:{col}4" for col in "ABC"] and lp["dynamic"] is True
    assert any("ranges follow the table" in n for n in lp["notes"])
    assert [c["cell"] for c in lp["cells"]] == ["Calculation!A2", "Calculation!B2", "Calculation!C2"]
    assert any("the table's header: the loop starts at 2" in n for n in lp["notes"])  # row 1 is the header: skipped by default, said
    assert any("Excel behaviour may differ from Mind behaviour" in n for n in lp["notes"]) and any("Input Manager only reach a grid flagged /Input" in n for n in lp["notes"])
    out = plan["outputs"][0]
    assert out["status"] == "ready" and out["items"] == 3 and out["cells"] == 6 and out["loop"] == "Policy_Data" and "written for the 3 loop values present now" in out["rows_note"]
    assert out["example"] == 'OUTPUT!A2 =MM_RESULT(Calculation!$E$10,"Policy_Data",1)'
    ops = {(o["sheet"], o["cell"]): o for o in plan["operations"]}
    assert ops[("OUTPUT", "B4")]["after"] == '=MM_RESULT(Calculation!$F$10,"Policy_Data",3)' and ops[("OUTPUT", "B4")]["before"] == 4
    assert ops[("Calculation", "A2")]["before"] == 103 and ops[("Calculation", "A2")]["op"] == "set_formula"
    assert all(o["action_id"] == "mind_loop" and o["rule_id"] == "LOOP-001" for o in plan["operations"])
    # the cleared area is listed with what an earlier run left there, but not emptied unless asked
    cl = plan["clears"][0]
    assert cl["resolved"] == "OUTPUT!A2:B100000" and cl["cells"] == 1 and cl["included"] is False and plan["counts"]["cleared"] == 0
    assert len(plan["operations"]) == 3 + 6 + 2 and plan["counts"] == {"loop_cells": 3, "result_cells": 6, "cleared": 0, "extra": 2}  # + the 2 header captions
    assert "Policy_Data is a loop of 3 values" in plan["how_it_runs"] and "Run_All is not needed" in plan["how_it_runs"]
    json.dumps(plan, default=str)
    text = describe(plan)
    assert "loop Policy_Data  driver Calculation!A2:C2  2 -> 4  [ready]" in text and "operations: 11" in text


def test_bounds_name_and_clearing_can_be_chosen(tmp_path, monkeypatch):
    plan = _plan_for(_row_loop_book(tmp_path, last_as_formula=False), tmp_path, ROW_LOOP_STYLE, monkeypatch, name="Policy", first=2, clear_old=True)
    lp = plan["loops"][0]
    assert lp["status"] == "ready" and lp["name"] == "Policy" and (lp["first"], lp["last"], lp["count"]) == (2, 4, 3) and lp["last_from"]["kind"] == "cell"
    assert lp["cells"][0]["formula"] == "=MM_LOOP(\"Policy\",'Policy Data'!A2:INDEX('Policy Data'!A:A,MM_LASTROW('Policy Data'!A2)))" and not any("header" in n for n in lp["notes"])
    ops = {(o["sheet"], o["cell"]): o for o in plan["operations"]}
    assert ops[("OUTPUT", "A2")]["after"] == '=MM_RESULT(Calculation!$E$10,"Policy",1)'
    assert ops[("OUTPUT", "A1")]["op"] == "set_value"  # the header row above the block, completed for Mind (no MM_RESULT there)
    # the macro clears from the first value's row on (A2:B100000): row 5 is emptied, row 1 is
    # outside that area and the written rows are left
    cleared = sorted(c for (sh, c), o in ops.items() if o["op"] == "clear_cell")
    assert plan["clears"][0]["resolved"] == "OUTPUT!A2:B100000" and cleared == ["A5"] and plan["clears"][0]["included"] is True and plan["counts"]["cleared"] == 1


def test_a_counter_driver_with_named_bounds(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    m = wb.active
    m.title = "Macro"
    m["A1"], m["B1"], m["A2"], m["B2"] = "Macro Start", 1, "Macro End", 19
    pr = wb.create_sheet("Projection")
    pr["B3"], pr["C3"] = "Policy run", 7
    for n, ref in (("MacroStart", "Macro!$B$1"), ("MacroEnd", "Macro!$B$2"), ("POLICY_ROW", "Projection!$C$3"), ("ByGroupPaster", "Macro!$V$13:$AE$140"), ("AggregateOutput", "Macro!$A$11:$J$138")):
        wb.defined_names[n] = DefinedName(n, attr_text=ref)
    p = tmp_path / "counter.xlsx"
    wb.save(p)
    wb.close()
    plan = _plan_for(p, tmp_path, COUNTER_STYLE, monkeypatch)
    lp = plan["loops"][0]
    assert lp["status"] == "ready" and lp["kind"] == "counter" and lp["name"] == "POLICY_ROW" and lp["cells"][0]["formula"] == '=MM_LOOP("POLICY_ROW",Macro!B2)'
    assert plan["outputs"][0]["status"] == "not_built" and "MM_RESULT(cell,\"POLICY_ROW\",k)" in plan["outputs"][0]["why"]
    # starting elsewhere than 1: the documented IndexStart / IndexEnd form, flagged for a check in Mind
    plan = _plan_for(p, tmp_path, COUNTER_STYLE, monkeypatch, first=3)
    lp = plan["loops"][0]
    assert lp["cells"][0]["formula"] == '=MM_LOOP("POLICY_ROW",Macro!B2,3,3,Macro!B2)' and any("IndexStart" in n for n in lp["notes"])
    assert lp["first_from"] == {"kind": "chosen", "value": 3}


def test_a_counted_loop_without_a_count_cell_asks_for_one(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "Assumption Setting"
    a["C8"], a["D8"] = "Projection year", 2025
    mi = wb.create_sheet("Model Inputs")
    mi["B173"], mi["C173"] = "Active Year", 2025
    for n, ref in (("proj_year", "'Assumption Setting'!$D$8"), ("active_year", "'Model Inputs'!$C$173"), ("out_block", "'Model Inputs'!$F$1:$G$2"), ("calc_block", "'Model Inputs'!$H$1:$I$2")):
        wb.defined_names[n] = DefinedName(n, attr_text=ref)
    p = tmp_path / "counted.xlsx"
    wb.save(p)
    wb.close()
    plan = _plan_for(p, tmp_path, COUNTED_STYLE, monkeypatch)
    lp = plan["loops"][0]
    assert lp["status"] == "needs_decision" and "count" in lp["why"] and plan["operations"] == []
    plan = _plan_for(p, tmp_path, COUNTED_STYLE, monkeypatch, count_cell="Model Inputs!D173")
    lp = plan["loops"][0]
    assert lp["status"] == "ready" and lp["kind"] == "derived"
    assert lp["cells"][0]["formula"] == "='Assumption Setting'!D8+MM_LOOP(\"active_year\",'Model Inputs'!D173)-1"
    extra = [o for o in plan["operations"] if o["op"] == "set_value"]
    assert extra == [{"action_id": "mind_loop", "rule_id": "LOOP-001", "op": "set_value", "sheet": "Model Inputs", "cell": "D173", "before": None, "after": 5, "note": "the count of the loop active_year (chosen cell)"}]


def test_nested_list_loops_and_a_block_placed_by_one_of_them(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    mi = wb.active
    mi.title = "Model Inputs"
    mi["B8"], mi["C8"] = "Scenario-->", "Base"
    mi["E8"], mi["F8"], mi["G8"] = "Base", "Up", "Down"
    mi["B9"], mi["C9"] = "Cluster", "North"
    vl = wb.create_sheet("Validation list")
    vl["B7"], vl["B8"] = "North", "South"
    calc = wb.create_sheet("Calc")
    calc["E10"], calc["F10"] = "='Model Inputs'!C8", "='Model Inputs'!C9"
    wb.create_sheet("Out")
    for n, ref in (("active_scenario", "'Model Inputs'!$C$8"), ("scenario_start", "'Model Inputs'!$E$8"), ("active_cluster", "'Model Inputs'!$C$9"), ("Cluster1", "'Validation list'!$B$7")):
        wb.defined_names[n] = DefinedName(n, attr_text=ref)
    p = tmp_path / "lists.xlsx"
    wb.save(p)
    wb.close()
    plan = _plan_for(p, tmp_path, LISTS_STYLE, monkeypatch)
    loops = {l["driver"]: l for l in plan["loops"]}
    sc, cl = loops["'Model Inputs'!C8"], loops["'Model Inputs'!C9"]
    # the two lists run to their last item: the ranges follow the lists (MM_ROW along a row, MM_COLUMN down a column)
    assert sc["status"] == "ready" and sc["kind"] == "list" and sc["cells"][0]["formula"] == "=MM_LOOP(\"active_scenario\",MM_ROW('Model Inputs'!E8))" and sc["cells"][0]["range_fixed"] == "'Model Inputs'!E8:G8"
    assert cl["status"] == "ready" and cl["kind"] == "list" and cl["cells"][0]["formula"] == "=MM_LOOP(\"active_cluster\",MM_COLUMN('Validation list'!B7))" and cl["cells"][0]["range_fixed"] == "'Validation list'!B7:B8"
    out = plan["outputs"][0]
    assert out["status"] == "not_built" and "2 nested loops" in out["why"] and '"active_scenario",k1,"active_cluster",k2' in out["why"]
    assert plan["counts"]["loop_cells"] == 2 and plan["counts"]["result_cells"] == 0


def test_without_a_loop_there_is_nothing_to_build(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    p = tmp_path / "plain.xlsx"
    wb.save(p)
    wb.close()
    analysis = build_analysis(p, tmp_path / "w", "plain")
    plan = build_loop_plan(analysis, build_model_map(analysis))
    assert plan["found"] is False and plan["operations"] == [] and "No run loop" in plan["notes"][0]


def test_uc1_plan(tmp_path):
    """The real row-per-record model (UC-1), when its copy is here: the macro starts on the header row, the loop on
    row 2; 39 MM_LOOP cells (the table's columns filled in rows 2..6), 50 MM_RESULT cells, 6 header captions above
    the block. Its sheet and name values come from the local expectations (tests/local_usecases.py)."""
    path, n = usecase_or_skip("UC-1")
    analysis = build_analysis(path, tmp_path / "w", "UC-1")
    plan = build_loop_plan(analysis, build_model_map(analysis))
    lp = plan["loops"][0]
    assert lp["status"] == "ready" and lp["name"] == n["loop_name"] and (lp["first"], lp["last"]) == (2, 6) and lp["header_skipped"] == 1 and lp["last_from"]["formula"] == n["last_formula"]
    # 39 columns: three columns of the table hold nothing in rows 2..6, so no loop is written for them
    assert lp["cells"][0]["formula"] == n["first_loop_formula"] and lp["cells"][0]["range_fixed"] == n["first_loop_range"] and len(lp["cells"]) == 39
    out = plan["outputs"][0]
    assert out["status"] == "ready" and out["items"] == 5 and out["cells"] == 50 and out["example"] == n["result_example"]
    assert len(plan["operations"]) == 95 and plan["counts"]["extra"] == 6  # 39 loop cells + 50 result cells + 6 header captions


def test_the_header_row_above_an_output_block_is_completed_for_mind(tmp_path, monkeypatch):
    """Seen in real Mind: a grid's first row is its header, an empty cell there splits the block and the
    first pass shows as a header. With first = 2 the block starts on row 2 and row 1 holds the macro's
    header-row pass (numbers here): every such cell gets the source's caption, else the source address."""
    plan = _plan_for(_row_loop_book(tmp_path), tmp_path, ROW_LOOP_STYLE, monkeypatch, first=2, last=4)
    out = plan["outputs"][0]
    assert out["status"] == "ready" and out["items"] == 3 and out["header_cells"] == ["OUTPUT!A1", "OUTPUT!B1"]
    heads = {o["cell"]: o for o in plan["operations"] if o["op"] == "set_value" and o["sheet"] == "OUTPUT"}
    assert set(heads) == {"A1", "B1"} and heads["A1"]["before"] == 1 and heads["A1"]["after"] and heads["A1"]["rule_id"] == "LOOP-001"
    assert plan["counts"]["extra"] == 2 and any("header row 1 above the block is completed" in n for n in plan["notes"])
    # a block that starts on row 1 has no header row: said, nothing written
    plan = _plan_for(_row_loop_book(tmp_path), tmp_path, ROW_LOOP_STYLE, monkeypatch, first=1, last=4)
    assert "header_cells" not in plan["outputs"][0] and plan["counts"]["extra"] == 0 and any("no header row above it" in n for n in plan["notes"])


def test_fixed_ranges_can_be_asked_for(tmp_path, monkeypatch):
    """`dynamic=False` (the API's fixed_ranges, the CLI's --fixed-ranges) writes the old X2:X6 ranges."""
    plan = _plan_for(_row_loop_book(tmp_path), tmp_path, ROW_LOOP_STYLE, monkeypatch, last=4, dynamic=False)
    lp = plan["loops"][0]
    assert lp["dynamic"] is False and [c["formula"] for c in lp["cells"]] == [f"=MM_LOOP(\"Policy_Data\",'Policy Data'!{col}2:{col}4)" for col in "ABC"]
    assert not any("follow the table" in n for n in lp["notes"])


def test_a_stale_loop_in_a_skipped_driver_cell_is_cleared(tmp_path, monkeypatch):
    """A column of the table that holds nothing gets no loop; a loop written there by an earlier build
    (another length once the table grows) is cleared."""
    import openpyxl

    p = _row_loop_book(tmp_path)
    wb = openpyxl.load_workbook(p)
    for r in range(2, 5):
        wb["Policy Data"].cell(r, 3).value = None  # column C of the table: empty in the loop rows
    wb["Calculation"]["C2"] = "=MM_LOOP(\"Policy_Data\",'Policy Data'!C2:C4)"  # left by an earlier build
    wb.save(p)
    wb.close()
    plan = _plan_for(p, tmp_path, ROW_LOOP_STYLE, monkeypatch, first=2, last=4)
    lp = plan["loops"][0]
    assert [c["ref"] for c in lp["cells"]] == ["A2", "B2"] and any("hold nothing" in n for n in lp["notes"])
    clears = [o for o in plan["operations"] if o["op"] == "clear_cell"]
    assert len(clears) == 1 and clears[0]["sheet"] == "Calculation" and clears[0]["cell"] == "C2" and "MM_LOOP" in clears[0]["before"]


def test_the_header_pass_is_kept_when_first_is_chosen_and_a_cut_short_loop_keeps_fixed_ranges(tmp_path, monkeypatch):
    plan = _plan_for(_row_loop_book(tmp_path), tmp_path, ROW_LOOP_STYLE, monkeypatch, first=1, last=4)
    lp = plan["loops"][0]
    assert lp["first"] == 1 and "header_skipped" not in lp and any("its header row" in n for n in lp["notes"])
    # the table runs to row 4: stopping at 3 keeps the rows of today only
    plan = _plan_for(_row_loop_book(tmp_path), tmp_path, ROW_LOOP_STYLE, monkeypatch, first=2, last=3)
    lp = plan["loops"][0]
    assert lp["dynamic"] is False and lp["cells"][0]["formula"] == "=MM_LOOP(\"Policy_Data\",'Policy Data'!A2:A3)"
    assert any("before the end of the table (row 4)" in n for n in lp["notes"]) and "rows_note" not in plan["outputs"][0]


def test_a_dropdown_on_a_loop_driver_is_found_and_removed(tmp_path):
    """1.11.3: Mind drops the formula of a cell with a dropdown, so the loop builder removes the dropdown
    (UC-3's drivers, and a probe workbook in Mind on 2026-10-16)."""
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    from app import loop_builder, prep

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model Inputs"
    for i, t in enumerate(["a", "b", "c"], start=2):
        ws.cell(i, 5, t)
    wb.defined_names["opts"] = DefinedName("opts", attr_text="'Model Inputs'!$E$2:$E$4")
    dv = DataValidation(type="list", formula1="opts")
    ws.add_data_validation(dv)
    dv.add("B2:D4")
    path = tmp_path / "dv.xlsx"
    wb.save(path)
    analysis = {"workbooks": [{"copy_path": str(path)}]}
    assert loop_builder._dropdown_at(analysis, "Model Inputs", "C3") == "a dropdown list (opts)"
    assert loop_builder._dropdown_at(analysis, "Model Inputs", "A1") is None
    wb = openpyxl.load_workbook(path)
    prep._remove_validation_openpyxl(wb["Model Inputs"], "C3")
    left = str(wb["Model Inputs"].data_validations.dataValidation[0].sqref).split()
    cells = {(c, r) for part in left for c in range(openpyxl.utils.cell.range_boundaries(part)[0], openpyxl.utils.cell.range_boundaries(part)[2] + 1)
             for r in range(openpyxl.utils.cell.range_boundaries(part)[1], openpyxl.utils.cell.range_boundaries(part)[3] + 1)}
    assert cells == {(c, r) for c in (2, 3, 4) for r in (2, 3, 4)} - {(3, 3)}


def test_a_dropdown_kept_in_the_extension_block_is_found(tmp_path):
    """A list taken from another sheet is kept by Excel in the sheet's extension block, which openpyxl does not read."""
    import zipfile

    from app import loop_builder

    path = tmp_path / "x14.xlsx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("xl/workbook.xml", '<workbook><sheets><sheet name="Drivers" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships><Relationship Id="rId1" Type="ws" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", '<worksheet><extLst><ext><x14:dataValidations count="1"><x14:dataValidation type="list" allowBlank="1">'
                   '<x14:formula1><xm:f>Lists!$B$2:$B$6</xm:f></x14:formula1><xm:sqref>E3 K3</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst></worksheet>')
    assert loop_builder._x14_dropdowns(str(path), "Drivers") == [("E3 K3", "Lists!$B$2:$B$6")]
    analysis = {"workbooks": [{"copy_path": str(path)}]}
    assert loop_builder._dropdown_at(analysis, "Drivers", "K3") == "a dropdown list (Lists!$B$2:$B$6)"
    assert loop_builder._dropdown_at(analysis, "Drivers", "H3") is None


def test_an_output_that_reads_the_cleared_area_gets_it_cleared():
    # UC-2 (1.11.11): Macro!A13 = Projection!CX15 + OFFSET(V13, 0, 10 * tariff): the block the macro
    # copies adds the per-tariff results it pasted earlier, which the macro clears before its loop
    from app import loop_builder

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Macro"
    ws["A13"] = "=Projection!CX15+OFFSET(V13,0,10*(MATCH(1,X1:X3,0)-1))"
    ws["B13"] = "=Projection!CY15"
    area = {"sheet": "Macro", "r1": 13, "c1": 22, "r2": 140, "c2": 281}
    assert loop_builder._reader_of_area(wb, "Macro!A13:B13", area) == "Macro!A13"
    assert loop_builder._reader_of_area(wb, "Macro!B13:B13", area) is None  # reads another sheet only
    assert loop_builder._reader_of_area(wb, "Macro!A13:B13", {**area, "sheet": "Other"}) is None
