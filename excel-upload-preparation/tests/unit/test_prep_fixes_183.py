"""1.8.3: three findings that stopped Mind on UC-2 and UC-4 now have a Prep action.

  FRM-006  =SUM('Line 1:>>'!E5)  -> =SUM('Line 1'!E5,'Line 2'!E5)   (UC-4: "Sheet '>> Reporting:>>>' not found")
  RSK-002  headers over every column an array formula covers      (UC-2: "exceeds the grid size")
  REF-001  a name that does not exist -> NA()                      (UC-2: "not a function : Step_increase_rates_array")
"""
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from app.config import load_config
from app.excel_com import com_available
from app.grids import all_grids
from app.inventory import build_analysis
from app.numbers_check import Moves, compare_workbooks
from app.prep import apply_operations, plan_actions
from app.rules_engine import RulesEngine
from app.validators import run_rule

needs_excel = pytest.mark.skipif(not com_available(), reason="needs a real Excel")


def _finding(analysis, rule_id):
    return run_rule(RulesEngine().get(rule_id), analysis, load_config())


def _plan(analysis):
    return {a["id"]: a for a in plan_actions(analysis, {"findings": []})}


def _three_d_model(path: Path, divider_holds_a_value: bool = False) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, scale in (("Line 1", 1), ("Line 2", 10)):
        ws = wb.create_sheet(name)
        ws["A1"], ws["B1"] = "Line", "2024"
        ws["A2"], ws["B2"] = "Premium", 100 * scale
        ws["A3"], ws["B3"] = "Claims", 60 * scale
    end = wb.create_sheet(">>")
    end["A1"] = "end of the lines of business"  # one text cell alone: no grid, Mind imports nothing
    if divider_holds_a_value:
        end["B2"] = "n/a"   # text alone: still no grid for Mind, but the 3-D reference does read it
    tot = wb.create_sheet("Total")
    tot["A1"], tot["B1"] = "Line", "2024"
    tot["A2"], tot["B2"] = "Premium", "=SUM('Line 1:>>'!B2)"
    tot["A3"], tot["B3"] = "Claims", "=SUM('Line 1:>>'!B3)"
    tot["A4"], tot["B4"] = "Lines", "=COUNT('Line 1:>>'!B2)"
    wb.save(path)
    wb.close()
    return path


@needs_excel
def test_3d_references_are_written_out_sheet_by_sheet_and_the_numbers_stay(tmp_path):
    src = _three_d_model(tmp_path / "m.xlsx")
    analysis = build_analysis(src, tmp_path / "w", "m")
    a = _plan(analysis)["expand_3d_references"]
    assert a["level"] == "blocking" and a["count"] == 3 and not a["skipped"]
    by = {o["cell"]: o for o in a["operations"]}
    assert by["B2"]["after"] == "=SUM('Line 1'!B2,'Line 2'!B2)" and by["B4"]["after"] == "=COUNT('Line 1'!B2,'Line 2'!B2)"
    assert ">>" in by["B2"]["note"]
    out = apply_operations(src, tmp_path / "prep", a["operations"])
    assert out["status"] == "APPLIED" and not out["failed"]
    res = compare_workbooks(src, Path(out["output_path"]), Moves())
    assert res["clean"] is True and res["counts"]["good_to_other"] == 0
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "a")
    assert _finding(after, "FRM-006")["status"] == "PASS"


def test_a_divider_sheet_that_holds_something_at_the_cell_is_not_dropped_silently(tmp_path):
    src = _three_d_model(tmp_path / "m.xlsx", divider_holds_a_value=True)
    analysis = build_analysis(src, tmp_path / "w", "m")
    a = _plan(analysis)["expand_3d_references"]
    # B2 and B4 read B2, where '>>' holds something: left for a person; B3 reads B3, where '>>' holds nothing
    assert [o["cell"] for o in a["operations"]] == ["B3"]
    assert len(a["skipped"]) == 2 and all(">>" in s and "by hand" in s for s in a["skipped"])


def _array_model(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Temp"
    ws["B4"] = "Per Policy Projection"
    ws.merge_cells("B4:E4")                                  # a merged header: covers B..E, nothing over F..L
    for i in range(10):
        ws[f"A{20 + i}"] = (i + 1) * 1.5
    ws["B5"], ws["B6"] = "ME charge", "Admin cost"
    ws["C5"] = ArrayFormula("C5:L5", "=TRANSPOSE(A20:A29)")  # ten columns, C..L
    ws["C6"] = ArrayFormula("C6:L6", "=TRANSPOSE(A20:A29)*2")
    wb.save(path)
    wb.close()
    return path


@needs_excel
def test_the_header_row_is_extended_over_the_array_and_the_grid_becomes_one(tmp_path):
    src = _array_model(tmp_path / "arr.xlsx")
    analysis = build_analysis(src, tmp_path / "w", "arr")
    assert _finding(analysis, "RSK-002")["status"] == "ERROR"
    plan = _plan(analysis)
    a = plan["cover_array_headers"]
    assert a["level"] == "blocking" and [o["cell"] for o in a["operations"]] == ["F4", "G4", "H4", "I4", "J4", "K4", "L4"]
    assert a["operations"][0]["after"] == "Period 4" and a["operations"][-1]["after"] == "Period 10"
    # the titles the grid detection would put on the array's loose cells lose to the headers: one write per cell
    assert not any(o["cell"] in ("F4", "G4") for o in plan["create_grid_titles"]["operations"])
    out = apply_operations(src, tmp_path / "prep", a["operations"])
    assert out["status"] == "APPLIED" and not out["failed"]
    res = compare_workbooks(src, Path(out["output_path"]), Moves())
    assert res["clean"] is True
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "a")
    assert _finding(after, "RSK-002")["status"] == "PASS"
    assert [g["ref"] for g in all_grids(after) if g["sheet"] == "Temp" and g["first_row"] == 4] == ["B4:L6"]


@needs_excel
def test_a_name_that_does_not_exist_becomes_na_where_it_stands(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["A1"] = "Y"
    ws["B1"], ws["C1"] = 1, 2
    ws["A2"] = '=IF(A1="S",HLOOKUP(2,Nope_rates_array,1,FALSE),B1+C1)'   # the branch Excel never takes: no #NAME? shows
    ws["A3"] = "=A2*2"
    p = tmp_path / "name.xlsx"
    wb.save(p)
    wb.close()
    analysis = build_analysis(p, tmp_path / "w", "n")
    assert _finding(analysis, "REF-001")["status"] == "ERROR"
    plan = _plan(analysis)
    a = plan["fix_broken_refs"]                                       # Excel takes HLOOKUP(..., NA(), ...): the name alone becomes NA()
    assert a["count"] == 1 and a["operations"][0]["after"] == '=IF(A1="S",HLOOKUP(2,NA(),1,FALSE),B1+C1)'
    out = apply_operations(p, tmp_path / "prep", a["operations"])
    assert out["status"] == "APPLIED"
    res = compare_workbooks(p, Path(out["output_path"]), Moves())
    assert res["clean"] is True
    assert _finding(build_analysis(Path(out["output_path"]), tmp_path / "after", "a"), "REF-001")["status"] == "PASS"


@needs_excel
def test_indirect_used_as_a_value_becomes_the_direct_reference(tmp_path):
    """1.8.4 -- FRM-007 (UC-4, 4,410 formulas: 'Unable to cast AMReference to IConvertible')."""
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    inp = wb.active
    inp.title = "Inputs"
    lob = wb.create_sheet("Line 1")
    lob["Z12"], lob["Z16"] = 10, 4
    inp["G710"], inp["G711"] = "Line 1", "Line 2"                 # Line 2 is not a sheet
    inp["AA710"] = "=IF(INDIRECT(\"'\"&$G710&\"'!$Z$16\")=0,0,66%*INDIRECT(\"'\"&$G710&\"'!$Z$12\")/INDIRECT(\"'\"&$G710&\"'!$Z$16\"))"
    inp["AA711"] = "=IFERROR(IF(INDIRECT(\"'\"&$G711&\"'!$Z$16\")=0,0,1),-1)"   # the missing sheet: an error today, NA() tomorrow, -1 both ways
    inp["B2"] = "=NOT(INDIRECT(\"Line\"&COLUMN()-1&\"_Boolean\"))*2"
    inp["B3"] = "=INDIRECT(\"'\"&$B$4&\"'!A1\")+1"                   # B4 is a formula with no stored value: by hand
    inp["B4"] = "=G710"
    inp["B5"] = True
    wb.defined_names["Line1_Boolean"] = DefinedName("Line1_Boolean", attr_text="Inputs!$B$5")
    p = tmp_path / "indirect.xlsx"
    wb.save(p)
    wb.close()
    analysis = build_analysis(p, tmp_path / "w", "i")
    a = _plan(analysis)["direct_references"]
    by = {o["cell"]: o for o in a["operations"]}
    assert a["level"] == "blocking" and set(by) == {"AA710", "AA711", "B2"}
    assert by["AA710"]["after"] == "=IF('Line 1'!$Z$16=0,0,66%*'Line 1'!$Z$12/'Line 1'!$Z$16)"
    assert by["AA711"]["after"] == "=IFERROR(IF(NA()=0,0,1),-1)" and "Line 2" in by["AA711"]["note"]
    assert by["B2"]["after"] == "=NOT(Line1_Boolean)*2"
    assert len(a["skipped"]) == 1 and "B3" in a["skipped"][0] and "cannot be worked out" in a["skipped"][0]
    out = apply_operations(p, tmp_path / "prep", a["operations"])
    assert out["status"] == "APPLIED" and not out["failed"]
    res = compare_workbooks(p, Path(out["output_path"]), Moves())
    assert res["clean"] is True and res["counts"]["good_to_other"] == 0
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "a")
    f = _finding(after, "FRM-007")
    assert f["status"] == "ERROR" and {s["cell"] for s in f["observed"]["sites"]} == {"B3"}     # only the one left for a person


def test_a_lone_text_cell_that_formulas_read_gets_a_title_of_its_own(tmp_path):
    """1.8.8 -- STR-008 (UC-4, 'Line 1'!G2: Mind imports no loose text cell; "Premium_Retained#"&$G$2 came back empty)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Line 1"
    ws["E2"], ws["G2"] = "Sheet name :", "Line 1"          # a label, a gap, the value: G2 stands alone
    ws["C6"] = 42                                           # a number standing alone: Mind imports it
    ws["J2"] = "unused"                                     # a lone text no formula reads
    ws["B10"] = "Note"                                      # a lone text on row... read, and nothing above? B9 is empty
    ws["E5"], ws["F5"] = "Line", "2024"
    ws["E6"], ws["F6"] = "Premium", 100
    ws["A20"] = '="Premium_Retained#"&$G$2'
    ws["A21"] = "=C6*2"
    ws["A22"] = "=B10&C6"
    p = tmp_path / "lone.xlsx"
    wb.save(p)
    wb.close()
    analysis = build_analysis(p, tmp_path / "w", "l")
    f = _finding(analysis, "STR-008")
    assert f["status"] == "ERROR" and {s["cell"] for s in f["observed"]["sites"]} == {"G2", "B10"}
    a = _plan(analysis)["title_lone_text_cells"]
    assert a["level"] == "blocking" and not a["skipped"]
    by = {o["cell"]: o for o in a["operations"]}
    assert by["G1"]["after"] == "#Sheet name value" and by["B9"]["after"] == "#Line 1 B10"
    out = apply_operations(p, tmp_path / "prep", a["operations"], prefer_excel=False)
    assert out["status"] == "APPLIED"
    after = build_analysis(Path(out["output_path"]), tmp_path / "after", "a")
    assert _finding(after, "STR-008")["status"] == "PASS"
    assert any(g["sheet"] == "Line 1" and g["ref"] == "G2" and g["name"] == "Sheet name value" for g in all_grids(after))
