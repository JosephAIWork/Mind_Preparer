"""1.11.6: the "Do everything" time estimate (app/estimate.py) and its endpoint."""
from fastapi.testclient import TestClient

from app.estimate import CAVEAT, OPEN_S_PER_MB, RECALC_AHEAD_S, SAVE_S_PER_MB, estimate_chain
from app.web import server


def test_uc5_numbers_give_the_range_it_took():
    # UC-5 (1.11.10, two sheets skipped so the stored errors are unknown): 170 MB, scan 270 s, 852
    # Prep changes, 8 pivot tables to values -- Fix all issues took 661 s, the whole Do everything 1,189 s
    r = estimate_chain(170, 270, 305_785, None, 852, True, worker=True, baseline_mb=170, pivots=[("pivot_to_values", 0)] * 8)
    fix = next(s for s in r["steps"] if s["id"] == "fixAll")
    assert fix["low_s"] <= 661 <= fix["high_s"]
    assert r["total"]["low_s"] <= 1189 <= r["total"]["high_s"]
    assert [s["id"] for s in r["steps"]] == ["fixAll", "prepare", "loop", "recalc"]
    assert r["total"]["low_s"] <= r["total"]["mid_s"] <= r["total"]["high_s"]
    assert r["caveat"] == CAVEAT and "depends" in CAVEAT


def test_a_clean_workbook_and_no_loop_cost_less():
    clean = estimate_chain(27, 300, 57_000, 0, 0, False)
    broken = estimate_chain(27, 300, 57_000, 500, 200, True)
    assert clean["total"]["mid_s"] < broken["total"]["mid_s"]
    assert next(s for s in clean["steps"] if s["id"] == "loop")["mid_s"] == 0


def test_an_unscanned_file_falls_back_to_its_formula_count():
    r = estimate_chain(10, None, 100_000, None)
    assert r["basis"]["scan_measured"] is False and r["basis"]["scan_s"] == 350


def test_the_endpoint_answers_after_a_scan(flagged_model_broken_xlsx):
    client = TestClient(server.app)
    with flagged_model_broken_xlsx.open("rb") as f:
        sid = client.post("/api/sessions", files={"file": (flagged_model_broken_xlsx.name, f, "application/octet-stream")}, data={"mode": "plan"}).json()["sessionId"]
    body = client.get(f"/api/sessions/{sid}/estimate?loop=false").json()
    assert body["basis"]["scan_measured"] and body["basis"]["stored_errors"] is not None
    assert next(s for s in body["steps"] if s["id"] == "loop")["mid_s"] == 0
    assert client.get("/api/sessions/nope/estimate").status_code == 404


def test_pivot_tables_to_formulas_are_priced_by_their_cells():
    plain = estimate_chain(179, 344, 305_785, 10, 8, True)
    pivots = estimate_chain(179, 344, 305_785, 10, 8, True, pivots=[("pivot_to_formulas", 2_000)] * 8)
    prep = lambda r: next(s for s in r["steps"] if s["id"] == "prepare")  # noqa: E731
    # 8 pivots of 2,000 value cells: 8 * (3 s + 20 s) on top of the round, not 8 * 0.02 s
    assert prep(pivots)["low_s"] - prep(plain)["low_s"] >= 8 * 20
    values = estimate_chain(179, 344, 305_785, 10, 8, True, pivots=[("pivot_to_values", 2_000)] * 8)
    assert prep(plain)["low_s"] < prep(values)["low_s"] < prep(pivots)["low_s"]


def test_a_pivot_too_big_for_formulas_is_priced_as_values():
    from app.estimate import pivot_seconds
    from app.pivots import MAX_FORMULA_CELLS

    # above MAX_FORMULA_CELLS the pivot becomes values without reading its cells one by one
    assert pivot_seconds([("pivot_to_formulas", 1_000_000)]) == pivot_seconds([("pivot_to_values", 1_000_000)])
    assert pivot_seconds([("pivot_to_formulas", MAX_FORMULA_CELLS)]) > pivot_seconds([("pivot_to_values", MAX_FORMULA_CELLS)])


def test_the_originals_recalculation_is_counted_when_it_has_not_run():
    recalc = lambda r: next(s for s in r["steps"] if s["id"] == "recalc")  # noqa: E731
    without = estimate_chain(179, 344, 305_785, 48, 0, True, worker=True)
    pending = estimate_chain(179, 344, 305_785, 48, 0, True, worker=True, baseline_mb=179)
    assert recalc(pending)["low_s"] == recalc(without)["low_s"] == 5  # the fixer's step may have paid it
    assert recalc(pending)["mid_s"] - recalc(without)["mid_s"] >= OPEN_S_PER_MB * 179 and "original" in recalc(pending)["why"]


def test_recalculate_is_ready_only_when_the_worker_works_it_out_ahead():
    recalc = lambda r: next(s for s in r["steps"] if s["id"] == "recalc")  # noqa: E731
    assert recalc(estimate_chain(179, 344, 305_785, 48, 0, True, worker=True))["mid_s"] == RECALC_AHEAD_S["mid"]
    # otherwise its own Excel opens and saves the version
    assert recalc(estimate_chain(179, 344, 305_785, 48, 0, True, worker=True, recalc_ahead=False))["mid_s"] > OPEN_S_PER_MB * 179 + SAVE_S_PER_MB * 179


def _book(tmp_path, name, fill):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "SYNTHETIC_TEST_FIXTURE"
    ws["A2"], ws["B2"] = 1, "=A2*2"
    other = wb.create_sheet("Other")
    other["A1"] = 5
    fill(wb)
    path = tmp_path / name
    wb.save(path)
    return path


def _stored(client, path, **data):
    with path.open("rb") as f:
        res = client.post("/api/sessions", files={"file": (path.name, f, "application/octet-stream")}, data={"mode": "plan", **data})
    assert res.status_code == 200, res.text
    return client.get(f"/api/sessions/{res.json()['sessionId']}/estimate").json()["basis"]["stored_errors"]


def test_no_stored_error_is_said_only_when_the_fixer_has_nothing_else_to_do(tmp_path):
    client = TestClient(server.app)
    clean = _book(tmp_path, "clean.xlsx", lambda wb: None)
    assert _stored(client, clean) == 0
    # a sheet the scan skipped (the fixer reads it), an error typed as a constant, a formula with #REF!
    assert _stored(client, clean, ignore_sheets='["Other"]') is None
    assert _stored(client, _book(tmp_path, "const.xlsx", lambda wb: wb["Other"].__setitem__("B1", "#N/A"))) is None
    assert _stored(client, _book(tmp_path, "ref.xlsx", lambda wb: wb["Other"].__setitem__("B1", "=#REF!+1"))) is None
