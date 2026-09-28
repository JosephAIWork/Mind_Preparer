"""1.7.2: the readiness verdict, the Prep gauge, action levels, and the two
endless-Prep-round causes (a trapped title on a grid's first row; a title
written into a merged cell). Fixtures are small synthetic workbooks."""
from pathlib import Path

import openpyxl
import pytest

from app.config import load_config
from app.modes import plan_mode
from app.prep import ACTIONS, _merged_guard, _read_back_guard, merged_conflict, merged_ranges, plan_actions
from app.readiness import compute_readiness, level_of, prep_progress
from app.rules_engine import RulesEngine

LABEL = "SYNTHETIC_TEST_FIXTURE"


def _run(tmp_path, source):
    return plan_mode.run(source, tmp_path / "work", load_config(), engine=RulesEngine())


def _by_id(plan):
    return {a["id"]: a for a in plan}


# --- A1: a '#Title' on the grid's own first row is not fixed by a row insert ------------
@pytest.fixture
def title_on_first_row_xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Params"
    ws["A1"] = LABEL
    ws["B4"] = "Abatement"  # a label ...
    ws["C4"] = "#Abatement"  # ... next to a title, on the same row as the block's first row
    ws["B5"], ws["C5"] = 1, 0.65
    ws["B6"], ws["C6"] = 2, 0.70
    p = tmp_path / "title_first_row.xlsx"
    wb.save(p)
    wb.close()
    return p


def test_trapped_title_on_first_row_is_left_for_review_not_row_inserted(tmp_path, title_on_first_row_xlsx):
    res = _run(tmp_path, title_on_first_row_xlsx)
    plan = _by_id(plan_actions(res["workbook_analysis"], res["validation_report"]))
    sep = plan["separate_merged_grids"]
    assert sep["count"] == 0, sep["operations"]
    assert any("first row" in s and "by hand" in s for s in sep["skipped"]), sep["skipped"]


# --- A2: a title aimed at a merged cell is skipped, with the merged range named ----------
@pytest.fixture
def merged_above_grid_xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Crit"
    ws["A1"] = LABEL
    ws.merge_cells("B5:D5")  # empty merged band right above the grid
    ws["C6"], ws["D6"] = "Head1", "Head2"
    ws["C7"], ws["D7"] = 1, 2
    ws["C8"], ws["D8"] = 3, 4
    p = tmp_path / "merged_above.xlsx"
    wb.save(p)
    wb.close()
    return p


def test_title_into_non_first_merged_cell_is_skipped(tmp_path, merged_above_grid_xlsx):
    res = _run(tmp_path, merged_above_grid_xlsx)
    analysis = res["workbook_analysis"]
    merged = merged_ranges(analysis)
    assert merged["Crit"] == [(5, 5, 2, 4)]
    assert merged_conflict(merged, "Crit", "C5") == "B5:D5"
    assert merged_conflict(merged, "Crit", "B5") is None  # the first cell is writable
    plan = _by_id(plan_actions(analysis, res["validation_report"]))
    titles = plan["create_grid_titles"]
    assert not any(o.get("cell") == "C5" for o in titles["operations"])
    assert any("merged range B5:D5" in s for s in titles["skipped"]), titles["skipped"]


# --- B1: broken references -- piecewise NA() vs whole-formula =NA() -------------------------
@pytest.fixture
def broken_refs_xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["A1"] = LABEL
    ws["A3"], ws["B3"] = "Key", "Value"
    ws["A4"], ws["B4"] = "x", "=#REF!+1"  # a value-like broken reference: patchable
    ws["A5"], ws["B5"] = "y", "=SUMIFS('Base'!#REF!,A:A,A5)"  # a range-like one: only the whole formula can go
    p = tmp_path / "broken_refs.xlsx"
    wb.save(p)
    wb.close()
    return p


def test_broken_range_reference_gets_its_own_whole_formula_action(tmp_path, broken_refs_xlsx):
    res = _run(tmp_path, broken_refs_xlsx)
    plan = _by_id(plan_actions(res["workbook_analysis"], res["validation_report"]))
    piecewise, whole = plan["fix_broken_refs"], plan["fix_broken_refs_whole"]
    assert [(o["cell"], o["after"]) for o in piecewise["operations"]] == [("B4", "=NA()+1")]
    assert [(o["cell"], o["after"]) for o in whole["operations"]] == [("B5", "=NA()")]
    assert any("B5" in s and "range" in s for s in piecewise["skipped"])
    assert piecewise["level"] == whole["level"] == "blocking"
    assert piecewise["default_on"] and whole["default_on"]
    assert whole["caution"]


# --- levels and priorities ------------------------------------------------------------------
def test_every_action_has_a_level_and_blocking_ones_are_on_by_default():
    for a in ACTIONS:
        assert a["level"] in ("blocking", "optional"), a["id"]
        if a["level"] == "blocking":
            assert a["default_on"], a["id"]


def test_findings_carry_their_rules_priority(tmp_path, broken_refs_xlsx):
    res = _run(tmp_path, broken_refs_xlsx)
    findings = res["validation_report"]["findings"]
    assert findings and all(f.get("priority") in ("REQUIRED", "RECOMMENDED", "INFORMATIONAL") for f in findings)
    ref = next(f for f in findings if f["rule_id"] == "REF-001")
    assert ref["status"] == "ERROR" and ref["priority"] == "REQUIRED" and level_of(ref) == "blocking"


# --- the verdict -------------------------------------------------------------------------------
def _report(*findings):
    return {"findings": list(findings)}


def _f(rule_id, status, priority, **kw):
    return {"rule_id": rule_id, "status": status, "priority": priority, "message": kw.pop("message", rule_id), **kw}


def test_readiness_blocked_then_unverified_then_ready():
    plan = [
        {"id": "fix_broken_refs", "rule_ids": ["REF-001"], "count": 0, "skipped": ["left"], "level": "blocking", "operations": []},
        {"id": "create_grid_titles", "rule_ids": ["STR-004"], "count": 5, "skipped": [], "level": "optional", "operations": []},
    ]
    blocked = compute_readiness(_report(_f("REF-001", "ERROR", "REQUIRED", observed={"sites": [1, 2, 3]}), _f("STR-004", "WARNING", "RECOMMENDED"), _f("READY-001", "NOT_SUPPORTED", "REQUIRED")), plan, None, "ver-001")
    assert blocked["state"] == "blocked" and blocked["blocking_count"] == 1 and blocked["optional_count"] == 1
    assert blocked["blocking"][0] == {**blocked["blocking"][0], "fix": "prep_skipped", "action_id": "fix_broken_refs", "sites": 3}
    assert blocked["manual_blocking_count"] == 1 and blocked["next_step"]["screen"] == "findings"

    unverified = compute_readiness(_report(_f("REF-001", "PASS", "REQUIRED"), _f("STR-004", "WARNING", "RECOMMENDED"), _f("READY-001", "NOT_SUPPORTED", "REQUIRED")), plan, None, "ver-002")
    assert unverified["state"] == "unverified" and unverified["next_step"]["screen"] == "recalculate"

    stale_recalc = {"version_id": "ver-001", "ran": True, "formula_errors": [], "addin_gap_errors": []}
    assert compute_readiness(_report(_f("READY-001", "NOT_SUPPORTED", "REQUIRED")), plan, stale_recalc, "ver-002")["state"] == "unverified"

    dirty = {"version_id": "ver-002", "ran": True, "formula_errors": [{"cell": "A1"}], "addin_gap_errors": []}
    r = compute_readiness(_report(_f("READY-001", "NOT_SUPPORTED", "REQUIRED")), plan, dirty, "ver-002")
    assert r["state"] == "unverified" and r["recalc"]["formula_errors"] == 1

    clean = {"version_id": "ver-002", "ran": True, "formula_errors": [], "addin_gap_errors": [{"cell": "B2"}]}
    r = compute_readiness(_report(_f("READY-001", "NOT_SUPPORTED", "REQUIRED"), _f("FMT-002", "WARNING", "RECOMMENDED")), plan, clean, "ver-002")
    assert r["state"] == "ready" and r["next_step"]["screen"] == "reports" and r["optional_count"] == 1


def test_error_from_a_non_required_rule_is_optional_not_blocking():
    r = compute_readiness(_report(_f("STR-002", "ERROR", "RECOMMENDED")), [], None, "v")
    assert r["state"] == "unverified" and r["optional_count"] == 1 and r["blocking_count"] == 0


def test_blocking_with_prep_operations_routes_to_prep():
    plan = [{"id": "fix_broken_refs", "rule_ids": ["REF-001"], "count": 4, "skipped": [], "level": "blocking", "operations": []}]
    r = compute_readiness(_report(_f("REF-001", "ERROR", "REQUIRED")), plan, None, "v")
    assert r["blocking"][0]["fix"] == "prep" and r["next_step"]["screen"] == "prep" and r["prep"]["blocking_ops"] == 4


# --- the Prep gauge -----------------------------------------------------------------------------
def _hist(totals):
    return [{"version_id": f"v{i}", "total_ops": t, "blocking_ops": 0, "optional_ops": t, "blocking_findings": 0, "by_action": {"create_grid_titles": t}} for i, t in enumerate(totals, 1)]


def test_prep_progress_reports_stall_only_when_three_rounds_do_not_go_down():
    assert prep_progress(_hist([33, 12, 8]))["stalled"] is False
    assert prep_progress(_hist([33, 12, 0]))["stalled"] is False
    stalled = prep_progress(_hist([12, 8, 9, 6, 7, 6, 7]))
    assert stalled["stalled"] is True and stalled["message"]
    assert len(stalled["entries"]) == 7 and stalled["entries"][-1]["version_id"] == "v7"


# --- the executor guards (fake COM objects) --------------------------------------------------
class _Cell:
    def __init__(self, address, formula=""):
        self.Address = address
        self.Formula = formula


class _Area:
    def __init__(self, first, span):
        self._first, self.Address = first, span

    def Cells(self, r, c):
        return _Cell(self._first)


class _Range:
    def __init__(self, merged, first="$B$5", span="$B$5:$D$5", formula=""):
        self.MergeCells = merged
        self._area = _Area(first, span)
        self._cell = _Cell(first, formula)

    def Cells(self, r, c):
        cell = _Cell(self._cell.Address, self._cell.Formula)
        cell.MergeArea = self._area
        return cell


def test_merged_guard_refuses_non_first_cell_and_accepts_the_first():
    with pytest.raises(ValueError, match="merged range B5:D5"):
        _merged_guard(_Range(True), "Crit", "C5")
    _merged_guard(_Range(True), "Crit", "B5")  # no error
    _merged_guard(_Range(False), "Crit", "C5")  # not merged: no error


def test_read_back_guard_catches_a_silent_no_op():
    with pytest.raises(ValueError, match="still empty"):
        _read_back_guard(_Range(False, formula=""), {"op": "set_value", "after": "#Title"}, "S", "C5")
    _read_back_guard(_Range(False, formula="#Title"), {"op": "set_value", "after": "#Title"}, "S", "C5")
    with pytest.raises(ValueError, match="still has content"):
        _read_back_guard(_Range(False, formula="x"), {"op": "clear_cell"}, "S", "C5")


# --- recalculation errors the original already had are the model's own -----------------------
def test_classify_against_original_separates_preexisting_from_new(tmp_path):
    from app.recalc import classify_against_original

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["A1"] = "#DIV/0!"  # cached error value, as Excel would have saved it
    ws["A2"] = "=B2/C2"
    ws["B5"] = "#N/A"
    p = tmp_path / "original.xlsx"
    wb.save(p)
    wb.close()
    errors = [
        {"sheet": "Calc", "cell": "A1", "error": "#DIV/0!", "formula": "=1/0"},  # same address: pre-existing
        {"sheet": "Calc", "cell": "B7", "error": "#N/A", "formula": "=NA()"},  # new
        {"sheet": "Calc", "cell": "A3", "error": "#DIV/0!", "formula": "=B2/C2"},  # moved by a row insert? no error cached there and A2 held no error value -> new
    ]
    out = classify_against_original(p, errors)
    assert out["compared"] is True and out["preexisting"] == 1 and out["new"] == 2
    assert [e["preexisting"] for e in out["errors"]] == [True, False, False]
    missing = classify_against_original(tmp_path / "nope.xlsx", errors)
    assert missing["compared"] is False and missing["new"] == 3


def test_readiness_is_ready_when_every_recalc_error_is_preexisting():
    rc = {"version_id": "v", "ran": True, "formula_errors": [{"cell": "A1"}, {"cell": "A2"}], "preexisting_errors": 2, "new_errors": 0, "addin_gap_errors": []}
    r = compute_readiness(_report(_f("READY-001", "NOT_SUPPORTED", "REQUIRED")), [], rc, "v")
    assert r["state"] == "ready" and r["recalc"]["preexisting_errors"] == 2 and "already in the original" in r["headline"]
    rc["new_errors"], rc["preexisting_errors"] = 1, 1
    r = compute_readiness(_report(_f("READY-001", "NOT_SUPPORTED", "REQUIRED")), [], rc, "v")
    assert r["state"] == "unverified" and r["next_step"]["label"] == "Review the new formula errors"


def test_next_step_goes_to_the_assistant_when_prep_work_targets_no_blocking_finding():
    # a blocking-level action still has work for a rule that is only a WARNING; the one blocking finding has no repair
    plan = [{"id": "separate_merged_grids", "rule_ids": ["STR-001"], "count": 3, "skipped": [], "level": "blocking", "operations": []}]
    r = compute_readiness(_report(_f("FRM-002", "ERROR", "REQUIRED"), _f("STR-001", "WARNING", "REQUIRED")), plan, None, "v")
    assert r["state"] == "blocked" and r["prep"]["blocking_ops"] == 3
    assert r["next_step"]["screen"] == "findings" and r["next_step"]["rule_id"] == "FRM-002"


# --- Excel: a title written into the FIRST cell of a merged range lands (and a label there can be cleared)
needs_excel = pytest.mark.skipif(not __import__("app.excel_com", fromlist=["com_available"]).com_available(), reason="requires pywin32 + an installed Excel")


@needs_excel
def test_excel_writes_into_a_merged_first_cell_and_refuses_the_others(tmp_path):
    from app.prep import apply_operations

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.merge_cells("B2:D2")
    ws["B2"] = "Caption"
    ws.merge_cells("B4:D4")
    p = tmp_path / "merged.xlsx"
    wb.save(p)
    wb.close()
    ops = [
        {"op": "set_value", "action_id": "t", "rule_id": "STR-004", "sheet": "S", "cell": "B2", "before": "Caption", "after": "#Caption"},
        {"op": "clear_cell", "action_id": "t", "rule_id": "STR-002", "sheet": "S", "cell": "B4", "before": None, "after": None},
        {"op": "set_value", "action_id": "t", "rule_id": "STR-004", "sheet": "S", "cell": "C4", "before": None, "after": "#Nope"},
    ]
    res = apply_operations(p, tmp_path / "work", ops, prefer_excel=True)
    assert res["method"] == "excel_com", res
    assert [(o["cell"]) for o in res["applied"]] == ["B2", "B4"]
    assert len(res["failed"]) == 1 and "merged range B4:D4" in res["failed"][0]["error"]
    out = openpyxl.load_workbook(res["output_path"])
    assert out["S"]["B2"].value == "#Caption" and out["S"]["C4"].value is None
    out.close()


def test_classify_errors_uses_the_recalculated_original_as_baseline():
    from app.recalc import classify_errors

    original_errors = [
        {"sheet": "S", "cell": "A1", "error": "#REF!", "formula": "=SUM('B'!#REF!)"},
        {"sheet": "S", "cell": "A2", "error": "#DIV/0!", "formula": "=B2/C2"},
    ]
    errors = [
        {"sheet": "S", "cell": "A1", "error": "#N/A", "formula": "=NA()"},  # same address: the fix replaced a broken formula
        {"sheet": "S", "cell": "A3", "error": "#DIV/0!", "formula": "=B2/C2"},  # shifted by a row insert, same formula text
        {"sheet": "S", "cell": "A9", "error": "#DIV/0!", "formula": "=B9/C9"},  # the original had a #REF! formula there (stale cached number)
        {"sheet": "S", "cell": "A5", "error": "#VALUE!", "formula": "=VALUE(B5)"},  # new
    ]
    out = classify_errors(errors, original_errors, {("S", "A9")})
    assert out["preexisting"] == 3 and out["new"] == 1
    assert [e["preexisting"] for e in out["errors"]] == [True, True, True, False]


def test_formulas_shown_to_the_user_drop_excel_storage_prefixes(tmp_path):
    """openpyxl reads `=_xlfn.NUMBERVALUE(A1)`; Excel shows `=NUMBERVALUE(A1)`.
    The cell window and the FRM-002 call sites show the Excel form, so a user
    who copies a formula into the assistant never carries the prefix along."""
    from app.inventory import build_analysis, cell_window
    from app.validators.formula import unsupported_functions

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["A1"] = "12"
    ws["B1"] = "=_xlfn.NUMBERVALUE(A1)"
    p = tmp_path / "prefix.xlsx"
    wb.save(p)
    analysis = build_analysis(p, tmp_path / "an", "t")
    assert cell_window(analysis, "Calc", "B1", 0, 0)["rows"][0]["cells"][0]["formula"] == "=NUMBERVALUE(A1)"
    rule = RulesEngine().rules["FRM-002"]
    out = unsupported_functions(rule, analysis, load_config())
    assert out["status"] == "ERROR"
    assert [s["formula"] for s in out["observed"]["call_sites"]] == ["=NUMBERVALUE(A1)"]


def test_assistant_gets_the_tool_verdict_on_function_support(tmp_path):
    """"Is IFNA supported?" must be answered from the Mind list the tool owns,
    not from general knowledge: the retrieval carries the verdict, the usage
    and the fact that `_xlfn.` is only Excel's storage prefix."""
    from app.chat_context import SYSTEM_PROMPT, function_support, retrieve
    from app.inventory import build_analysis

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["A1"] = 3
    ws["B1"] = "=_xlfn.IFNA(A1,0)"
    ws["C1"] = "=_xlfn.NUMBERVALUE(\"1\")"
    p = tmp_path / "fn.xlsx"
    wb.save(p)
    analysis = build_analysis(p, tmp_path / "an", "t")
    report = {"findings": []}
    block = retrieve("Is _xlfn.IFNA a problem for Mind? and NUMBERVALUE?", analysis, report)
    assert "## Function support" in block
    # IFNA is NOT on the Mind list (IFERROR is): the assistant must say so, and give the behaviour-preserving equivalent
    assert "IFNA: NOT on the Mind supported-function list" in block and "used in 1 formula(s), first at Calc!B1" in block
    assert "IF(ISNA(x),alt,x)" in block
    assert "NUMBERVALUE: NOT on the Mind supported-function list" in block and 'IF(x="",0,VALUE(x))' in block
    assert "IFERROR: on the Mind supported-function list" in function_support(["IFERROR"], analysis)
    assert "storage prefix, not a defect" in block
    # ordinary words are not functions
    assert function_support(["problem", "the"], analysis) == ""
    assert "MM_LOOP" in function_support(["mm_loop"], analysis)
    assert "never hedge" in SYSTEM_PROMPT and "storage prefixes" in SYSTEM_PROMPT
    # the general rule: no conclusion about Mind from anything but the Mind documentation the tool holds
    assert "must come from the Mind documentation the tool holds" in SYSTEM_PROMPT and "is NOT evidence about Mind" in SYSTEM_PROMPT


def test_spilled_range_references_are_blocking_and_frozen_by_prep(tmp_path):
    """Mind rejects INDEX(A1#, ...) with 'Unsupported formula: ANCHORARRAY()'.
    FRM-003 must fail (ERROR, blocking) and Prep must replace each reference
    by the fixed range the spill covers today, on the same sheet or across
    sheets, leaving unresolvable ones for review."""
    from openpyxl.worksheet.formula import ArrayFormula

    from app.inventory import build_analysis
    from app.prep import ACTION_LEVELS, freeze_spill_refs, plan_freeze_spill_refs
    from app.validators.formula import frm_003

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model Inputs"
    for r in range(1, 4):
        ws.cell(r, 2, r)
    ws["A1"] = ArrayFormula("A1:A3", "=B1:B3*2")  # a spill anchor covering A1:A3
    ws["E169"] = "=INDEX(_xlfn.ANCHORARRAY(A1),2)"  # how Excel stores INDEX(A1#,2)
    ws["E171"] = "=SUM(A1#)"  # the visible form
    ws["E173"] = '=SUM(Z9#)&"A1#"'  # anchor unknown -> review; the string literal stays untouched
    other = wb.create_sheet("Cluster")
    other["BG74"] = "=SUMPRODUCT('Model Inputs'!$A$1#)"
    p = tmp_path / "spill.xlsx"
    wb.save(p)
    analysis = build_analysis(p, tmp_path / "an", "t")

    rule = RulesEngine().rules["FRM-003"]
    out = frm_003(rule, analysis, load_config())
    assert out["status"] == "ERROR" and "ANCHORARRAY()" in out["message"]
    assert level_of({"rule_id": "FRM-003", "status": out["status"], "priority": rule["priority"]}) == "blocking"

    ops, skipped = plan_freeze_spill_refs(analysis, {"findings": []})
    after = {(o["sheet"], o["cell"]): o["after"] for o in ops}
    assert after[("Model Inputs", "E169")] == "=INDEX(A1:A3,2)"
    assert after[("Model Inputs", "E171")] == "=SUM(A1:A3)"
    assert after[("Cluster", "BG74")] == "=SUMPRODUCT('Model Inputs'!$A$1:$A$3)"  # absolute stays absolute
    notes = {(o["sheet"], o["cell"]): o["note"] for o in ops}
    assert notes[("Cluster", "BG74")].startswith("'Model Inputs'!$A$1# -> 'Model Inputs'!$A$1:$A$3")
    assert all(o["rule_id"] == "FRM-003" and o["action_id"] == "freeze_spill_refs" for o in ops)
    assert len(skipped) == 1 and "E173" in skipped[0] and "Z9#" in skipped[0]
    assert ACTION_LEVELS["freeze_spill_refs"] == "blocking"
    # the rewrite never touches string literals
    fixed, unresolved = freeze_spill_refs('=A1#&"A1#"', "Model Inputs", {("Model Inputs", "A1"): "A1:A3"})
    assert fixed == '=A1:A3&"A1#"' and unresolved == []
