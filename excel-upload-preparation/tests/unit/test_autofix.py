"""1.8.0: the automatic formula-error fixer (app/autofix.py).

The pure parts (R1C1 reading, the range index, the strategy ladder) run
everywhere; the engine itself needs a real Excel and is skipped without one."""
from pathlib import Path

import openpyxl
import pytest

from app import autofix
from app.autofix import AutoFixConfig, CellIndex, candidate_fixes, parse_r1c1, r1c1_to_a1, rectangles, ref_rect, run_autofix
from app.excel_com import com_available

needs_excel = pytest.mark.skipif(not com_available(), reason="needs a real Excel")


# --- reading R1C1 formulas -----------------------------------------------------------------
def test_relative_and_absolute_references_resolve_from_the_cell():
    p = parse_r1c1("=R[-1]C+RC[3]-R[-1]C[3]")
    assert [ref_rect(r, 100, 10) for r in p.refs] == [(99, 99, 10, 10), (100, 100, 13, 13), (99, 99, 13, 13)]
    assert not p.dynamic and not p.external and not p.broken and not p.names
    p = parse_r1c1("=SUM(R5C:R[2]C[1])+R1C1")
    assert [ref_rect(r, 100, 10) for r in p.refs] == [(5, 102, 10, 11), (1, 1, 1, 1)]


def test_whole_columns_rows_sheets_and_3d_references():
    p = parse_r1c1("=SUMIFS('BASE Polices'!#REF!,OUTPUT!C[13],'C20'!R[-1]C[-2])")
    assert p.broken
    specs = [(r[0], ref_rect(r, 20, 5)) for r in p.refs]
    assert (("sheet", "OUTPUT"), (1, autofix.MAX_ROW, 18, 18)) in specs
    assert (("sheet", "C20"), (19, 19, 3, 3)) in specs
    p = parse_r1c1("=SUM('LoB 1:>>'!RC)")
    assert p.refs[0][0] == ("3d", "LoB 1", ">>") and ref_rect(p.refs[0], 7, 8) == (7, 7, 8, 8)
    p = parse_r1c1("=SUM(R[2]:R[3])")
    assert ref_rect(p.refs[0], 10, 4) == (12, 13, 1, autofix.MAX_COL)


def test_external_links_names_functions_and_dynamic_references():
    p = parse_r1c1(r"='C:\Users\a\Documents\[PVFP Inputs_v4.0.xlsx]BP by LoB'!R[-98]C[-2]")
    assert p.external and p.refs[0][0] == ("ext",)
    p = parse_r1c1('=IF(INDIRECT("\'"&RC7&"\'!$Z$16")=0,0,66%*R[277]C26)')
    assert p.dynamic and "INDIRECT" in p.funcs
    assert [ref_rect(r, 10, 3) for r in p.refs] == [(10, 10, 7, 7), (287, 287, 26, 26)]  # nothing read out of the string
    p = parse_r1c1("=VLOOKUP(Policy_tariff,Assumptions_array,MATCH(R[-3]C,Assumption_headings,0),FALSE)")
    assert p.names == ("ASSUMPTIONS_ARRAY", "ASSUMPTION_HEADINGS", "POLICY_TARIFF") and {"VLOOKUP", "MATCH"} <= p.funcs
    p = parse_r1c1("=ROUND(RATE(1,2,3)*COUNT(RC[1]:RC[3]),2)+_xlfn.XLOOKUP(RC1,C1,C2)")
    assert {"ROUND", "RATE", "COUNT", "XLOOKUP"} <= p.funcs and not p.names  # RATE( / ROUND( / COUNT( are not references
    assert len(p.refs) == 4


def test_r1c1_text_becomes_the_a1_text_people_read():
    assert r1c1_to_a1("=R[-1]C+RC[3]-R[-1]C[3]", 100, 10) == "=J99+M100-M99"
    assert r1c1_to_a1("=IFERROR(RC[-4]/RC[-3],0)", 8, 9) == "=IFERROR(E8/F8,0)"
    assert r1c1_to_a1("=SUM(R5C2:R9C2)*'My sheet'!R1C1", 3, 3) == "=SUM($B$5:$B$9)*'My sheet'!$A$1"
    assert r1c1_to_a1("=SUMIFS(OUTPUT!C[13],OUTPUT!C1,RC1)", 20, 5) == "=SUMIFS(OUTPUT!R:R,OUTPUT!$A:$A,$A20)"
    assert r1c1_to_a1('=IF(RC[-1]="R1C1 text",R[1]C,0)', 2, 2) == '=IF(A2="R1C1 text",B3,0)'  # strings are left alone
    assert r1c1_to_a1("=SUM(R[2]:R[3])", 10, 4) == "=SUM(12:13)"


# --- the range index -------------------------------------------------------------------------
def test_cell_index_answers_rectangle_queries_from_either_side():
    idx = CellIndex([(0, 5, 3), (0, 5, 9), (0, 700, 3), (1, 1, 1)])
    assert idx.any_in(0, 5, 5, 3, 3) and not idx.any_in(0, 5, 5, 4, 4)
    assert idx.any_in(0, 1, 10_000, 9, 9) and idx.any_in(0, 700, 700, 1, 16_000)
    assert not idx.any_in(0, 6, 699, 1, 16_000) and not idx.any_in(2, 1, 9, 1, 9)
    assert not idx.any_in(0, 5, 5, 3, 3, exclude=(5, 3)) and idx.any_in(0, 5, 5, 1, 9, exclude=(5, 3))
    assert sorted(idx.cells_in(0, 1, 1000, 1, 5)) == [(5, 3), (700, 3)]


def test_rectangles_cover_blocks_with_few_pieces():
    assert rectangles([(1, 1), (1, 2), (2, 1), (2, 2), (3, 1), (5, 5)]) == [(1, 1, 2, 2), (3, 1, 3, 1), (5, 5, 5, 5)]
    block = [(r, c) for r in range(10, 60) for c in range(4, 104)]
    assert rectangles(block) == [(10, 4, 59, 103)]


# --- the strategy ladder ---------------------------------------------------------------------
def test_a_live_formula_is_kept_and_only_given_a_value_to_fall_back_on():
    fixes = candidate_fixes("=+RC[-4]/RC[-3]", parse_r1c1("=+RC[-4]/RC[-3]"), "number")
    assert [(f["kind"], f["content"]) for f in fixes] == [("formula", "=IFERROR(+RC[-4]/RC[-3],0)"), ("formula", '=IFERROR(+RC[-4]/RC[-3],"")')]
    fixes = candidate_fixes('=RC[-1]&"x"', parse_r1c1('=RC[-1]&"x"'), "text")
    assert fixes[0]["content"] == '=IFERROR(RC[-1]&"x","")'
    fixes = candidate_fixes("=RC[-1]=RC[-2]", parse_r1c1("=RC[-1]=RC[-2]"), "bool")
    assert fixes[0]["content"] == "=IFERROR(RC[-1]=RC[-2],FALSE)"
    # the assistant's choice goes first, the rest of the ladder stays behind it
    fixes = candidate_fixes("=+RC[-4]/RC[-3]", parse_r1c1("=+RC[-4]/RC[-3]"), "number", preferred='""')
    assert [f["content"] for f in fixes] == ['=IFERROR(+RC[-4]/RC[-3],"")', "=IFERROR(+RC[-4]/RC[-3],0)"]


def test_a_dead_formula_is_replaced_and_a_broken_one_loses_its_ref_literal():
    dead = candidate_fixes("=+NA()", parse_r1c1("=+NA()"), "number")
    assert [(f["kind"], f["content"]) for f in dead] == [("value", 0), ("value", "")]
    broken = candidate_fixes("=SUM('BASE Polices'!#REF!)+RC[1]", parse_r1c1("=SUM('BASE Polices'!#REF!)+RC[1]"), "number")
    assert broken[0]["content"] == "=IFERROR(SUM(NA())+RC[1],0)" and "#REF!" not in broken[0]["content"]
    endpoint = candidate_fixes("=SUM(R1C1:#REF!)", parse_r1c1("=SUM(R1C1:#REF!)"), "number")
    assert [f["kind"] for f in endpoint] == ["value", "value"]  # the range can never be rebuilt: no IFERROR around a #REF!


# --- the engine, in a real Excel -----------------------------------------------------------
def _book(tmp_path: Path, fill) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    fill(wb, ws)
    p = tmp_path / "model.xlsx"
    wb.save(p)
    wb.close()
    return p


def _values(path: Path) -> dict[str, object]:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return {f"{ws.title}!{c.coordinate}": c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None}
    finally:
        wb.close()


@needs_excel
def test_roots_are_fixed_and_everything_downstream_clears(tmp_path):
    def fill(wb, ws):
        ws["A1"], ws["B1"], ws["C1"] = "Premium", "Policies", "Average"
        for r, (a, b) in enumerate([(100, 4), (50, 0), (30, 3)], start=2):
            ws[f"A{r}"], ws[f"B{r}"] = a, b
            ws[f"C{r}"] = f"=A{r}/B{r}"          # C3 is the root: #DIV/0!
            ws[f"D{r}"] = f"=C{r}*2"             # D3 only repeats it
        ws["C5"] = "=SUM(C2:C4)"                 # ... and so does the total
        ws["F2"] = "=A2+A3"                      # a good number, untouched

    src = _book(tmp_path, fill)
    res = run_autofix(src, tmp_path / "work")
    assert res["status"] == "clean" and res["errors_before"] == 3 and res["errors_after"] == 0
    assert res["cells_rewritten"] == 1 and res["passes"] == 1  # one root; the two cells downstream are not touched
    assert res["fixes"][0]["cell"] == "C3" and res["fixes"][0]["after"] == "=IFERROR(A3/B3,0)"
    after = _values(Path(res["output_path"]))
    assert after["Calc!C3"] == 0 and after["Calc!D3"] == 0 and after["Calc!C5"] == 35 and after["Calc!F2"] == 150
    assert res["recalc"]["status"] == "PASS" and res["recalc"]["formula_errors"] == []
    assert res["applied"][0]["op"] == "set_formula" and res["applied"][0]["action_id"] == "autofix"
    assert _values(src)["Calc!A2"] == 100  # the source is never modified


@needs_excel
def test_a_fix_that_changes_a_good_number_is_rolled_back_and_another_one_tried(tmp_path):
    def fill(wb, ws):
        ws["A1"], ws["B1"] = 10, 0
        ws["C1"] = "=A1/B1"                      # #DIV/0!
        ws["A2"], ws["B2"] = 8, 2
        ws["C2"] = "=A2/B2"
        ws["E1"] = "=COUNT(C1:C2)"               # 1 today: counts numbers only -> a 0 in C1 would make it 2
        ws["E2"] = "=E1*100"

    src = _book(tmp_path, fill)
    res = run_autofix(src, tmp_path / "work")
    assert res["status"] == "clean" and res["rolled_back"] >= 1
    assert res["fixes"][0]["after"] == '=IFERROR(A1/B1,"")'  # 0 was tried first and rolled back; empty text keeps the count
    after = _values(Path(res["output_path"]))
    assert after["Calc!E1"] == 1 and after["Calc!E2"] == 100


@needs_excel
def test_an_error_that_cannot_be_fixed_without_changing_a_number_is_left_for_a_person(tmp_path):
    def fill(wb, ws):
        ws["A1"], ws["B1"] = 10, 0
        ws["C1"] = "=A1/B1"                      # #DIV/0!
        ws["E1"] = "=IF(ISERROR(C1),1,2)"        # 1 today; any value in C1 turns it into 2
        ws["A3"], ws["B3"] = 5, 0
        ws["C3"] = "=A3/B3"                      # an independent root: fixed all the same

    src = _book(tmp_path, fill)
    res = run_autofix(src, tmp_path / "work")
    assert res["status"] == "partial" and res["errors_before"] == 2 and res["errors_after"] == 1
    left = res["left"][0]
    assert left["cell"] == "C1" and "good value" in left["why"]
    after = _values(Path(res["output_path"]))
    assert after["Calc!E1"] == 1 and after["Calc!C3"] == 0


@needs_excel
def test_layers_of_errors_take_several_passes_and_blocks_are_written_as_one(tmp_path):
    def fill(wb, ws):
        ws["A1"] = "rate"
        for c in range(2, 42):                   # 40 columns sharing one R1C1 formula
            col = openpyxl.utils.get_column_letter(c)
            ws[f"{col}2"] = 0
            ws[f"{col}3"] = f"=1/{col}2"                          # layer 1: #DIV/0!
            ws[f"{col}4"] = f"=IF({col}3=0,1/{col}2,{col}3)"      # repeats row 3's error; once row 3 is 0 it divides by zero itself
        ws["A6"] = "=SUM(B4:AO4)"

    src = _book(tmp_path, fill)
    res = run_autofix(src, tmp_path / "work")
    assert res["status"] == "clean" and res["errors_before"] == 81 and res["passes"] == 2
    assert res["cells_rewritten"] == 80 and len(res["applied"]) == 2  # two blocks of 40, one write each
    assert {g["cells"] for g in res["fixes"]} == {40}


@needs_excel
def test_a_clean_workbook_is_left_alone(tmp_path):
    def fill(wb, ws):
        ws["A1"], ws["A2"] = 1, 2
        ws["A3"] = "=A1+A2"

    res = run_autofix(_book(tmp_path, fill), tmp_path / "work")
    assert res["status"] == "unchanged" and res["output_path"] is None and res["errors_before"] == 0


@needs_excel
def test_array_formulas_and_dead_formulas_and_error_constants(tmp_path):
    def fill(wb, ws):
        ws["A1"], ws["A2"], ws["A3"] = 4, 0, 2
        ws["B1"] = "=10/A1"                      # becomes the 3-cell array formula =10/A1:A3 below; its middle cell is #DIV/0!
        ws["D1"] = "=NA()"                       # nothing left to compute
        ws["D2"] = "=D1+1"
        ws["F1"] = "#N/A"                        # an error value sitting in a data cell
        ws["F2"] = "=F1*2"

    src = _book(tmp_path, fill)
    # openpyxl writes the error text as a string: make F1 a real error value and B1:B3 a real array through Excel
    from app.excel_com import close_quietly, excel_session, open_for_write, save_in_place

    with excel_session() as excel:
        def prepare():
            wb = open_for_write(excel, src)
            try:
                ws = wb.Worksheets("Calc")
                ws.Range("B1:B3").FormulaArray = "=10/A1:A3"
                ws.Range("F1").Formula = "#N/A"
                ws = None
                save_in_place(wb, src)
            finally:
                close_quietly(wb)

        prepare()
    res = run_autofix(src, tmp_path / "work")
    assert res["status"] == "clean", res.get("left")
    by_cell = {g["cell"]: g for g in res["fixes"]}
    assert by_cell["B1"]["kind"] == "array" and by_cell["B1"]["after"] == "=IFERROR(10/A1:A3,0)"
    assert by_cell["D1"]["after"] == 0 and by_cell["F1"]["strategy"] == "error value cleared"
    after = _values(Path(res["output_path"]))
    assert after["Calc!B1"] == 2.5 and after["Calc!B2"] == 0 and after["Calc!B3"] == 5 and after["Calc!D2"] == 1 and after["Calc!F2"] == 0
