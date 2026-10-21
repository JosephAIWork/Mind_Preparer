"""1.8.0: a recalculation finds its error cells in the values it has already
read, computes their addresses, and asks Excel about array formulas for a
whole sheet (then a row) at a time. Asking about each error cell -- Cells,
Address, HasArray -- took ten minutes for a model with 65,883 error cells,
twice (the original is recalculated too): 25 minutes before the first
Recalculate of that model answered."""
import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from app import recalc
from app.excel_com import com_available

needs_excel = pytest.mark.skipif(not com_available(), reason="needs a real Excel")


@needs_excel
def test_error_cells_keep_their_addresses_and_their_array(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plain"                                       # a used range that does not start at A1
    ws["C3"] = 0
    ws["D3"] = "=1/C3"
    ws["AB7"] = "=D3+1"                                      # a two-letter column
    ws["E5"] = '=VLOOKUP("x",C3:C4,1,FALSE)'
    ws["F5"] = "=C3+1"                                       # fine
    arr = wb.create_sheet("Arr")
    arr["A5"], arr["A6"] = 1, 2
    arr["C5"] = ArrayFormula("C5:C9", "=A5:A6*2")            # C7:C9 overflow to #N/A
    arr["E7"] = "=1/0"                                       # an ordinary error on a row that also holds array cells
    arr["G2"] = ArrayFormula("G2", "=SUM(1/(A5:A6-A5:A6))")  # an array formula in one cell
    wb.create_sheet("Clean")["B2"] = "=1+1"
    path = tmp_path / "scan.xlsx"
    wb.save(path)
    wb.close()

    rc = recalc.recalculate(path)
    assert rc["ran"] is True and rc["status"] == "ERROR"
    got = {(e["sheet"], e["cell"]): e for e in rc["formula_errors"]}
    assert set(got) == {("Plain", "D3"), ("Plain", "AB7"), ("Plain", "E5"), ("Arr", "C7"), ("Arr", "C8"), ("Arr", "C9"), ("Arr", "E7"), ("Arr", "G2")}
    assert got[("Plain", "D3")] == {"sheet": "Plain", "cell": "D3", "error": "#DIV/0!", "formula": "=1/C3"}
    assert got[("Plain", "AB7")]["error"] == "#DIV/0!" and got[("Plain", "AB7")]["formula"] == "=D3+1"
    assert got[("Plain", "E5")]["error"] == "#N/A"
    assert all(got[("Arr", c)].get("array") == "C5:C9" for c in ("C7", "C8", "C9"))   # can only be changed as a whole
    assert "array" not in got[("Arr", "E7")] and "array" not in got[("Arr", "G2")]


class _Cell:
    def __init__(self, sheet, row, col):
        self.row, self.col = row, col
        self.HasArray = (row, col) in sheet.arrays
        self.CurrentArray = type("A", (), {"Address": sheet.arrays.get((row, col), "")})()


class _Span:
    def __init__(self, sheet, a, b):
        inside = [(a.row, c) in sheet.arrays for c in range(a.col, b.col + 1)]
        self.HasArray = True if all(inside) else (False if not any(inside) else None)


class _Sheet:
    """Counts what is asked of Excel."""

    def __init__(self, arrays):
        self.arrays = arrays
        self.cells_asked: list[tuple[int, int]] = []
        self.rows_asked: list[int] = []

    def Cells(self, row, col):
        self.cells_asked.append((row, col))
        return _Cell(self, row, col)

    def Range(self, a, b):
        self.rows_asked.append(a.row)
        return _Span(self, a, b)


def test_a_sheet_without_array_formulas_costs_one_question():
    sheet = _Sheet({})
    cells = [(r, c, "#REF!", "=x") for r in range(1, 400) for c in range(1, 30)]
    assert recalc._array_formulas(sheet, type("U", (), {"HasArray": False})(), cells) == {}
    assert sheet.cells_asked == [] and sheet.rows_asked == []


def test_only_the_rows_that_hold_an_array_are_looked_at_cell_by_cell():
    sheet = _Sheet({(7, 3): "$C$5:$C$9", (8, 3): "$C$5:$C$9"})
    cells = [(5, 2, "#N/A", "=a"), (5, 9, "#N/A", "=a"), (7, 3, "#N/A", "=b"), (7, 5, "#DIV/0!", "=1/0"), (8, 3, "#N/A", "=b")]
    out = recalc._array_formulas(sheet, type("U", (), {"HasArray": None})(), cells)
    assert out == {(7, 3): "C5:C9", (8, 3): "C5:C9"}
    assert sheet.rows_asked == [5, 7, 8]
    assert set(sheet.cells_asked) <= {(5, 2), (5, 9), (7, 3), (7, 5), (8, 3)}
    # row 5 holds no array: its two cells were only used to draw the row's span, never asked one by one
    assert sheet.cells_asked.count((5, 2)) == 1 and sheet.cells_asked.count((5, 9)) == 1
    assert sheet.cells_asked.count((7, 5)) == 2                  # the span's end, then the cell itself


# --- 1.11.x: pre-existing errors follow the original's cells through Prep's row inserts ----------
def _down_one_row_from_10(sheet, row, col):
    """A title row inserted at row 10 of S, and S renamed T (as Moves.forward answers)."""
    if sheet == "S":
        return "T", row + (row >= 10), col
    return sheet, row, col


def test_an_old_error_moved_down_by_a_title_row_stays_preexisting():
    original = [{"sheet": "S", "cell": "C10", "error": "#DIV/0!", "formula": "=A10/B10"}]
    # the original's error is now at C11 (formula text shifted too), and a cell the preparation
    # broke landed on the old address C10
    now = [
        {"sheet": "T", "cell": "C11", "error": "#DIV/0!", "formula": "=A11/B11"},
        {"sheet": "T", "cell": "C10", "error": "#VALUE!", "formula": "=A9+B9"},
    ]
    out = recalc.classify_errors(now, original, {("S", "$D$12")}, _down_one_row_from_10)
    assert [e["preexisting"] for e in out["errors"]] == [True, False]
    assert (out["preexisting"], out["new"]) == (1, 1)
    # the original's #REF! formulas move as well
    ref = recalc.classify_errors([{"sheet": "T", "cell": "D13", "error": "#REF!", "formula": "=#REF!"}], [], {("S", "$D$12")}, _down_one_row_from_10)
    assert ref["preexisting"] == 1
    # without the moves, as before: matched by address only
    same = recalc.classify_errors(now, original)
    assert same["preexisting"] == 0
    # a good cell from C9 pushed onto C10 now reads the old error's text (=A10/B10): with the moves
    # known it is not taken for the old error (the text match is only the fallback without them)
    pushed = [{"sheet": "T", "cell": "C10", "error": "#DIV/0!", "formula": "=A10/B10"}]
    assert recalc.classify_errors(pushed, original, None, _down_one_row_from_10)["new"] == 1


def test_the_cached_value_fallback_moves_the_original_too(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["C10"] = "#DIV/0!"  # what openpyxl reads back as the cached value of an error cell
    p = tmp_path / "original.xlsx"
    wb.save(p)
    now = [{"sheet": "T", "cell": "C11", "error": "#DIV/0!", "formula": "=A11/B11"}, {"sheet": "T", "cell": "C10", "error": "#DIV/0!", "formula": "=A9/B9"}]
    out = recalc.classify_against_original(p, now, _down_one_row_from_10)
    assert [e["preexisting"] for e in out["errors"]] == [True, False]
