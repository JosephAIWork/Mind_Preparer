import pytest

from app.config import load_config
from app.inventory import build_analysis
from app.rules_engine import RulesEngine
from app.validators import run_rule


@pytest.fixture(scope="module")
def engine():
    return RulesEngine()


@pytest.fixture(scope="module")
def config():
    return load_config()


def _finding(engine, config, analysis, rule_id):
    rule = engine.get(rule_id)
    assert rule is not None, f"{rule_id} missing from loaded rules"
    return run_rule(rule, analysis, config)


def _statuses(engine, config, analysis, rule_ids):
    return {rid: _finding(engine, config, analysis, rid)["status"] for rid in rule_ids}


# --- loops (1.1.0 behaviour preserved) ------------------------------------------
def test_loop_inventory_and_naming_pass_on_clean_fixture(tmp_path, engine, config, loops_ok_xlsx):
    analysis = build_analysis(loops_ok_xlsx, tmp_path / "work", "t1")
    assert _statuses(engine, config, analysis, ["LOOP-001", "LOOP-002", "LOOP-003", "LBL-001"]) == {
        "LOOP-001": "PASS", "LOOP-002": "PASS", "LOOP-003": "PASS", "LBL-001": "PASS"}


def test_loop_validators_fail_on_broken_fixture(tmp_path, engine, config, loops_broken_xlsx):
    analysis = build_analysis(loops_broken_xlsx, tmp_path / "work", "t2")
    naming = _finding(engine, config, analysis, "LOOP-002")
    assert naming["status"] == "ERROR"
    assert "scenario" in str(naming["observed"]).lower()
    assert _finding(engine, config, analysis, "LOOP-003")["status"] == "ERROR"
    assert _finding(engine, config, analysis, "LBL-001")["status"] == "ERROR"


# --- formulas -----------------------------------------------------------------------
def test_unsupported_functions_flags_unregistered_mm_call(tmp_path, engine, config, unregistered_function_xlsx):
    analysis = build_analysis(unregistered_function_xlsx, tmp_path / "work", "t3")
    finding = _finding(engine, config, analysis, "FRM-002")
    assert finding["status"] == "ERROR"
    assert "MM_TOTALLYMADEUP" in finding["observed"]["unregistered_mm_functions"]


def test_unsupported_functions_passes_on_registered_calls(tmp_path, engine, config, registered_functions_only_xlsx):
    analysis = build_analysis(registered_functions_only_xlsx, tmp_path / "work", "t4")
    assert _finding(engine, config, analysis, "FRM-002")["status"] == "PASS"


def test_looplabels_is_registered_via_the_kb_registry(tmp_path, engine, config, loops_ok_xlsx):
    """1.2.0 flagged MM_LOOPLABELS because doc 02 has no entry for it; the KB
    registry mined from CompleteMindDocn.docx (1.3.0) documents it."""
    analysis = build_analysis(loops_ok_xlsx, tmp_path / "work", "t4b")
    finding = _finding(engine, config, analysis, "FRM-002")
    assert finding["status"] == "PASS"
    assert finding["observed"]["unregistered_mm_functions"] == []


def test_native_function_off_the_kb_list_and_vba_udf_are_flagged(tmp_path, engine, config, array_formula_xlsx):
    analysis = build_analysis(array_formula_xlsx, tmp_path / "work", "t4c")
    frm = _finding(engine, config, analysis, "FRM-002")
    # XLOOKUP is not on the KB page; it is accepted on the evidence recorded for it, and said so
    assert "XLOOKUP" in frm["observed"]["accepted_beyond_kb_list"]
    assert "XLOOKUP" not in frm["observed"]["undocumented_native_functions"]
    assert "MYMACROFUNCTION" not in frm["observed"]["undocumented_native_functions"]
    assert frm["status"] != "ERROR"
    udf = _finding(engine, config, analysis, "FORMULA-002")
    assert udf["status"] == "WARNING"  # no VBA project in this fixture -> unknown function, not a confirmed UDF
    assert "MYMACROFUNCTION" in udf["observed"]["udf_candidates"]
    dyn = _finding(engine, config, analysis, "FRM-003")
    assert dyn["status"] == "WARNING"
    assert dyn["observed"]["counts"]["implicit_intersection"] == 1
    assert dyn["observed"]["counts"]["array_formulas"] == 1


def test_array_formula_past_grid_boundary_is_oversized(tmp_path, engine, config, array_formula_xlsx):
    analysis = build_analysis(array_formula_xlsx, tmp_path / "work", "t4d")
    finding = _finding(engine, config, analysis, "RSK-002")
    assert finding["status"] == "ERROR"
    assert finding["observed"]["oversized_array_formulas"][0]["cell"] == "C5"


def test_vba_presence(tmp_path, engine, config, plain_grid_xlsx, macro_xlsm):
    clean = build_analysis(plain_grid_xlsx, tmp_path / "work1", "t5")
    assert _finding(engine, config, clean, "FRM-004")["status"] == "PASS"
    with_vba = build_analysis(macro_xlsm, tmp_path / "work2", "t6")
    # Mind ignores VBA it can't run -- presence alone isn't a blocker, only WARNING.
    assert _finding(engine, config, with_vba, "FRM-004")["status"] == "WARNING"


# --- format / structure / risk ------------------------------------------------------
def test_format_and_structure_inventory_warn_on_messy_fixture(tmp_path, engine, config, messy_workbook_xlsx):
    analysis = build_analysis(messy_workbook_xlsx, tmp_path / "work", "t7")
    assert _statuses(engine, config, analysis, ["FMT-004", "FMT-005", "FMT-007", "RSK-003", "STR-007"]) == {
        "FMT-004": "WARNING", "FMT-005": "WARNING", "FMT-007": "WARNING", "RSK-003": "WARNING", "STR-007": "WARNING"}


def test_clean_fixture_passes_format_and_risk_checks(tmp_path, engine, config, plain_grid_xlsx):
    analysis = build_analysis(plain_grid_xlsx, tmp_path / "work", "t8")
    assert _statuses(engine, config, analysis, ["FMT-001", "FMT-002", "FMT-003", "FMT-004", "FMT-006", "RSK-003", "RSK-004", "MMX-002"]) == {
        "FMT-001": "PASS", "FMT-002": "PASS", "FMT-003": "PASS", "FMT-004": "PASS", "FMT-006": "PASS", "RSK-003": "PASS", "RSK-004": "PASS", "MMX-002": "PASS"}


def test_recalculation_always_not_supported(tmp_path, engine, config, plain_grid_xlsx):
    analysis = build_analysis(plain_grid_xlsx, tmp_path / "work", "t9")
    assert _finding(engine, config, analysis, "READY-001")["status"] == "NOT_SUPPORTED"


def test_every_active_rule_has_a_real_validator(engine):
    """1.3.0: no active rule falls through to the 'Not implemented' path any more."""
    missing = [r["id"] for r in engine.active_rules() if RulesEngine.resolve_validator(r["validation"]["implementation"]) is None]
    assert missing == []


def test_unresolvable_implementation_still_reports_on_unevaluable(tmp_path, engine, config, plain_grid_xlsx):
    """The honest fallback (SKILL.md rule #8) is still wired for any future rule."""
    analysis = build_analysis(plain_grid_xlsx, tmp_path / "work", "t10")
    rule = dict(engine.get("RES-001"))
    rule["validation"] = {**rule["validation"], "implementation": "validators.loop.does_not_exist"}
    finding = run_rule(rule, analysis, config)
    assert finding["status"] == rule["on_unevaluable"]
    assert "Not implemented" in finding["message"]


# --- the KB-backed validators on a well-formed model -----------------------------------
GOOD_EXPECTATIONS = {
    "STR-001": "PASS", "STR-003": "PASS", "STR-004": "PASS", "STR-005": "PASS", "STR-006": "PASS",
    "INP-001": "PASS", "INP-002": "PASS", "INP-003": "PASS", "INP-004": "PASS", "INP-005": "PASS", "INP-006": "PASS",
    "EXP-001": "PASS", "EXP-002": "PASS", "EXP-003": "PASS", "EXP-004": "PASS",
    "PAR-001": "PASS", "PAR-002": "PASS", "PAR-003": "PASS", "PAR-004": "PASS",
    "PRJ-001": "PASS", "PRJ-002": "PASS", "PRJ-003": "WARNING", "PRJ-005": "PASS", "PRJ-006": "PASS",
    "RES-001": "PASS", "RES-002": "PASS", "RES-004": "PASS", "RES-005": "PASS",
    "RZS-001": "PASS", "RZS-002": "PASS", "RZS-003": "PASS", "RZS-004": "PASS", "RZS-005": "PASS", "RZS-007": "PASS",
    "RSK-001": "PASS", "RSK-002": "PASS",
    "LKP-001": "PASS", "LKP-002": "PASS", "LKP-003": "PASS", "LKP-004": "PASS", "LKP-005": "PASS",
    "CAL-001": "PASS", "CAL-002": "PASS", "CAL-004": "PASS", "CAL-005": "PASS",
    "MMX-001": "PASS", "MMX-002": "PASS", "MMX-003": "PASS",
    "FLG-001": "PASS", "GRP-001": "PASS", "BKP-001": "PASS", "HID-001": "PASS", "RNG-001": "PASS", "UNQ-001": "PASS", "TRN-001": "PASS", "SIM-001": "PASS",
    "DBG-003": "PASS", "DBG-004": "PASS",
}


def test_well_formed_model_passes_the_kb_backed_validators(tmp_path, engine, config, flagged_model_xlsx):
    analysis = build_analysis(flagged_model_xlsx, tmp_path / "work", "good")
    actual = {}
    for rid, expected in GOOD_EXPECTATIONS.items():
        f = _finding(engine, config, analysis, rid)
        actual[rid] = (f["status"], f["message"][:120])
    mismatches = {rid: actual[rid] for rid, expected in GOOD_EXPECTATIONS.items() if actual[rid][0] != expected}
    assert mismatches == {}


BROKEN_EXPECTATIONS = {
    "STR-001": "WARNING",  # '#Second' trapped inside the first grid
    "STR-003": "ERROR",    # mixed header row (Key, Key, 42) on a grid whose columns Mind matches by header
    "INP-005": "ERROR",    # /Reorder without /Input, duplicate headers
    "FLG-001": "WARNING",  # /Inpt is not a documented flag
    "PAR-002": "ERROR",    # 'Kind' instead of 'Type'
    "PRJ-001": "ERROR",    # two /ProjectSettings grids
    "PRJ-006": "ERROR",    # EnableDebugMode: a setting name Mind rejects
    "PRJ-002": "ERROR",    # 'Setting' header, 'sometimes' as Locked
    "UNQ-001": "ERROR",
    "EXP-003": "ERROR",    # unknown column, non-export GridName, bad boolean
    "EXP-004": "WARNING",  # {Bogus} tag
    "RES-002": "ERROR",    # "scenario" vs "Scenario"
    "RES-004": "ERROR",    # same loop twice
    "RZS-002": "ERROR",    # MM_SETSIZE nested in SUM
    "RZS-003": "ERROR",    # 'leftover' in the destination cells
    "RZS-005": "WARNING",  # single-member resize group
    "GRP-001": "ERROR",    # x is not an integer
    "BKP-001": "ERROR",    # no destination, button inside the source grid
    "HID-001": "ERROR",    # 'maybe'
    "SIM-001": "ERROR",    # 'lots'
    "TRN-001": "ERROR",    # 'english'
    "RNG-001": "ERROR",    # MM_RANGE shares its grid with AA5
    "LKP-002": "ERROR",    # range starts on a data row
    "LKP-003": "ERROR",    # 'Result' is not a header
    "LKP-005": "WARNING",  # plain MM_READTABLE
}


def test_broken_model_fails_each_kb_backed_validator(tmp_path, engine, config, flagged_model_broken_xlsx):
    analysis = build_analysis(flagged_model_broken_xlsx, tmp_path / "work", "bad")
    actual = {}
    for rid, expected in BROKEN_EXPECTATIONS.items():
        f = _finding(engine, config, analysis, rid)
        actual[rid] = (f["status"], f["message"][:160])
    mismatches = {rid: actual[rid] for rid, expected in BROKEN_EXPECTATIONS.items() if actual[rid][0] != expected}
    assert mismatches == {}


def test_par_003_and_par_004_report_bad_types_and_ranges(tmp_path, engine, config, flagged_model_broken_xlsx):
    analysis = build_analysis(flagged_model_broken_xlsx, tmp_path / "work", "par")
    # PAR-002 already fails on the header; PAR-003/004 still inspect the rows they can.
    t = _finding(engine, config, analysis, "PAR-003")
    assert t["status"] == "ERROR" and "numbre" in str(t["observed"])
    p = _finding(engine, config, analysis, "PAR-004")
    assert p["status"] == "ERROR"


def test_ref_001_flags_broken_defined_name_references(tmp_path, engine, config):
    """REF-001: a cell referencing a #REF! defined name is an ERROR blocker."""
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "#Data"
    ws["A2"] = "x"
    ws["A3"] = "=COLUMNS(NB_Breakdown)"
    dn = DefinedName("NB_Breakdown", attr_text="#REF!")
    try:
        wb.defined_names["NB_Breakdown"] = dn
    except Exception:
        wb.defined_names.add(dn)
    p = tmp_path / "broken_name.xlsx"
    wb.save(p)
    wb.close()
    analysis = build_analysis(p, tmp_path / "w", "b")
    f = _finding(engine, config, analysis, "REF-001")
    assert f["status"] == "ERROR"
    assert "NB_Breakdown" in str(f["observed"]["broken_names"])
    assert any(s["cell"] == "A3" for s in f["observed"]["sites"])


def test_a_native_function_the_mind_documentation_does_not_list_is_a_warning_not_a_blocker(tmp_path, engine, config):
    """The KB page names what Mind supports and says nothing about the rest:
    a function missing from it is undocumented, not refused. Only an MM_
    name missing from the MM_ registry stays an error."""
    import openpyxl

    from app.readiness import level_of

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["A1"] = "12"
    ws["B1"] = "=_xlfn.NUMBERVALUE(A1)"
    ws["B2"] = "=_xlfn.IFNA(A1,0)+INDIRECT(\"A1\")"
    p = tmp_path / "undocumented.xlsx"
    wb.save(p)
    wb.close()
    f = _finding(engine, config, build_analysis(p, tmp_path / "w", "u"), "FRM-002")
    assert f["status"] == "WARNING" and level_of(f) == "optional"
    assert f["observed"]["undocumented_native_functions"] == ["NUMBERVALUE"]
    assert set(f["observed"]["accepted_beyond_kb_list"]) == {"IFNA", "INDIRECT"}
    assert "neither documented as supported nor as refused" in f["message"] and "last updated 2021-11-17" in f["message"]
    ws_mm = openpyxl.Workbook()
    ws_mm.active["A1"] = "=MM_TOTALLYMADEUP(1)+_xlfn.NUMBERVALUE(\"1\")"
    p2 = tmp_path / "mm.xlsx"
    ws_mm.save(p2)
    ws_mm.close()
    assert _finding(engine, config, build_analysis(p2, tmp_path / "w2", "m"), "FRM-002")["status"] == "ERROR"


def test_every_function_accepted_beyond_the_kb_page_names_its_evidence():
    from app.validators.formula import _confirmed, confirmed_functions

    entries = confirmed_functions()
    assert {"FILTER", "IFNA", "INDIRECT", "XLOOKUP"} <= set(entries)
    assert all(e["evidence"] in ("MIND_CONVERSION", "USER_PROVIDED") and e.get("detail") and e.get("date") for e in entries.values())
    assert _confirmed()["kb_page"]["says_about_unlisted_functions"] == "nothing"


def test_filter_is_treated_as_mind_supported():
    """FILTER is a native function the KB scrape omits but real Mind conversion accepts
    (proven by uploading the Shlomo model); FRM-002 must not flag it as unsupported."""
    from app.validators.formula import _supported_native_functions, MIND_CONFIRMED_SUPPORTED

    supported = _supported_native_functions()
    assert "FILTER" in supported
    assert MIND_CONFIRMED_SUPPORTED <= supported
    assert len(supported) > 100  # the KB list still loads


def test_rep_001_parse_total_recognises_pure_totals():
    from app.validators.structure import _parse_total

    assert _parse_total("=SUM(A1:A10)") == ("sum", "A1:A10")
    assert _parse_total("=SUM($B$2:$B$5)") == ("sum", "B2:B5")
    assert _parse_total("=A1+A2+A3") == ("chain", ["A1", "A2", "A3"])
    # not pure totals
    assert _parse_total("=SUM(A1:A10)+B1") is None
    assert _parse_total("=SUM(Sheet2!A1:A10)") is None
    assert _parse_total("=A1*2") is None
    assert _parse_total("=A1") is None


# --- 1.7.3: two refusals Mind reported on a real model that the app had let through ------------
def test_a_name_that_does_not_exist_is_a_broken_reference(tmp_path, engine, config):
    """Mind: "No function found ... (not a function : Semi_dynamic_increase_rates_array)"
    -- the formula referred to a name the workbook does not define (a misspelling
    of an existing one). Names of any alphabet, '?' in a name, table names,
    sheet-qualified references and LET variables are not names to flag."""
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["A1"], ws["A2"], ws["A3"] = "h", 1, 2
    for name, ref in (("Semi_dynamic_increase_rates", "Calc!$A$2:$A$3"), ("Rate_λG", "Calc!$A$2"), ("Full_Run?", "Calc!$A$3")):
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    ws["C1"], ws["D1"], ws["C2"], ws["D2"] = "K", "V", 1, 2
    ws.add_table(Table(displayName="Rates", ref="C1:D2"))
    ws["F1"] = "=HLOOKUP(1,Semi_dynamic_increase_rates_array,1,FALSE)"  # misspelt
    ws["F2"] = "=Rate_λG+Full_Run?+SUM(Rates[V])+ROWS(Rates)+Calc!A2+TRUE"  # all fine
    ws["F3"] = "=LET(x,A2,x*Nope)"  # LET variables: not workbook names (RSK-004)
    p = tmp_path / "names.xlsx"
    wb.save(p)
    wb.close()
    f = _finding(engine, config, build_analysis(p, tmp_path / "w", "n"), "REF-001")
    assert f["status"] == "ERROR"
    assert f["observed"]["undefined_names"] == {"Semi_dynamic_increase_rates_array": {"cells": 1, "closest": "Semi_dynamic_increase_rates"}}
    assert [s["cell"] for s in f["observed"]["sites"]] == ["F1"]
    assert "closest existing name: Semi_dynamic_increase_rates" in f["message"] and "#NAME?" in f["message"]


def test_a_merged_header_spans_its_columns_and_an_array_past_it_exceeds_the_grid(tmp_path, engine, config):
    """Mind: "The array formula on sheet Temp, cell F5 exceeds the grid size. Ensure
    that the column headers cover the entire array formula." The header was one
    cell merged over B4:E4; the arrays ran to DZ. Read as one cell, the merge
    hid the grid boundary the converter applies."""
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula

    def model(merge_to: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Temp"
        ws["B4"] = "Per Policy Projection"
        ws.merge_cells(f"B4:{merge_to}4")
        for r, label in ((5, "ME charge"), (6, "Admin cost")):
            ws[f"B{r}"] = label
            ws[f"C{r}"] = ArrayFormula(f"C{r}:H{r}", "=TRANSPOSE($A$10:$A$15)")
        for r in range(10, 16):
            ws[f"A{r}"] = r
        p = tmp_path / f"merged_{merge_to}.xlsx"
        wb.save(p)
        wb.close()
        return build_analysis(p, tmp_path / f"w_{merge_to}", "m")

    short = model("E")
    grids = [(g["display_name"], g["ref"]) for g in short["workbooks"][0]["sheets"][0]["grids"]]
    assert ("untitled B4", "B4:E6") in grids  # the merge counts for B..E
    f = _finding(engine, config, short, "RSK-002")
    assert f["status"] == "ERROR"
    first = f["observed"]["oversized_array_formulas"][0]
    assert (first["cell"], first["array_ref"], first["first_cell_outside"]) == ("C5", "C5:H5", "F5")
    assert "column headers cover the entire array formula" in f["message"]
    wide = model("H")  # headers cover every column of the arrays: nothing exceeds
    assert _finding(engine, config, wide, "RSK-002")["status"] == "PASS"


# --- 1.7.3: two refusals from Mind's runner / compiler, detected before upload -----------------
def test_circular_references_are_found_and_reported_as_chains(tmp_path, engine, config):
    """Mind: "Run error: Circular reference found" (WGM input!I11). A cycle through
    cells, a range and a defined name is found; an acyclic model passes."""
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WGM input"
    ws["A1"] = "h"
    ws["I11"] = "=J11*2"  # I11 -> J11 -> K11 -> (range) I11 : a 3-cell cycle
    ws["J11"] = "=K11+1"
    ws["K11"] = "=SUM(I10:I12)"
    ws["B2"] = "=Total+1"  # B2 -> name Total -> C2 -> B2 : a cycle through a defined name
    ws["C2"] = "=B2*3"
    wb.defined_names["Total"] = DefinedName("Total", attr_text="'WGM input'!$C$2")
    ws["D2"] = "=D2"  # a cell reading itself
    ws["E2"] = "=SUM(A1:A5)+INDIRECT(\"F2\")"  # no static cycle; the dynamic target is not followed
    other = wb.create_sheet("Calc")
    other["A1"] = "='WGM input'!I11"  # reads a cycle but is not part of it
    # ROW($A10) uses A10 as a coordinate, not its value: A10 <-> B10 is not a cycle (a real model's pattern)
    ws["A10"] = "=IF(B10=1,1,0)"
    ws["B10"] = '=INDIRECT("Cluster"&(ROW($A10)-ROW($A$7)))'
    # CI14 <-> CJ14 through IF branches that never evaluate together (a real model's pattern): a conditional cycle
    ws["F1"] = "=IF($G$9=1,H1*2,IFERROR(H1/G1,0))"
    ws["G1"] = "=IF(H1=0,0,IF($G$9=1,H1/F1,H1))"
    ws["H1"] = 3
    # OFFSET($G452,,n) in G452 starts from G452 without reading it (a real model's pattern): not a cycle
    ws["G5"] = "=OFFSET($G5,,$H$1+1)"
    wb.calculation.iterate = True
    p = tmp_path / "circular.xlsx"
    wb.save(p)
    wb.close()
    f = _finding(engine, config, build_analysis(p, tmp_path / "w", "c"), "FRM-005")
    assert f["status"] == "ERROR"
    obs = f["observed"]
    assert obs["cycle_count"] == 4 and obs["cells_in_cycles"] == 8 and obs["iterative_calculation"] is True
    assert obs["unconditional_cycles"] == 3 and obs["conditional_cycles"] == 1
    assert [c["unconditional"] for c in obs["cycles"]] == [True, True, True, False]
    assert obs["cycles"][-1]["chain"] == ["WGM input!F1", "WGM input!G1", "WGM input!F1"]
    assert "1 more cycle(s) close only through an IF/IFERROR/CHOOSE branch" in f["message"]
    chains = {tuple(c["chain"]) for c in obs["cycles"]}
    assert ("WGM input!I11", "WGM input!J11", "WGM input!K11", "WGM input!I11") in chains
    assert ("WGM input!B2", "WGM input!C2", "WGM input!B2") in chains
    assert ("WGM input!D2", "WGM input!D2") in chains
    assert "WGM input!A10" not in str(chains) and "WGM input!G5" not in str(chains)
    assert obs["dynamic_reference_formulas_not_followed"] == 3
    assert "Circular reference found" in f["message"] and "MM_ITERATIONS" in f["message"]
    assert "Calc!A1" not in str(chains)

    clean = openpyxl.Workbook()
    ws = clean.active
    ws.title = "S"
    ws["A1"], ws["A2"], ws["A3"] = 1, "=A1*2", "=SUM(A1:A2)"
    p2 = tmp_path / "acyclic.xlsx"
    clean.save(p2)
    clean.close()
    assert _finding(engine, config, build_analysis(p2, tmp_path / "w2", "a"), "FRM-005")["status"] == "PASS"

    # a cell reading its own address on a run-time sheet (INDIRECT + ADDRESS(ROW(),COLUMN())): Mind stopped there
    own = openpyxl.Workbook()
    ws = own.active
    ws.title = "WGM input"
    ws["A1"] = "Run A"
    ws["I11"] = '=INDIRECT("'"&$A$1&"'!"&ADDRESS(ROW(),COLUMN()),TRUE)'
    ws["I12"] = '=INDIRECT("'"&$A$1&"'!"&ADDRESS(ROW(),COLUMN()),TRUE)'
    own.create_sheet("Run A")["I11"] = 5
    p4 = tmp_path / "self_address.xlsx"
    own.save(p4)
    own.close()
    e = _finding(engine, config, build_analysis(p4, tmp_path / "w4", "o"), "FRM-005")
    assert e["status"] == "ERROR" and e["observed"]["self_addressing_indirect"] == {"count": 2, "by_sheet": {"WGM input": 2}, "first": e["observed"]["self_addressing_indirect"]["first"]}
    assert e["location"] == {"sheet": "WGM input", "cell": "I11"} and "read their own address" in e["message"]

    # only a conditional cycle: a warning, not a blocker -- Excel computes it, Mind's answer is not documented
    soft = openpyxl.Workbook()
    ws = soft.active
    ws.title = "S"
    ws["A1"], ws["B1"], ws["C1"] = "=IF($C$1=1,B1*2,0)", "=IF($C$1=1,0,A1+1)", 1
    p3 = tmp_path / "conditional.xlsx"
    soft.save(p3)
    soft.close()
    w = _finding(engine, config, build_analysis(p3, tmp_path / "w3", "s"), "FRM-005")
    assert w["status"] == "WARNING" and w["observed"]["conditional_cycles"] == 1 and "not documented" in w["message"]


def test_three_d_references_are_flagged_with_the_explicit_formula_proposed(tmp_path, engine, config):
    """Mind: "Sheet '>> Reporting:>>>' not found in workbook ... on compiling formula:
    SUM('>> Reporting:>>>'!RC)". The sheets between the two tabs are listed and the
    per-sheet formula is proposed; nothing is written."""
    import openpyxl

    wb = openpyxl.Workbook()
    first = wb.active
    first.title = ">> Reporting"
    for name in ("LoB A", "LoB B", ">>>"):
        wb.create_sheet(name)
    total = wb.create_sheet("Reporting_LoB Total")
    for ws in wb.worksheets[:4]:
        ws["E5"] = 1
    total["E5"] = "=SUM('>> Reporting:>>>'!E5)"
    total["E19"] = "=AVERAGE('>> Reporting:>>>'!E5:E7)/2"
    total["G37"] = "=SUM('>> Reporting:Nope'!G37)"  # an endpoint that does not exist
    total["H1"] = "=SUM('LoB A'!E5)"  # an ordinary sheet reference: not 3-D
    p = tmp_path / "3d.xlsx"
    wb.save(p)
    wb.close()
    f = _finding(engine, config, build_analysis(p, tmp_path / "w", "d"), "FRM-006")
    assert f["status"] == "ERROR"
    sites = {s["cell"]: s for s in f["observed"]["sites"]}
    assert set(sites) == {"E5", "E19", "G37"}
    assert sites["E5"]["references"][0]["sheets"] == [">> Reporting", "LoB A", "LoB B", ">>>"]
    assert sites["E5"]["suggested_formula"] == "=SUM('>> Reporting'!E5,'LoB A'!E5,'LoB B'!E5,'>>>'!E5)"
    assert sites["E19"]["suggested_formula"] == "=AVERAGE('>> Reporting'!E5:E7,'LoB A'!E5:E7,'LoB B'!E5:E7,'>>>'!E5:E7)/2"
    assert sites["G37"]["suggested_formula"] is None and "Nope" in sites["G37"]["issue"]
    assert "not found in workbook" in f["message"] and "No automatic repair" in f["message"]
