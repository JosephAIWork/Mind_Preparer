"""1.8.0: FMT-002 is done in the workbook's style table (app/theme_colors.py).

Going through the cells in Excel never finished (Excel keeps the default text
colour as a theme colour when a cell is assigned the colour it already shows,
so every Prep round planned the same cells again) and took an hour for one
124,000-cell sheet. The style table has a few hundred fonts and fills."""
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color

from app import prep
from app.excel_com import close_quietly, com_available, excel_session
from app.inventory import _color_is_theme
from app.theme_colors import explicit_styles, explicit_theme_colors_in_package, theme_palette, tinted

needs_excel = pytest.mark.skipif(not com_available(), reason="needs a real Excel")

# What Excel itself showed (Interior.Color) for a theme colour with a tint.
EXCEL_SAYS = [
    ("FFFFFF", -0.249977111117893, "BFBFBF"), ("FFFFFF", -0.499984740745262, "808080"), ("FFFFFF", 0.33, "FFFFFF"),
    ("000000", 0.7999816888943144, "CCCCCC"), ("000000", 0.3999755851924192, "666666"), ("000000", -0.499984740745262, "000000"),
    ("000000", 0.8999908444471572, "E6E6E6"), ("000000", 0.0499893185216834, "0D0D0D"), ("000000", 0.33, "555555"),
    ("EEECE1", 0.7999816888943144, "FBFBF9"), ("EEECE1", -0.249977111117893, "C4BD97"), ("EEECE1", -0.499984740745262, "948A54"),
    ("EEECE1", 0.0499893185216834, "EFEDE2"), ("1F497D", 0.7999816888943144, "C5D9F1"), ("1F497D", 0.3999755851924192, "538DD5"),
    ("1F497D", -0.249977111117893, "16365C"), ("1F497D", 0.8999908444471572, "E2ECF8"), ("1F497D", 0.33, "4081D0"),
    ("4F81BD", 0.7999816888943144, "DCE6F1"), ("4F81BD", 0.3999755851924192, "95B3D7"), ("4F81BD", -0.249977111117893, "366092"),
    ("4F81BD", -0.499984740745262, "244062"), ("4F81BD", 0.0499893185216834, "5686C0"), ("4F81BD", 0.33, "89ABD3"),
    ("9BBB59", 0.7999816888943144, "EBF1DE"), ("9BBB59", 0.3999755851924192, "C4D79B"), ("9BBB59", -0.249977111117893, "76933C"),
    ("9BBB59", 0.8999908444471572, "F5F9EE"), ("0000FF", 0.7999816888943144, "CCCCFF"), ("0000FF", -0.249977111117893, "0000BF"),
    ("0000FF", 0.0499893185216834, "0D0DFF"), ("800080", 0.7999816888943144, "FFB3FF"), ("800080", 0.3999755851924192, "FF1AFF"),
    ("800080", -0.499984740745262, "400040"), ("800080", 0.33, "FF00FF"),
    # the Office 2013+ blue, whose shades everybody has seen in a table
    ("5B9BD5", 0.7999816888943144, "DDEBF7"), ("5B9BD5", 0.5999938962981048, "BDD7EE"), ("5B9BD5", 0.3999755851924192, "9BC2E6"),
    ("5B9BD5", -0.249977111117893, "2F75B5"), ("5B9BD5", -0.499984740745262, "1F4E78"),
]


def test_a_tint_gives_the_colour_excel_shows():
    assert [(base, tint, tinted(base, tint)) for base, tint, _ in EXCEL_SAYS] == EXCEL_SAYS
    assert tinted("4f81bd", 0.0) == "4F81BD"       # no tint: the theme colour itself
    assert tinted("123456", 1.0) == "FFFFFF" and tinted("123456", -1.0) == "000000"


THEME = """<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:themeElements><a:clrScheme name="Office">
<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1><a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
<a:dk2><a:srgbClr val="1F497D"/></a:dk2><a:lt2><a:srgbClr val="EEECE1"/></a:lt2>
<a:accent1><a:srgbClr val="4F81BD"/></a:accent1><a:accent2><a:srgbClr val="C0504D"/></a:accent2><a:accent3><a:srgbClr val="9BBB59"/></a:accent3>
<a:accent4><a:srgbClr val="8064A2"/></a:accent4><a:accent5><a:srgbClr val="4BACC6"/></a:accent5><a:accent6><a:srgbClr val="F79646"/></a:accent6>
<a:hlink><a:srgbClr val="0000FF"/></a:hlink><a:folHlink><a:srgbClr val="800080"/></a:folHlink></a:clrScheme></a:themeElements></a:theme>"""


def test_the_palette_is_in_the_order_cells_number_it():
    palette = theme_palette(THEME)
    assert palette[:5] == ["FFFFFF", "000000", "EEECE1", "1F497D", "4F81BD"]  # 0 = light background, 1 = dark text
    assert palette[10:] == ["0000FF", "800080"] and len(palette) == 12
    assert theme_palette(THEME.replace(' lastClr="000000"', "")) is not None  # a system colour without its last value
    assert theme_palette("<a:theme/>") is None                                # no scheme: the workbook is left alone


STYLES = """<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<fonts count="4"><font><sz val="11"/><color theme="1"/><name val="Calibri"/></font>
<font><b/><color theme="4" tint="-0.249977111117893"/></font><font><color rgb="FF112233"/></font><font><color indexed="8"/></font></fonts>
<fills count="3"><fill><patternFill patternType="none"/></fill>
<fill><patternFill patternType="solid"><fgColor theme="4" tint="0.79998168889431442"/><bgColor indexed="64"/></patternFill></fill>
<fill><patternFill patternType="solid"><fgColor theme="0"/><bgColor theme="99"/></patternFill></fill></fills>
<borders count="1"><border><left style="thin"><color theme="4"/></left></border></borders>
<dxfs count="1"><dxf><font><color theme="5"/></font><fill><patternFill><bgColor theme="6"/></patternFill></fill></dxf></dxfs>
</styleSheet>"""


def test_fonts_and_fills_of_the_style_table_become_rgb_and_the_rest_is_left_byte_for_byte():
    out, counts = explicit_styles(STYLES, theme_palette(THEME))
    assert counts == {"fonts": 2, "fills": 2}
    assert '<font><sz val="11"/><color rgb="FF000000"/><name val="Calibri"/></font>' in out      # the default text colour
    assert '<font><b/><color rgb="FF366092"/></font>' in out                                     # theme + tint
    assert '<color rgb="FF112233"/>' in out and '<color indexed="8"/>' in out                   # explicit already
    assert '<fgColor rgb="FFDCE6F1"/><bgColor indexed="64"/>' in out
    assert '<fgColor rgb="FFFFFFFF"/><bgColor theme="99"/>' in out                               # a theme index nobody has: left
    tail = STYLES[STYLES.index("<borders"):]
    assert out.endswith(tail)                                                                   # borders, conditional formats: not this rule
    assert explicit_styles(out, theme_palette(THEME)) == (out, {"fonts": 0, "fills": 0})         # once is enough


def _book(path: Path) -> list[str]:
    """Cells in the default text colour (the case that never converged), theme
    fills with tints, a colour that is explicit already, a theme-coloured border."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    cells = []
    for r in range(2, 10):
        for c in range(2, 7):
            cell = ws.cell(r, c, r * c)
            cell.font = Font(color=Color(theme=1), bold=(r == 2))
            if c != 4:
                cell.fill = PatternFill("solid", fgColor=Color(theme=4 + c % 3, tint=0.5999938962981048 if r % 2 else -0.249977111117893))
            cells.append(cell.coordinate)
    ws["D5"].font = Font(color="FF112233")
    ws["B2"].border = Border(left=Side(style="thin", color=Color(theme=5)))
    ws["H1"] = "=SUM(B2:F9)"
    wb.create_sheet("Other")["A1"] = "text"
    wb["Other"]["A1"].font = Font(color=Color(theme=3))
    wb.save(path)
    wb.close()
    return cells


def _theme_cells(path: Path) -> list[str]:
    """The cells FMT-002 would flag (the detection of app/inventory.py)."""
    wb = openpyxl.load_workbook(path)
    try:
        found = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None or not c.has_style:
                        continue
                    fill = c.fill
                    if _color_is_theme(c.font.color if c.font is not None else None) or (fill is not None and fill.fill_type is not None and _color_is_theme(fill.fgColor)):
                        found.append(f"{ws.title}!{c.coordinate}")
        return found
    finally:
        wb.close()


def test_only_the_style_table_of_the_package_changes(tmp_path):
    path = tmp_path / "book.xlsx"
    cells = _book(path)
    assert sorted(_theme_cells(path)) == sorted([f"Calc!{c}" for c in cells if c != "D5"] + ["Other!A1"])  # D5: explicit already, no fill
    with zipfile.ZipFile(path) as z:
        before = {n: z.read(n) for n in z.namelist()}
    counts = explicit_theme_colors_in_package(path)
    assert counts["fonts"] >= 2 and counts["fills"] >= 4
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        after = {n: z.read(n) for n in z.namelist()}
    assert list(after) == list(before)                             # same parts, same order
    assert [n for n in before if before[n] != after[n]] == ["xl/styles.xml"]
    assert _theme_cells(path) == []
    wb = openpyxl.load_workbook(path)
    ws = wb["Calc"]
    assert ws["B2"].font.color.rgb == "FF000000" and ws["B2"].font.bold                      # the default text colour, as RGB
    assert ws["B3"].fill.fgColor.rgb == "FF" + tinted("9BBB59", 0.5999938962981048)          # theme 6 = accent3 (c=2: 4 + 2 % 3)
    assert ws["D5"].font.color.rgb == "FF112233"
    assert ws["B2"].border.left.color.theme == 5                                             # a border is not this rule's business
    assert ws["H1"].value == "=SUM(B2:F9)"
    wb.close()
    again = path.read_bytes()
    assert explicit_theme_colors_in_package(path) == {"fonts": 0, "fills": 0} and path.read_bytes() == again  # nothing left: nothing written


def test_a_file_that_is_not_a_package_is_left_alone(tmp_path):
    other = tmp_path / "notes.xlsx"
    other.write_text("not a zip", encoding="utf-8")
    assert explicit_theme_colors_in_package(other) is None
    assert other.read_text(encoding="utf-8") == "not a zip"


@needs_excel
def test_apply_converts_every_theme_colour_and_keeps_the_look(tmp_path):
    src = tmp_path / "src" / "book.xlsx"
    src.parent.mkdir()
    cells = _book(src)
    shown: dict[str, tuple[int, int]] = {}

    def look(path: Path, into: dict) -> None:
        with excel_session() as excel:
            def go() -> None:
                wb = excel.Workbooks.Open(str(path), 0, True)
                try:
                    ws = wb.Worksheets("Calc")
                    for ref in ("B2", "B3", "C4", "D5", "E6", "F9"):
                        cell = ws.Range(ref)
                        into[ref] = (int(cell.Font.Color), int(cell.Interior.Color))
                        cell = None
                    ws = None
                finally:
                    close_quietly(wb)

            go()

    look(src, shown)
    op = {"op": "explicit_colors", "action_id": "explicit_colors", "rule_id": "FMT-002", "sheet": "Calc", "cells": cells,
          "before": f"{len(cells)} cell(s) with theme colours", "after": "same colours as explicit RGB", "note": "requires Excel"}
    told = []
    res = prep.apply_operations(src, tmp_path / "work", [op], progress=lambda stage, message, fraction=None, **facts: told.append(message))
    assert res["status"] == "APPLIED" and res["method"] == "excel_com" and res["verified_opens_in_excel"] is True
    assert "style table" in res["applied"][0]["note"] and res["change_log_entry"]["applied"][0]["cells"] == len(cells)
    assert any("style table" in m for m in told) and not any("colours of cell" in m for m in told)   # no cell-by-cell work
    out = Path(res["output_path"])
    assert _theme_cells(out) == []                    # the default text colour included, on every sheet -- nothing to plan next round
    after: dict[str, tuple[int, int]] = {}
    look(out, after)
    assert after == shown                             # Excel shows exactly the colours it showed before
    assert _theme_cells(src) != []                    # the source is never modified
