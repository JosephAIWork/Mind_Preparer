"""1.8.0: did the preparation change what the model computes? (app/numbers_check.py)

Found on a real model (PVFP): the standard Prep inserted title rows on sheets a
3-D reference reads by position, and 2,462 computed values changed -- 1,781 of
them into errors. Nothing in the app said so."""
from pathlib import Path

import openpyxl
import pytest

from app.autofix import AutoFixConfig, run_autofix
from app.excel_com import com_available
from app.numbers_check import Moves, compare_workbooks
from app.prep import apply_operations

needs_excel = pytest.mark.skipif(not com_available(), reason="needs a real Excel")


def test_a_cell_is_followed_through_inserts_and_renames_apply_after_apply():
    moves = Moves([
        [{"op": "insert_row", "sheet": "S", "row": 5}, {"op": "insert_row", "sheet": "S", "row": 2}, {"op": "insert_column", "sheet": "S", "column": "B"},
         {"op": "rename_sheet", "sheet": "S", "after": "T"}, {"op": "set_value", "sheet": "S", "cell": "A1", "after": "#x"}],
        [{"op": "insert_row", "sheet": "T", "row": 3}],
        [{"op": "set_formula", "sheet": "T", "cell": "A1", "after": "=1"}],  # moves nothing: not a batch
    ])
    assert moves.summary() == {"row_inserts": 3, "column_inserts": 1, "sheet_renames": 1}
    assert moves.forward("S", 1, 1) == ("T", 1, 1)      # above every insert, left of the column
    assert moves.forward("S", 2, 2) == ("T", 4, 3)      # row 2 -> 3 (insert at 2), -> 4 (second Apply, insert at 3); column B -> C
    assert moves.forward("S", 5, 3) == ("T", 8, 4)      # two inserts at or above it, then one more
    assert moves.forward("Other", 9, 9) == ("Other", 9, 9)


def _model(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, scale in (("LoB 1", 1), ("LoB 2", 10)):
        ws = wb.create_sheet(name)
        ws["A1"], ws["B1"] = "Line", "2024"
        ws["A2"], ws["B2"] = "Premium", 100 * scale
        ws["A3"], ws["B3"] = "Claims", 60 * scale
        ws["A4"], ws["B4"] = "Ratio", "=B3/B2"
    end = wb.create_sheet(">>")
    end["A1"] = "end of the lines of business"
    tot = wb.create_sheet("Total")
    tot["A2"], tot["B2"] = "Premium", "=SUM('LoB 1:>>'!B2)"
    tot["A3"], tot["B3"] = "Claims", "=SUM('LoB 1:>>'!B3)"
    tot["A4"], tot["B4"] = "Ratio", "=B3/B2"
    tot["A6"], tot["B6"] = "Lines", "=COUNT('LoB 1:>>'!B4)"
    wb.save(path)
    wb.close()
    return path


@needs_excel
def test_a_row_inserted_under_a_3d_reference_is_caught_and_a_harmless_one_is_not(tmp_path):
    src = _model(tmp_path / "model.xlsx")
    # a title row on top of 'LoB 1' only: Total still sums B2 and B3 of every sheet -- other lines now, on 'LoB 1'
    bad = apply_operations(src, tmp_path / "bad", [{"op": "insert_row", "sheet": "LoB 1", "row": 1, "action_id": "t", "rule_id": "T"}])
    res = compare_workbooks(src, Path(bad["output_path"]), Moves([bad["applied"]]))
    assert res["ran"] and res["clean"] is False and res["counts"]["good_to_other"] >= 2 and res["moves"]["row_inserts"] == 1
    assert "changed" in res["verdict"] and res["by_sheet"]["Total"]["good_to_other"] >= 2
    moved = {s["original"]: (s["was"], s["is"]) for s in res["samples"]["good_to_other"]}
    assert moved["Total!B3"] == ("660", "700")  # Claims: 60 + 600 before; now 'LoB 1'!B3 is the Premium line (100) + 600
    # the same row on 'Total', which nothing reads by position: every value is where it was, one row lower
    good = apply_operations(src, tmp_path / "good", [{"op": "insert_row", "sheet": "Total", "row": 1, "action_id": "t", "rule_id": "T"}])
    res = compare_workbooks(src, Path(good["output_path"]), Moves([good["applied"]]))
    assert res["clean"] is True and res["counts"]["compared"] == 6 and res["counts"]["same"] == 6 and "no computed value" in res["verdict"]
    # without the change log the cells are compared at the wrong addresses -- the log is what makes the check exact
    assert compare_workbooks(src, Path(good["output_path"]), Moves())["counts"]["missing"] > 0


@needs_excel
def test_an_error_the_preparation_made_is_not_given_a_fallback_by_the_fixer(tmp_path):
    src = _model(tmp_path / "model.xlsx")
    # "Prep" empties the premium of LoB 1 (stands for any change that breaks a formula): its ratio becomes #DIV/0!
    broken = apply_operations(src, tmp_path / "broken", [{"op": "set_value", "sheet": "LoB 1", "cell": "B2", "after": 0, "action_id": "t", "rule_id": "T"}])
    out = Path(broken["output_path"])
    numbers = compare_workbooks(src, out, Moves([broken["applied"]]))
    assert numbers["counts"]["good_to_error"] == 1 and numbers["regressions"] == [("LoB 1", 4, 2, 0.6)]
    res = run_autofix(out, tmp_path / "fix", AutoFixConfig(regressions=numbers["regressions"]))
    assert res["status"] == "stuck" and res["regression_cells"] == 1 and res["errors_after"] == 1 and res["output_path"] is None
    assert "it was 0.6 in the original workbook" in res["left"][0]["why"] and "preparation broke" in res["summary"]
    # the same workbook taken as it is, with no original to compare with: the fixer has no reason to hold back
    assert run_autofix(out, tmp_path / "fix2")["status"] == "clean"


def test_a_value_that_only_holds_the_workbooks_own_folder_is_not_a_change():
    """=CELL("filename") and the paths built on it differ between the original and
    its copy only because they sit in different folders (Palermo: 'Saved down'!C2)."""
    from app.numbers_check import _without_location

    original = Path(r"C:\Users\x\usecases\palermo.xlsm")
    copy = Path(r"C:\Temp\mind_ready_b1\apply\v3\palermo.xlsm")
    assert _without_location("C:\\Users\\x\\usecases\\", original) == _without_location("C:\\Temp\\mind_ready_b1\\apply\\v3\\", copy)
    assert _without_location(r"C:\Users\x\usecases\[palermo.xlsm]Control", original) == _without_location(r"C:\Temp\mind_ready_b1\apply\v3\[palermo.xlsm]Control", copy)
    assert _without_location("Premium", original) == "Premium" and _without_location(12.5, original) == 12.5
    assert _without_location("Total 2024", original) != _without_location("Total 2025", copy)
